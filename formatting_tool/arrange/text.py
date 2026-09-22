"""Aligning the text INSIDE an element, as opposed to where the element sits.

Section 9 of the pipeline document, and the subtlest part of the original for
a reason that survives the port intact: "the text of a shape" is not one
thing. A table is a shape whose text lives on cells, a mixed run has no single
font, and a paragraph's alignment is a different setting from the anchor that
decides which end of the box the paragraph starts at.

FOUR SETTINGS, AND THEY ARE NOT INTERCHANGEABLE.

  vertical anchor     which end of the box the block of text sits against.
                      `a:bodyPr/@anchor` on a shape, `a:tcPr/@anchor` on a
                      table cell -- two different attributes, and writing the
                      shape one on a cell changes nothing PowerPoint reads.
  horizontal anchor   whether the block is centred across the box's width
                      rather than filling it. `a:bodyPr/@anchorCtr`. Almost
                      nobody sets it and it is invisible until somebody does.
  insets              how far in from the box edge the copy starts. The
                      setting `space.text_insets` reports on: a row of cards
                      whose boxes are exactly right and whose copy is not.
  paragraph alignment flush left, centred, flush right, justified. Per
                      paragraph, which is why a bilingual shape can be right
                      about one line and wrong about the next.

TWO TRAVERSALS, AND THEY DIFFER ON PURPOSE. Reading samples: the first cell of
a table stands in for the table, and the first run of the first paragraph
stands in for the shape, because a range-level font on a mixed run answers
with nothing at all and a first character answers with something comparable.
Writing visits everything: every cell, every paragraph. That asymmetry is the
C#'s and it is what makes "make these the same" idempotent -- read one value,
write it everywhere, read the same value back.
"""

from __future__ import annotations

import logging
from typing import Any, Iterator, Optional, Sequence

from .actions import ArrangeResult
from .elements import EMU_PER_INCH

log = logging.getLogger(__name__)

# The words this tool speaks. `ShapeProfile.vertical_anchor` and
# `ParagraphProfile.alignment` already carry these, so an engine that spoke
# python-pptx enums would need translating at every call site inside the tool
# and at none outside it.
ANCHORS = ("top", "middle", "bottom")
ALIGNMENTS = ("left", "center", "right", "justify", "distribute")

_ANCHOR_WORDS = {"TOP": "top", "MIDDLE": "middle", "CENTER": "middle",
                 "CENTRE": "middle", "BOTTOM": "bottom"}
_ALIGN_WORDS = {"LEFT": "left", "CENTER": "center", "CENTRE": "center",
                "RIGHT": "right", "JUSTIFY": "justify",
                "DISTRIBUTE": "distribute"}


def _safe(getter, default=None):
    """A read that is allowed to fail.

    Every read here goes through this, for the reason the C# gives for its
    bare `catch {}`: a property a given shape type does not expose is an
    expected condition and not an error. A picture has no text frame and a
    group has no paragraphs.
    """
    try:
        return getter()
    except Exception:
        return default


# --------------------------------------------------------------------------- #
# Traversal
# --------------------------------------------------------------------------- #

def table_of(shape: Any) -> Any:
    """The table on a graphic frame, or None. C#: shape.HasTable."""
    if not _safe(lambda: shape.has_table, False):
        return None
    return _safe(lambda: shape.table)


def text_frames(shape: Any) -> list:
    """Every text frame this shape writes through, in row-major order.

    One for an ordinary shape. One per cell for a table, because a table
    restyled cell by cell is the only table that comes out uniform -- the C#
    walks `Cell(row, col).Shape.TextFrame2` for exactly this, and a fixer that
    wrote the graphic frame instead would change nothing at all.
    """
    table = table_of(shape)
    if table is None:
        frame = _safe(lambda: shape.text_frame)
        return [frame] if frame is not None else []
    return [cell.text_frame for cell in _cells(table)]


def frame_targets(shape: Any) -> list:
    """Whatever carries the anchor and the insets for this shape.

    NOT THE SAME OBJECTS AS `text_frames`, and the difference is a real bug
    waiting to happen. A table cell's anchor and insets live on `a:tcPr` --
    `@anchor`, `@marL`, `@marR`, `@marT`, `@marB` -- and python-pptx exposes
    them on the cell. The cell's own `text_frame` has a `a:bodyPr` too, it
    accepts the same writes, and PowerPoint ignores every one of them.
    """
    table = table_of(shape)
    if table is None:
        frame = _safe(lambda: shape.text_frame)
        return [frame] if frame is not None else []
    return list(_cells(table))


def _cells(table: Any) -> Iterator[Any]:
    """Every cell of a table, row-major, skipping the ones a merge covers."""
    for row in _safe(lambda: list(table.rows), []) or []:
        for cell in _safe(lambda: list(row.cells), []) or []:
            if _safe(lambda: cell.is_spanned, False):
                continue
            yield cell


def paragraphs(shape: Any) -> list:
    """Every paragraph in every frame, which is the write-side set."""
    found: list = []
    for frame in text_frames(shape):
        found.extend(_safe(lambda frame=frame: list(frame.paragraphs), []) or [])
    return found


def primary_frame(shape: Any) -> Any:
    """The frame that stands in for the shape when reading. C#: GetPrimaryTextFrame."""
    frames = text_frames(shape)
    return frames[0] if frames else None


def primary_target(shape: Any) -> Any:
    """The anchor/inset carrier that stands in for the shape when reading."""
    targets = frame_targets(shape)
    return targets[0] if targets else None


def primary_paragraph(shape: Any) -> Any:
    """The first paragraph that has any copy in it, or the first at all."""
    frame = primary_frame(shape)
    if frame is None:
        return None
    items = _safe(lambda: list(frame.paragraphs), []) or []
    for paragraph in items:
        if (_safe(lambda p=paragraph: p.text, "") or "").strip():
            return paragraph
    return items[0] if items else None


def first_run(shape: Any) -> Any:
    """The first run of the first paragraph with copy in it.

    THE FIRST-CHARACTER SAMPLE, and the reason the C# narrows to
    `Characters[1,1]` before reading a font at all. A range spanning runs in
    two typefaces answers `Font.Name` with an empty string, so a comparison
    built on it says two shapes differ when neither of them has a font worth
    differing about. One character always answers, and it answers the same way
    every time, which is the whole requirement for matching.

    The cost is written down in the C# too and it is real: two shapes whose
    first characters agree and whose later runs do not will compare as equal.
    """
    paragraph = primary_paragraph(shape)
    if paragraph is None:
        return None
    for run in _safe(lambda: list(paragraph.runs), []) or []:
        if (_safe(lambda r=run: r.text, "") or "").strip():
            return run
    runs = _safe(lambda: list(paragraph.runs), []) or []
    return runs[0] if runs else None


# --------------------------------------------------------------------------- #
# Reading
# --------------------------------------------------------------------------- #

def anchor_of(shape: Any) -> Optional[str]:
    """Where the text sits in the box, as one word, or None where unstated.

    None is not "top" even though PowerPoint draws it that way. A caller
    deciding whether it is safe to grow a box wants the difference between a
    shape that says top and a shape that never said -- the same distinction
    `extract.deck_reader._anchor` keeps, and for the same caller.
    """
    target = primary_target(shape)
    if target is None:
        return None
    return _anchor_word(_safe(lambda: target.vertical_anchor))


def anchor_centred(shape: Any) -> Optional[bool]:
    """`a:bodyPr/@anchorCtr`: is the block centred across the box's width."""
    body = _body_pr(primary_frame(shape))
    if body is None:
        return None
    raw = body.get("anchorCtr")
    return None if raw is None else raw in ("1", "true")


def insets_of(shape: Any) -> Optional[tuple]:
    """The four insets in inches: left, top, right, bottom.

    The order `ShapeProfile.text_margins` uses, so a value read here can be
    compared with one the extractor read without either end reordering it.
    """
    target = primary_target(shape)
    if target is None:
        return None
    values = []
    for side in ("margin_left", "margin_top", "margin_right", "margin_bottom"):
        found = _safe(lambda side=side: getattr(target, side))
        if found is None:
            return None
        values.append(round(int(found) / EMU_PER_INCH, 4))
    return tuple(values)


def alignment_of(shape: Any) -> Optional[str]:
    """The first paragraph's alignment as a word, or None where unstated."""
    paragraph = primary_paragraph(shape)
    if paragraph is None:
        return None
    return _align_word(_safe(lambda: paragraph.alignment))


def reading_order_of(shape: Any) -> Optional[bool]:
    """`a:pPr/@rtl` on the first paragraph: which way the line runs.

    Not a nicety on an Arabic deck and not the same thing as the font. The
    letters shape and join whatever this says; what it decides is where the
    full stop, the brackets and the numbers go, which is why a deck missing it
    looks very nearly right.
    """
    paragraph = primary_paragraph(shape)
    if paragraph is None:
        return None
    properties = _safe(lambda: paragraph._p.find(_qn("a:pPr")))
    if properties is None:
        return None
    raw = properties.get("rtl")
    return None if raw is None else raw in ("1", "true")


# --------------------------------------------------------------------------- #
# Writing
# --------------------------------------------------------------------------- #

def set_anchor(shape: Any, anchor: str) -> int:
    """Anchor the text top, middle or bottom. Returns frames changed."""
    from pptx.enum.text import MSO_ANCHOR  # noqa: PLC0415 - lazy, optional dep

    wanted = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
              "bottom": MSO_ANCHOR.BOTTOM}.get(_word(anchor, _ANCHOR_WORDS))
    if wanted is None:
        return 0

    changed = 0
    for target in frame_targets(shape):
        if _safe(lambda t=target: t.vertical_anchor) == wanted:
            continue
        try:
            target.vertical_anchor = wanted
        except Exception:
            log.debug("could not anchor %r", shape, exc_info=True)
            continue
        changed += 1
    return changed


def set_anchor_centred(shape: Any, centred: bool) -> int:
    """Centre the text block across the box's width, or stop centring it."""
    changed = 0
    for frame in text_frames(shape):
        body = _body_pr(frame)
        if body is None:
            continue
        was = body.get("anchorCtr")
        wanted = "1" if centred else "0"
        if was == wanted:
            continue
        body.set("anchorCtr", wanted)
        changed += 1
    return changed


def set_insets(shape: Any, insets_in: Sequence[float]) -> int:
    """Set the four insets, in inches, left/top/right/bottom.

    Table-aware, which the fixer this replaces was not: a table's insets are
    on its cells, so setting them on the graphic frame wrote a value nothing
    reads and reported a correction that had not happened.
    """
    if insets_in is None or len(insets_in) != 4:
        return 0
    from pptx.util import Emu  # noqa: PLC0415 - lazy, optional dep

    left, top, right, bottom = (
        Emu(int(round(float(value) * EMU_PER_INCH))) for value in insets_in
    )
    changed = 0
    for target in frame_targets(shape):
        try:
            target.margin_left = left
            target.margin_top = top
            target.margin_right = right
            target.margin_bottom = bottom
        except Exception:
            log.debug("could not set insets on %r", shape, exc_info=True)
            continue
        changed += 1
    return changed


def set_alignment(shape: Any, alignment: str, *, only=None) -> int:
    """Set every paragraph's alignment. Returns paragraphs changed.

    `only` is a predicate over paragraphs, for the callers that must not touch
    all of them: the Arabic half of a bilingual shape, the header row of a
    table. Without it this writes the lot, which is what "make this shape's
    text flush right" means and what a designer gets from the ribbon.
    """
    wanted = _align_enum(alignment)
    if wanted is None:
        return 0

    changed = 0
    for paragraph in paragraphs(shape):
        if only is not None and not only(paragraph):
            continue
        if _safe(lambda p=paragraph: p.alignment) == wanted:
            continue
        try:
            paragraph.alignment = wanted
        except Exception:
            log.debug("could not align a paragraph on %r", shape, exc_info=True)
            continue
        changed += 1
    return changed


def set_reading_order(shape: Any, rtl: bool, *, only=None) -> int:
    """Mark paragraphs right-to-left, or left-to-right. Paragraphs changed."""
    wanted = "1" if rtl else "0"
    changed = 0
    for paragraph in paragraphs(shape):
        if only is not None and not only(paragraph):
            continue
        properties = _safe(lambda p=paragraph: p._p.get_or_add_pPr())
        if properties is None:
            continue
        was = properties.get("rtl")
        if was is not None and (was in ("1", "true")) == bool(rtl):
            continue
        properties.set("rtl", wanted)
        changed += 1
    return changed


def align_text(
    shape: Any,
    *,
    anchor: Optional[str] = None,
    centred: Optional[bool] = None,
    insets_in: Optional[Sequence[float]] = None,
    alignment: Optional[str] = None,
    rtl: Optional[bool] = None,
    only=None,
) -> ArrangeResult:
    """Set any combination of the four text-frame settings in one call.

    The entry point a caller wants when it is correcting "how the text sits in
    this box" rather than one attribute of it. Every argument left None is
    left alone, so a call that only anchors does only that.
    """
    result = ArrangeResult()
    if shape is None or not text_frames(shape):
        result.message = "this shape has no text frame to align"
        return result

    if anchor is not None and set_anchor(shape, anchor):
        result.record(f"anchored the text {_word(anchor, _ANCHOR_WORDS)}")
    if centred is not None and set_anchor_centred(shape, centred):
        result.record(
            "centred the text block across the box" if centred
            else "stopped centring the text block across the box"
        )
    if insets_in is not None and set_insets(shape, insets_in):
        result.record(f"set its text insets to {_inset_words(insets_in)}")
    if alignment is not None:
        count = set_alignment(shape, alignment, only=only)
        if count:
            result.record(
                f"set {count} paragraph(s) {_word(alignment, _ALIGN_WORDS)} aligned"
            )
    if rtl is not None:
        count = set_reading_order(shape, rtl, only=only)
        if count:
            way = "right-to-left" if rtl else "left-to-right"
            result.record(f"marked {count} paragraph(s) {way}")

    result.success = True
    return result


# --------------------------------------------------------------------------- #
# Words and enums
# --------------------------------------------------------------------------- #

def _word(value: Any, table: dict) -> Optional[str]:
    if value is None:
        return None
    return table.get(str(value).strip().upper())


def _anchor_word(value: Any) -> Optional[str]:
    """An MSO_ANCHOR, however it renders, as one of `ANCHORS`."""
    if value is None:
        return None
    name = str(value).split(" ")[0].upper()
    return _ANCHOR_WORDS.get(name)


def _align_word(value: Any) -> Optional[str]:
    if value is None:
        return None
    name = str(value).split(" ")[0].upper()
    return _ALIGN_WORDS.get(name)


def _align_enum(alignment: Any):
    from pptx.enum.text import PP_ALIGN  # noqa: PLC0415 - lazy, optional dep

    word = _word(alignment, _ALIGN_WORDS)
    return {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
        "distribute": PP_ALIGN.DISTRIBUTE,
    }.get(word)


def _body_pr(frame: Any):
    """`a:bodyPr`, made if the frame has none."""
    body = _safe(lambda: frame._txBody)
    if body is None:
        return None
    found = body.find(_qn("a:bodyPr"))
    if found is not None:
        return found
    return _safe(lambda: body.get_or_add_bodyPr())


def _qn(tag: str):
    from pptx.oxml.ns import qn  # noqa: PLC0415 - lazy, oxml internals

    return qn(tag)


def _inset_words(insets: Sequence[float]) -> str:
    """The four insets as a sentence, collapsed where they agree.

    The same spelling `rules.textframe._inset_words` writes, so a finding and
    the sentence reporting its correction describe the inset the same way.
    """
    left, top, right, bottom = (round(float(value), 2) for value in insets)
    if left == right and top == bottom:
        return f"{left:g}in at the sides and {top:g}in top and bottom"
    return f"{left:g}/{top:g}/{right:g}/{bottom:g}in"
