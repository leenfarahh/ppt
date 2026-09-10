"""Stage 1: read a .pptx into a DeckProfile.

This is the only module that knows about python-pptx. Everything downstream
works off the DeckProfile dataclasses, which keeps the rules testable without
fixture decks and keeps the AI payload independent of the parser.

python-pptx is imported lazily so that `formatting-tool --help` and the unit
tests still run in an environment where it is not installed.
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Optional
from xml.etree import ElementTree

from .. import svgicon
from ..models import (
    DeckProfile,
    Geometry,
    LayoutProfile,
    ParagraphProfile,
    RunProfile,
    ShapeProfile,
    SlideProfile,
    TextRole,
    placeholder_token,
)

_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_DRAWINGML_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

# Placeholder types that map onto a brand text role. python-pptx reports these
# as PP_PLACEHOLDER members, rendered as "SUBTITLE (4)"; the lookup is on the
# exact name, never a substring -- "TITLE" is a substring of "SUBTITLE".
_PLACEHOLDER_ROLES: dict[str, TextRole] = {
    "TITLE": TextRole.TITLE,
    "CENTER_TITLE": TextRole.TITLE,
    "VERTICAL_TITLE": TextRole.TITLE,
    "SUBTITLE": TextRole.SUBTITLE,
    "BODY": TextRole.BODY,
    "VERTICAL_BODY": TextRole.BODY,
    "OBJECT": TextRole.BODY,
    "FOOTER": TextRole.FOOTER,
    "SLIDE_NUMBER": TextRole.FOOTER,
    "DATE": TextRole.FOOTER,
}


class DeckReadError(RuntimeError):
    """Raised when a deck cannot be opened or parsed."""


def read_deck(path: str | Path) -> DeckProfile:
    """Open a .pptx and return its full profile."""
    path = Path(path)
    if not path.exists():
        raise DeckReadError(f"deck not found: {path}")

    try:
        from pptx import Presentation  # noqa: PLC0415 - lazy heavy dependency

        from .pptx_speedup import apply as _speed_up  # noqa: PLC0415

        # Once python-pptx is loaded, and not before: the library is imported
        # lazily so that `--help` does not pay for it, and this would undo that.
        _speed_up()
    except ImportError as exc:  # pragma: no cover
        raise DeckReadError(
            "python-pptx is required to read decks (pip install python-pptx)"
        ) from exc

    try:
        prs = Presentation(str(path))
    except Exception as exc:  # python-pptx raises a range of types here
        raise DeckReadError(f"could not open {path.name}: {exc}") from exc

    theme_fonts, theme_colors = _read_theme(prs)

    deck = DeckProfile(
        path=str(path),
        width_in=_inches(prs.slide_width),
        height_in=_inches(prs.slide_height),
        layouts=_read_layouts(prs),
        theme_fonts=theme_fonts,
        theme_colors=theme_colors,
    )

    for index, slide in enumerate(prs.slides, start=1):
        deck.slides.append(_read_slide(slide, index))
    return deck


# --------------------------------------------------------------------------- #
# Slides and shapes
# --------------------------------------------------------------------------- #

def _read_slide(slide: Any, number: int) -> SlideProfile:
    profile = SlideProfile(
        number=number,
        slide_id=_safe(lambda: slide.slide_id),
        layout_name=_safe(lambda: slide.slide_layout.name),
        hidden=_safe(lambda: slide.element.get("show")) == "0",
    )
    if _safe(lambda: slide.has_notes_slide):
        profile.notes = _safe(lambda: slide.notes_slide.notes_text_frame.text) or ""
    for shape in slide.shapes:
        profile.shapes.append(_read_shape(shape))
    return profile


class _Frame:
    """The transform from a shape's own coordinates into slide coordinates.

    A shape inside a group does not store where it sits on the slide. It stores
    where it sits in the group's *child* space, an arbitrary coordinate system
    the group declares with chOff and chExt, which the group then maps onto the
    slide through its own off and ext. Read `shape.left` on a grouped shape and
    you get a number in that private system: on one real deck a badge whose
    slide position is 8.06in reports 7.63in, and its neighbour in the same group
    reports a width of 1.09in while occupying 0.64in.

    Left uncomposed, every geometric rule that descends into a group is
    comparing inches from two different coordinate systems, which is how a
    misplaced element inside a group stays invisible to a checker that is
    otherwise measuring correctly.

    The identity frame is the slide itself, so a top-level shape is unaffected.
    """

    __slots__ = ("dx", "dy", "sx", "sy")

    def __init__(self, dx: float = 0.0, dy: float = 0.0, sx: float = 1.0, sy: float = 1.0):
        self.dx, self.dy, self.sx, self.sy = dx, dy, sx, sy

    def apply(self, left: float, top: float, width: float, height: float):
        return (
            self.dx + left * self.sx,
            self.dy + top * self.sy,
            width * self.sx,
            height * self.sy,
        )

    def descend(self, group: Any) -> "_Frame":
        """The frame for the children of `group`, expressed in slide space."""
        xfrm = _group_xfrm(group)
        if xfrm is None:
            return self
        off, ext = xfrm.find(_DRAWINGML_NS + "off"), xfrm.find(_DRAWINGML_NS + "ext")
        child_off = xfrm.find(_DRAWINGML_NS + "chOff")
        child_ext = xfrm.find(_DRAWINGML_NS + "chExt")
        if off is None or ext is None or child_off is None or child_ext is None:
            return self
        try:
            cx, cy = int(child_ext.get("cx")), int(child_ext.get("cy"))
            # A zero child extent is a degenerate group. Scaling by it would
            # divide by zero; leaving the frame alone keeps the child's own
            # numbers, which is no worse than before.
            sx = self.sx * (int(ext.get("cx")) / cx) if cx else self.sx
            sy = self.sy * (int(ext.get("cy")) / cy) if cy else self.sy
            gx = self.dx + int(off.get("x")) * self.sx
            gy = self.dy + int(off.get("y")) * self.sy
            return _Frame(
                gx - int(child_off.get("x")) * sx,
                gy - int(child_off.get("y")) * sy,
                sx,
                sy,
            )
        except (TypeError, ValueError):
            return self


def _group_xfrm(group: Any) -> Any:
    """The group's *own* transform, off its grpSpPr.

    Searched one level down rather than with `.//xfrm`, which is a document
    order descendant search and would happily return the first child shape's
    transform for any group whose grpSpPr carries none.
    """
    element = _safe(lambda: group._element)
    if element is None:
        return None
    for child in element:
        if str(child.tag).endswith("}grpSpPr") or str(child.tag) == "grpSpPr":
            return child.find(_DRAWINGML_NS + "xfrm")
    return None


def _read_shape(shape: Any, frame: Optional[_Frame] = None) -> ShapeProfile:
    frame = frame or _Frame()
    shape_type = str(_safe(lambda: shape.shape_type) or "UNKNOWN")
    placeholder_type = None
    placeholder_idx = None
    if _safe(lambda: shape.is_placeholder):
        placeholder_type = str(_safe(lambda: shape.placeholder_format.type) or "")
        placeholder_idx = _safe(lambda: shape.placeholder_format.idx)

    fill_hex, fill_theme = _fill_hex(shape)
    line_hex, line_theme = _line_hex(shape)
    graphic_colors = svgicon.colors_of(shape)

    profile = ShapeProfile(
        shape_id=_safe(lambda: shape.shape_id) or -1,
        name=_safe(lambda: shape.name) or "",
        shape_type=shape_type,
        geometry=_geometry(shape, frame),
        placeholder_type=placeholder_type,
        placeholder_idx=placeholder_idx if placeholder_idx is None else int(placeholder_idx),
        role=_role_for(placeholder_type, _safe(lambda: shape.name) or ""),
        fill_hex=fill_hex,
        fill_theme=fill_theme,
        line_hex=line_hex,
        line_theme=line_theme,
        graphic_colors=graphic_colors,
        is_picture="PICTURE" in shape_type,
        is_group="GROUP" in shape_type,
        alt_text=_alt_text(shape),
    )

    if profile.is_picture:
        profile.image_sha1 = _image_sha1(shape)

    if _safe(lambda: shape.has_text_frame):
        frame = shape.text_frame
        profile.text = _safe(lambda: frame.text) or ""
        profile.word_wrap = _safe(lambda: frame.word_wrap)
        profile.autofit = _stringify(_safe(lambda: frame.auto_size))
        profile.inset_top_in = _inches(_safe(lambda: frame.margin_top))
        profile.inset_bottom_in = _inches(_safe(lambda: frame.margin_bottom))
        for paragraph in frame.paragraphs:
            profile.paragraphs.append(_read_paragraph(paragraph))

    if profile.is_group:
        inner = frame.descend(shape)
        for child in _safe(lambda: list(shape.shapes)) or []:
            profile.children.append(_read_shape(child, inner))

    return profile


def _alt_text(shape: Any) -> str:
    """cNvPr/@descr, "" when the shape carries none.

    Read off the element rather than through python-pptx, which exposes this
    only on pictures (`shape._element` is the same object either way, and every
    shape kind carries a cNvPr). Searched from the shape root because the tag
    sits under nvSpPr, nvPicPr, nvGrpSpPr or nvGraphicFramePr depending on what
    the shape is, and the first match in document order is the shape's own.
    """
    element = _safe(lambda: shape._element)
    if element is None:
        return ""
    cNvPr = element.find(".//" + _P_NS + "cNvPr")
    if cNvPr is None:
        return ""
    return cNvPr.get("descr") or ""


def _geometry(shape: Any, frame: _Frame) -> Geometry:
    """The shape's box in slide inches, whatever frame it was authored in."""
    left, top, width, height = frame.apply(
        float(_safe(lambda: shape.left) or 0),
        float(_safe(lambda: shape.top) or 0),
        float(_safe(lambda: shape.width) or 0),
        float(_safe(lambda: shape.height) or 0),
    )
    return Geometry(
        left_in=_inches(left),
        top_in=_inches(top),
        width_in=_inches(width),
        height_in=_inches(height),
        rotation=float(_safe(lambda: shape.rotation) or 0.0),
    )


def _read_paragraph(paragraph: Any) -> ParagraphProfile:
    return ParagraphProfile(
        text=_safe(lambda: paragraph.text) or "",
        level=int(_safe(lambda: paragraph.level) or 0),
        alignment=_stringify(_safe(lambda: paragraph.alignment)),
        rtl=_paragraph_rtl(paragraph),
        space_before_pt=_points(_safe(lambda: paragraph.space_before)),
        space_after_pt=_points(_safe(lambda: paragraph.space_after)),
        line_spacing=_line_spacing(_safe(lambda: paragraph.line_spacing)),
        runs=[_read_run(run) for run in _safe(lambda: paragraph.runs) or []],
    )


def _paragraph_rtl(paragraph: Any) -> Optional[bool]:
    """`a:pPr/@rtl`, or None when the paragraph does not say.

    python-pptx has no accessor for it, so it is read off the element. None is
    a real answer and not a failure: a paragraph that says nothing inherits,
    and an Arabic paragraph that inherits left-to-right is the defect
    `typography.rtl_not_set` reports.
    """
    try:
        properties = paragraph._p.find(f"{_A_NS}pPr")
        if properties is None:
            return None
        value = properties.get("rtl")
        if value is None:
            return None
        return value in ("1", "true")
    except Exception:
        return None


def _read_run(run: Any) -> RunProfile:
    font = _safe(lambda: run.font)
    color_hex, is_theme, theme_slot = _font_color(font)
    return RunProfile(
        text=_safe(lambda: run.text) or "",
        font_name=_safe(lambda: font.name),
        size_pt=_points(_safe(lambda: font.size)),
        bold=_safe(lambda: font.bold),
        italic=_safe(lambda: font.italic),
        underline=_safe(lambda: font.underline),
        color_hex=color_hex,
        color_is_theme=is_theme,
        color_theme=theme_slot,
        language=_stringify(_safe(lambda: font.language_id)),
    )


def _role_for(placeholder_type: Optional[str], shape_name: str) -> TextRole:
    """Best-effort role assignment.

    A placeholder tells us its role outright. A loose text box does not, so we
    fall back to the shape name, which designers usually leave meaningful
    ("Title 1", "Subtitle 2") -- and which is itself a signal worth reporting
    when it disagrees with the geometry.
    """
    token = placeholder_token(placeholder_type)
    if token and token in _PLACEHOLDER_ROLES:
        return _PLACEHOLDER_ROLES[token]

    lowered = shape_name.lower()
    if "subtitle" in lowered:
        return TextRole.SUBTITLE
    if "title" in lowered:
        return TextRole.TITLE
    if "caption" in lowered:
        return TextRole.CAPTION
    return TextRole.UNKNOWN


# --------------------------------------------------------------------------- #
# Layouts
# --------------------------------------------------------------------------- #

def _read_layouts(prs: Any) -> list[LayoutProfile]:
    """Read every layout in the file, shapes and all.

    Names alone are not enough for either job that needs layouts: checking a
    layout carries its header and footer furniture means looking at its
    shapes, and choosing which layout a messy slide belongs on means comparing
    placeholder structure.
    """
    layouts: list[LayoutProfile] = []
    for index, layout in enumerate(_layouts(prs)):
        profile = LayoutProfile(
            name=_safe(lambda: layout.name) or f"layout {index + 1}",
            index=index,
        )
        for shape in _safe(lambda: list(layout.shapes)) or []:
            profile.shapes.append(_read_shape(shape))
        layouts.append(profile)
    return layouts


# --------------------------------------------------------------------------- #
# Colour and image helpers
# --------------------------------------------------------------------------- #

def _font_color(font: Any) -> tuple[Optional[str], bool, Optional[str]]:
    """Return (hex, resolves_through_theme, theme slot).

    A theme-bound colour has no literal RGB on the run, so the slot it binds
    to is the only thing there is to record. It is NOT "correct by
    construction", which is what this used to assume: it is correct only if
    the theme behind it is the master's. A deck built from another file
    resolves every one of these through its own theme and renders off-brand
    while each run looks blameless on its own.
    """
    color = _safe(lambda: font.color)
    if color is None:
        return None, False, None
    rgb = _safe(lambda: color.rgb)
    if rgb is not None:
        return str(rgb).upper(), False, None
    theme_color = _safe(lambda: color.theme_color)
    if theme_color is not None:
        return None, True, _theme_slot(theme_color)
    return None, False, None


# python-pptx names the slots after the OOXML enum; the theme part names them
# as they appear in `a:clrScheme`. Same twelve slots, two spellings, and the
# comparison has to happen in one of them.
_SLOTS = {
    "DARK_1": "dk1", "LIGHT_1": "lt1", "DARK_2": "dk2", "LIGHT_2": "lt2",
    "TEXT_1": "dk1", "BACKGROUND_1": "lt1",
    "TEXT_2": "dk2", "BACKGROUND_2": "lt2",
    "ACCENT_1": "accent1", "ACCENT_2": "accent2", "ACCENT_3": "accent3",
    "ACCENT_4": "accent4", "ACCENT_5": "accent5", "ACCENT_6": "accent6",
    "HYPERLINK": "hlink", "FOLLOWED_HYPERLINK": "folHlink",
}


def _theme_slot(theme_color: Any) -> Optional[str]:
    """A python-pptx theme colour -> the name the theme part uses for it."""
    name = _stringify(theme_color)
    if not name:
        return None
    # "ACCENT_1 (5)" -> "ACCENT_1"
    head = str(name).split(" ")[0].upper()
    return _SLOTS.get(head)


def _fill_hex(shape: Any) -> tuple[Optional[str], Optional[str]]:
    """A shape's fill as (hex, theme slot). Both None when it has no fill."""
    fill = _safe(lambda: shape.fill)
    if fill is None:
        return None, None
    if str(_safe(lambda: fill.type) or "") in ("None", "BACKGROUND (5)"):
        return None, None
    fore = _safe(lambda: fill.fore_color)
    if fore is None:
        return None, None
    rgb = _safe(lambda: fore.rgb)
    if rgb is not None:
        return str(rgb).upper(), None
    return None, _theme_slot(_safe(lambda: fore.theme_color))


def _line_hex(shape: Any) -> tuple[Optional[str], Optional[str]]:
    color = _safe(lambda: shape.line.color)
    if color is None:
        return None, None
    rgb = _safe(lambda: color.rgb)
    if rgb is not None:
        return str(rgb).upper(), None
    return None, _theme_slot(_safe(lambda: color.theme_color))


def _image_sha1(shape: Any) -> Optional[str]:
    """Identify a picture by content so the same logo matches across decks."""
    sha1 = _safe(lambda: shape.image.sha1)
    if sha1:
        return str(sha1)
    blob = _safe(lambda: shape.image.blob)
    return hashlib.sha1(blob).hexdigest() if blob else None


# --------------------------------------------------------------------------- #
# Theme
# --------------------------------------------------------------------------- #

def _read_theme(prs: Any) -> tuple[dict[str, str], dict[str, str]]:
    """Pull the major/minor typefaces and colour scheme out of theme1.xml.

    python-pptx exposes no theme API, so this reads the related theme part
    directly. It degrades to empty dicts rather than failing the run: a missing
    theme costs precision in the colour and font rules, not correctness.
    """
    xml = _theme_xml(prs)
    if xml is None:
        return {}, {}

    try:
        root = ElementTree.fromstring(xml)
    except ElementTree.ParseError:
        return {}, {}

    fonts: dict[str, str] = {}
    for key, tag in (("major", "majorFont"), ("minor", "minorFont")):
        node = root.find(f".//{_DRAWINGML_NS}{tag}/{_DRAWINGML_NS}latin")
        if node is not None and node.get("typeface"):
            fonts[key] = node.get("typeface", "")

    colors: dict[str, str] = {}
    scheme = root.find(f".//{_DRAWINGML_NS}clrScheme")
    if scheme is not None:
        for child in scheme:
            name = child.tag.replace(_DRAWINGML_NS, "")
            srgb = child.find(f"{_DRAWINGML_NS}srgbClr")
            if srgb is not None and srgb.get("val"):
                colors[name] = srgb.get("val", "").upper()
                continue
            sys_clr = child.find(f"{_DRAWINGML_NS}sysClr")
            if sys_clr is not None and sys_clr.get("lastClr"):
                colors[name] = sys_clr.get("lastClr", "").upper()
    return fonts, colors


def _theme_xml(prs: Any) -> Optional[bytes]:
    for master in _safe(lambda: list(prs.slide_masters)) or []:
        rels = _safe(lambda: master.part.rels)
        if not rels:
            continue
        for rel in _safe(lambda: list(rels.values())) or []:
            if "theme" in str(_safe(lambda: rel.reltype) or "").lower():
                blob = _safe(lambda: rel.target_part.blob)
                if blob:
                    return blob
    return None


def _layouts(prs: Any) -> list[Any]:
    layouts: list[Any] = []
    for master in _safe(lambda: list(prs.slide_masters)) or []:
        layouts.extend(_safe(lambda: list(master.slide_layouts)) or [])
    return layouts


# --------------------------------------------------------------------------- #
# Unit helpers
# --------------------------------------------------------------------------- #

def _inches(length: Any) -> float:
    """python-pptx Length -> inches. Missing geometry reads as 0.0."""
    if length is None:
        return 0.0
    inches = _safe(lambda: length.inches)
    if inches is not None:
        return round(float(inches), 4)
    return round(float(length) / 914400.0, 4)  # EMU per inch


def _points(length: Any) -> Optional[float]:
    if length is None:
        return None
    points = _safe(lambda: length.pt)
    if points is not None:
        return round(float(points), 2)
    if isinstance(length, (int, float)):
        return round(float(length) / 12700.0, 2)  # EMU per point
    return None


def _line_spacing(value: Any) -> Optional[float]:
    """Line spacing is either a multiple (float) or an absolute Length."""
    if value is None:
        return None
    if isinstance(value, float):
        return round(value, 3)
    return _points(value)


def _stringify(value: Any) -> Optional[str]:
    if value is None:
        return None
    return str(value)


def _safe(getter) -> Any:
    """Read an attribute off a python-pptx object, tolerating its exceptions.

    python-pptx raises rather than returning None in several places (a
    ColorFormat with no RGB, a shape with no fill, a picture whose image part
    is missing). A messy deck hits all of them, and none should abort a run.
    """
    try:
        return getter()
    except Exception:
        return None
