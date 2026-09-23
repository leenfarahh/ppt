"""Rebuild a messy deck on top of the approved master.

The master is opened as the base file, its own sample slides are dropped, and
every source slide is recreated on the layout it belongs to. Content moves
across; direct formatting does not.

That last point is the whole reason this exists. Re-pointing a slide at a
different layout in place changes nothing visible, because a layout supplies
only what a slide has not already set for itself, and a messy deck has set all
of it. Recreating the slide from the layout is what PowerPoint's Reset Slide
does, and it is the only way the master's geometry and type actually take
effect.

What crosses, and what does not:

    placeholder copy    moved into the matching placeholder on the new layout,
                        keeping paragraph level and bold/italic/underline
                        emphasis, dropping typeface, size and colour so that
                        the layout supplies them
    loose shapes        transplanted as they are, position included; they were
                        never governed by a layout and are not now
    pictures            transplanted with the image itself, so crops and
                        effects survive. A picture PLACEHOLDER is frozen
                        first: the frame and the shape it is cropped to are
                        the old layout's, so they are written onto the shape
                        before it moves and it stops being a placeholder.
                        See `rebuild.pictures` -- without it a circular
                        portrait arrives square and at the origin
    charts              copied as one element, chart part, workbook and all,
                        at their own position -- never claimed, never cut
                        into pieces, never re-placed. See `rebuild.charts`
    SmartArt, media,    left behind and reported by name, because their content
    embedded objects    lives in parts this cannot rebuild, and a silently
                        broken diagram is worse than a missing one

Nothing is dropped quietly. Every slide, every shape and every omission ends up
in the RebuildResult.
"""

from __future__ import annotations

import copy
import hashlib
import logging
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any, Collection, Optional, Sequence

from ..classify import SlideKind
from ..extract import read_deck
from ..models import (
    DeckProfile,
    LayoutChoice,
    LayoutProfile,
    RuleTuning,
    SlideProfile,
)
from . import charts as chartparts
from . import master_apply
from . import rtl
from .matcher import FAMILIES, LATENT, LayoutMatch, choose_layout
from .pictures import freeze_slide, reads_as_a_note

log = logging.getLogger(__name__)

_R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


class RebuildError(RuntimeError):
    """Raised when the rebuild cannot start or cannot be written."""


@dataclass
class DroppedShape:
    slide: int
    name: str
    reason: str

    def __str__(self) -> str:
        return f"slide {self.slide}: {self.name!r} ({self.reason})"


@dataclass
class SlideRecord:
    """What happened to one slide, in enough detail to check the work."""

    number: int
    source_layout: Optional[str]
    target_layout: str
    basis: str
    confident: bool
    kind: SlideKind = SlideKind.UNKNOWN     # what the slide is for
    layout_kind: SlideKind = SlideKind.UNKNOWN
    filled: list[str] = field(default_factory=list)
    transplanted: list[str] = field(default_factory=list)
    unfilled: list[str] = field(default_factory=list)
    dropped: list[DroppedShape] = field(default_factory=list)
    # Left exactly as it arrived, because no layout in the master fits it and
    # the nearest one wrecked it. The master is NOT on this slide; a designer
    # has been asked to place it by hand, in a comment on the slide itself.
    # See `rebuild.quarantine`.
    quarantined: bool = False
    # What the restyle would have cost, in square inches of visible damage.
    damage_before: float = 0.0
    damage_after: float = 0.0


@dataclass
class RebuildResult:
    master: str
    deck: str
    output: Path
    slides: list[SlideRecord] = field(default_factory=list)
    sample_slides_removed: int = 0
    size_note: Optional[str] = None
    # Which route did the work: "powerpoint" when PowerPoint applied the
    # layouts itself, "xml" when this module rebuilt the file. Reported because
    # the two differ in a way a designer sees -- the XML route drops what it
    # cannot recreate faithfully and lists it, the PowerPoint route drops
    # nothing -- and a reader comparing two runs needs to know which they have.
    applied_by: str = "xml"
    # PowerPoint route only: slides left on another design, and how many
    # masters the output carries as a result.
    stragglers: list[int] = field(default_factory=list)
    masters: int = 1
    # Shapes turned round so a right-to-left deck reads that way. Zero on an
    # English deck, where turning it round would be the defect.
    mirrored: int = 0
    # Why no slide was handed back, when the route could not do it. Set on the
    # XML route only, and only where there was something it would have
    # considered: carrying a slide over verbatim means not restyling it, and
    # that route builds the output from the master rather than editing the
    # deck, so there is no untouched slide for it to keep. Said out loud
    # rather than left as a silent difference between two hosts.
    quarantine_note: Optional[str] = None

    @property
    def unmatched(self) -> list[SlideRecord]:
        """Slides placed on a layout that was only the least-bad option."""
        return [record for record in self.slides if not record.confident]

    @property
    def quarantined(self) -> list[SlideRecord]:
        """Slides handed back untouched for a designer to place by hand.

        A subset of `unmatched`: the matcher's doubt is the gate and the
        measured damage is the verdict, so every one of these was doubted and
        only the wrecked ones are here.
        """
        return [record for record in self.slides if record.quarantined]

    @property
    def dropped(self) -> list[DroppedShape]:
        return [shape for record in self.slides for shape in record.dropped]

    @property
    def layouts_used(self) -> dict[str, int]:
        counts: dict[str, int] = {}
        for record in self.slides:
            counts[record.target_layout] = counts.get(record.target_layout, 0) + 1
        return counts


def rebuild(
    master: str | Path,
    deck: str | Path,
    out: str | Path,
    tuning: Optional[RuleTuning] = None,
    route: str = "auto",
    seen: Optional[Sequence[LayoutChoice]] = None,
    master_profile: Optional[DeckProfile] = None,
    deck_profile: Optional[DeckProfile] = None,
    mirror: Optional[bool] = None,
    roles: Optional[Any] = None,
) -> RebuildResult:
    """Rebuild `deck` onto `master` and write the result to `out`.

    `master_profile` and `deck_profile` let a caller that has already read
    those files hand them over instead of paying for the parse twice.

    `seen` is what the AI layer read off the rendered slides: which layout
    each one belongs on. It outranks the structural fit and not a layout-name
    match; see `matcher.choose_layout`.

    `roles` is what the same layer read off the same renders about the SHAPES:
    which box is the subtitle, which is a source line, which shapes are a
    chart, which are drawn furniture. It decides what may be claimed into a
    layout region and what has to survive the move untouched -- see
    `fill_runs`. Without it both routes fill by geometry alone, which is how a
    footnote ends up under the title and a chevron banner is deleted with the
    words it carried.

    `route` is "auto" to let PowerPoint do it where it can and fall back to
    rebuilding the file where it cannot, "powerpoint" to require PowerPoint, or
    "xml" to require the rebuild. Not a preference knob: the two routes produce
    genuinely different files -- see rebuild.master_apply -- so anything
    asserting about one has to say which.
    """
    from pptx import Presentation  # noqa: PLC0415 - lazy heavy dependency

    master, deck, out = Path(master), Path(deck), Path(out)
    tuning = tuning or RuleTuning()

    # Re-read only what the caller has not already read. Reading a 24MB deck
    # costs about eight seconds, and the pipeline has both files parsed by the
    # time it gets here; doing it again was a quarter of the local work.
    master_profile = master_profile or read_deck(master)
    deck_profile = deck_profile or read_deck(deck)
    if not master_profile.layouts:
        raise RebuildError(
            f"{master.name} defines no slide layouts, so there is nothing to "
            "rebuild onto"
        )

    plans = _plans(deck_profile, master_profile, tuning, seen)

    # Decided from the deck rather than asked for, unless the caller says
    # otherwise. A master is drawn for one reading direction and applying it
    # to a deck that reads the other way gets the type right and the layout
    # backwards, which is not a thing anybody wants by default.
    if mirror is None:
        mirror = deck_profile.rtl
    if mirror:
        log.info(
            "%s reads right to left; the rebuilt slides will be mirrored so "
            "the content does too", deck.name,
        )

    # PowerPoint first where it can run. Assigning CustomLayout runs its own
    # placeholder matching, which is what actually moves a slide's content into
    # the new layout's placeholders, and nothing is re-serialised by us, so the
    # class of corruption the XML route kept producing cannot arise.
    if route in ("auto", "powerpoint"):
        if master_apply.available():
            # The slides the matcher admits it is guessing about. Only these
            # are measured before and after the swap, and only one the restyle
            # actually wrecked is handed back -- doubt is the gate, damage is
            # the verdict. See `rebuild.quarantine`.
            doubtful = {n for n, match in plans.items() if not match.confident}
            applied = master_apply.apply_master(
                deck, master, out, {n: m.name for n, m in plans.items()},
                mirror=mirror, doubtful=doubtful, roles=roles,
            )
            if applied.fatal is None:
                result = _result_from_powerpoint(master, deck, out, plans, applied)
                # After PowerPoint, not instead of it: its own matching has
                # already moved every placeholder it could, and this fills what
                # is left from the slide's loose runs. Never fatal -- a failure
                # here leaves the deck exactly as PowerPoint wrote it, which is
                # what this route produced before.
                #
                # NOT on a quarantined slide. That slide was deliberately left
                # as the designer sent it, and filling its runs would claim
                # its boxes and delete them -- undoing the one thing the
                # quarantine promised and doing it to the slide least able to
                # take it.
                spared = {record.number for record in result.quarantined}
                try:
                    _record_filled(
                        result, *fill_runs(out, skip=spared, roles=roles)
                    )
                except Exception:
                    log.warning(
                        "could not fill the layouts' repeated regions from the "
                        "deck's own runs; the slides keep their loose boxes",
                        exc_info=True,
                    )
                _reset_at_the_end(out, mirror, skip=spared)
                return result
            if route == "powerpoint":
                raise RebuildError(applied.fatal)
            log.warning(
                "PowerPoint could not apply the master (%s); rebuilding the "
                "file instead, which drops what it cannot recreate",
                applied.fatal,
            )
        elif route == "powerpoint":
            raise RebuildError(
                "this host cannot drive PowerPoint, so the master cannot be "
                "applied that way; use route='xml' or 'auto'"
            )

    base = Presentation(str(master))
    source = Presentation(str(deck))

    result = RebuildResult(master=master.name, deck=deck.name, output=out)
    result.sample_slides_removed = _drop_slides(base)
    result.size_note = _size_note(base, source)

    canvas_in = (deck_profile.width_in, deck_profile.height_in)
    for profile, src_slide in zip(deck_profile.slides, source.slides):
        result.slides.append(
            _rebuild_slide(
                profile, src_slide, base, plans[profile.number],
                here=roles.for_slide(profile.number) if roles else None,
                canvas_in=canvas_in,
            )
        )

    if mirror:
        result.mirrored = rtl.mirror_presentation(base)

    # The XML route rebuilds the output from the master, so there is no
    # original slide sitting in it to keep -- the quarantine is a property of
    # editing the deck in place, which only the PowerPoint route does. The
    # slides it would have considered are still named, so a designer reading
    # this report is not told less than one reading the other route's.
    if result.unmatched:
        result.quarantine_note = (
            f"{len(result.unmatched)} slide(s) were placed on a layout that "
            f"was only the least-bad option, and this route cannot hand one "
            f"back: it builds the output from the master rather than editing "
            f"the deck, so there is no untouched slide to keep. Re-run on a "
            f"Windows host with desktop PowerPoint to have those measured and "
            f"left as they arrived. Slides: "
            f"{', '.join(str(r.number) for r in result.unmatched)}."
        )
        log.warning("%s", result.quarantine_note)

    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        base.save(str(out))
    except Exception as exc:
        raise RebuildError(f"could not write {out}: {exc}") from exc

    _reset_at_the_end(out, mirror)

    log.info(
        "rebuilt %d slide(s) from %s onto %s; %d shape(s) left behind",
        len(result.slides),
        deck.name,
        master.name,
        len(result.dropped),
    )
    return result


def _reset_at_the_end(
    out: Path, mirror: bool, skip: Collection[int] = (),
) -> None:
    """The last thing either route does: hand every placeholder back to its
    layout.

    NOT ON A MIRRORED DECK. Turning a right-to-left deck round works by writing
    an explicit position onto the placeholders that carry one -- see
    `rebuild.rtl` -- and this exists to take explicit positions off. Run after
    a mirror it would undo it, silently and completely. An RTL deck keeps the
    geometry the mirror gave it, which is the behaviour it had before this
    existed.

    Never fatal. A reset that fails leaves the deck as the route wrote it.
    """
    if mirror:
        log.info(
            "not resetting the placeholders: this deck was mirrored to read "
            "right to left, and the mirror is written as explicit positions"
        )
        return
    try:
        reset_layouts(out, skip=skip)
    except Exception:
        log.warning(
            "could not reset the slides onto their layouts' geometry; they "
            "keep the boxes they were rebuilt with", exc_info=True,
        )


def _plans(
    deck_profile: DeckProfile,
    master_profile: DeckProfile,
    tuning: RuleTuning,
    seen: Optional[Sequence[LayoutChoice]] = None,
) -> dict[int, LayoutMatch]:
    """The layout every slide should land on, decided once.

    Decided here rather than inside either route, so both apply the same picks
    and a report from one is comparable with a report from the other.
    """
    by_slide = {choice.slide: choice for choice in seen or ()}
    return {
        profile.number: choose_layout(
            profile, master_profile.layouts, tuning.layout_match_floor,
            deck=deck_profile, seen=by_slide.get(profile.number),
        )
        for profile in deck_profile.slides
    }


def _result_from_powerpoint(
    master: Path,
    deck: Path,
    out: Path,
    plans: dict[int, LayoutMatch],
    applied: master_apply.MasterApplyResult,
) -> RebuildResult:
    """The PowerPoint route's outcome, in the shape the report already reads.

    `filled` and `transplanted` stay empty on purpose. They are the XML route's
    account of which placeholder it copied what into, and PowerPoint does not
    report its matching decisions; claiming a shape landed somewhere specific
    would be inventing detail. What matters is stated: the layout applied, and
    whether it applied at all.
    """
    result = RebuildResult(
        master=master.name, deck=deck.name, output=out,
        applied_by="powerpoint",
        stragglers=applied.stragglers,
        masters=applied.masters,
        mirrored=applied.mirrored,
    )
    for outcome in applied.outcomes:
        match = plans.get(outcome.number)
        result.slides.append(
            SlideRecord(
                number=outcome.number,
                source_layout=None,
                target_layout=outcome.target_layout or "(none)",
                basis=(
                    outcome.detail if outcome.forced
                    else match.basis if match else "none"
                ),
                # A forced slide is on the master but not on the layout
                # anybody chose for it, so it is never reported as confident
                # however sure the matcher was of the pick that was missing.
                confident=bool(
                    match and match.confident and outcome.applied
                    and not outcome.forced
                ),
                kind=match.kind if match else SlideKind.UNKNOWN,
                layout_kind=match.layout_kind if match else SlideKind.UNKNOWN,
                quarantined=outcome.quarantined,
                damage_before=outcome.damage_before,
                damage_after=outcome.damage_after,
            )
        )
    log.info(
        "PowerPoint applied %s to %d of %d slide(s) of %s; %d left on another "
        "design, %d master(s) in the output",
        master.name, applied.applied, len(applied.outcomes), deck.name,
        len(applied.stragglers), applied.masters,
    )
    for failure in applied.failed:
        log.warning("slide %d not restyled: %s", failure.number, failure.detail)
    for forced in applied.forced:
        log.warning(
            "slide %d placed on %r: %s",
            forced.number, forced.target_layout, forced.detail,
        )
    if applied.quarantined:
        log.warning(
            "%d slide(s) were left exactly as they arrived because no layout "
            "in %s fits them and the nearest one wrecked them: %s. Each "
            "carries a PowerPoint comment asking a designer to place it.",
            len(applied.quarantined), master.name,
            ", ".join(str(o.number) for o in applied.quarantined),
        )
    return result


# --------------------------------------------------------------------------- #
# One slide
# --------------------------------------------------------------------------- #

def _rebuild_slide(
    profile: SlideProfile,
    src_slide: Any,
    base: Any,
    match: Optional[LayoutMatch],
    here: Optional[Any] = None,
    canvas_in: tuple[float, float] = (0.0, 0.0),
) -> SlideRecord:
    # The match arrives decided. Both routes read the same plan, so a report
    # from one is comparable with a report from the other, and re-deciding
    # here would quietly ignore what the AI layer read off the render.
    new_slide = base.slides.add_slide(_layout_object(base, match))

    record = SlideRecord(
        number=profile.number,
        source_layout=profile.layout_name,
        target_layout=match.name,
        basis=match.basis,
        confident=match.confident,
        kind=match.kind,
        layout_kind=match.layout_kind,
    )

    # Before anything claims a placeholder: a picture placeholder's frame and
    # its cropped-to shape are the OLD layout's, and both are about to stop
    # existing. See `pictures.freeze_slide`.
    frozen = freeze_slide(src_slide)
    if frozen:
        log.debug(
            "slide %d: kept the original frame of %s",
            profile.number, ", ".join(frozen),
        )

    pool = _placeholder_pool(new_slide)
    canvas = _canvas_of(base)

    # The placeholder pass first, and on its own, because it is the stronger
    # claim: a shape that was already a placeholder is saying which region it
    # belongs in, and the runs below are an inference from how a slide is
    # arranged. Deciding it all up front also means the transplant loop can be
    # told what has already found a home, rather than working it out twice.
    # READ ONCE AND KEEP THE OBJECTS. python-pptx builds a fresh proxy on
    # every `shapes` access, so the same shape is a different Python object
    # each time it is walked and `id()` is only stable within one list. Keyed
    # across two walks the claims match nothing, silently: the file is written
    # unchanged and the log still says what was decided.
    shapes = list(src_slide.shapes)
    charts_here = chartparts.chart_boxes(shapes)

    claims: dict[int, Any] = {}
    for shape in shapes:
        if not _is_placeholder(shape) or chartparts.is_chart(shape):
            # A chart in a content placeholder is still a chart. Claimed, it
            # would take the region and then fail to copy into it.
            continue
        target = _claim(pool, shape, canvas)
        if target is not None:
            claims[id(shape)] = target

    # A source line, a chart's parts and drawn furniture are never claimed:
    # see `fill_runs`, which makes the same exclusion on the other route.
    loose = [
        shape for shape in shapes
        if not _is_placeholder(shape) and _has_copy(shape)
        and not _leave_alone(shape, here, canvas_in, charts_here)
    ]
    run_claims, furniture, moves = claim_runs(loose, pool, shapes)
    claims.update(run_claims)
    subtitle = _subtitle_text(here)
    claims.update(claim_drawn_over(
        [s for s in loose if id(s) not in claims], pool, subtitle=subtitle,
    ))
    claims.update(claim_subtitle(
        [s for s in loose if id(s) not in claims], pool, subtitle,
    ))
    spare = {id(shape) for shape in furniture}
    written = layout_writes(new_slide.slide_layout)

    for shape in shapes:
        target = claims.get(id(shape))
        if target is not None:
            # Only counted as filled if the copy actually went. A region that
            # cannot hold text -- a table, chart or media slot, all of them
            # "content" in `FAMILIES` -- would otherwise be reported as filled
            # while the words were neither written into it nor carried over,
            # because this branch `continue`s past the transplant below.
            if _copy_text(shape, target):
                record.filled.append(_name_of(target))
                continue
            log.warning(
                "slide %d: %r was claimed by %r, which cannot hold text; it "
                "is carried over at its own position instead",
                profile.number, _name_of(shape), _name_of(target),
            )
        # NOTHING PROTECTED IS DROPPED BY ANY ROUTE. Keeping it out of `loose`
        # stops it being claimed into a region; it does not stop
        # `furniture_of` sweeping it up as the run's own drawing, or
        # `echoes_layout` reading it as a title the layout already writes.
        # Both drop the shape, and a chevron banner dropped as a run's drawing
        # is exactly as gone as one dropped with its copy.
        if _leave_alone(shape, here, canvas_in, charts_here):
            _transplant(shape, src_slide, new_slide, record)
            continue
        if id(shape) in spare:
            record.dropped.append(
                DroppedShape(
                    slide=profile.number,
                    name=_name_of(shape),
                    reason=RUN_DRAWING,
                )
            )
            continue
        if echoes_layout(shape, written):
            record.dropped.append(
                DroppedShape(
                    slide=profile.number,
                    name=_name_of(shape),
                    reason=LAYOUT_ECHO,
                )
            )
            continue
        _transplant(shape, src_slide, new_slide, record)
        _reposition(new_slide, moves.get(id(shape)))

    # Placeholders nothing claimed are left in place rather than deleted. They
    # do not render in a show or in print, and they are the visible answer to
    # "where should this loose text box actually go", which is the one
    # judgement this cannot make for a designer.
    record.unfilled = [_name_of(placeholder) for placeholder in pool]

    lift_imagery(new_slide, _canvas_in(base))

    _copy_notes(src_slide, new_slide)
    if profile.hidden:
        new_slide._element.set("show", "0")
    return record


def _layout_object(base: Any, match: LayoutMatch) -> Any:
    """The live layout object behind the matched LayoutProfile.

    Found by position rather than by name: a file can carry two slide masters
    that use the same layout name, and the profile's index is what actually
    identifies which one was chosen.
    """
    layouts = [
        layout for master in base.slide_masters for layout in master.slide_layouts
    ]
    return layouts[match.layout.index]


# --------------------------------------------------------------------------- #
# Placeholders
# --------------------------------------------------------------------------- #

def _is_placeholder(shape: Any) -> bool:
    try:
        return bool(shape.is_placeholder)
    except Exception:
        return False


def _placeholder_pool(slide: Any) -> list[Any]:
    """Placeholders on the new slide, in document order, none claimed yet."""
    try:
        return list(slide.placeholders)
    except Exception:
        return []


def _claim(
    pool: list[Any], shape: Any, canvas: tuple[int, int] = (0, 0)
) -> Optional[Any]:
    """Take the unclaimed placeholder of the same family nearest this shape.

    NEAREST, not first, and the difference is a slide arriving scrambled. The
    pool is in the layout's document order, which is the order the designer's
    XML happens to carry and has nothing to do with the order a reader sees.
    On an agenda of eleven numbered items -- twenty-two placeholders, a number
    and a label each -- taking the first free one of the family put item 01 in
    the last slot on the slide, shifted every other item up by one, and dropped
    a two-digit number into a slot sized for something else, where it wrapped
    to "0 / 3".

    Matched on the box instead: the item that sits top-left on the old slide
    goes in the placeholder that sits top-left on the new one. Overlap decides
    it, with the distance between centres breaking ties, so a slide whose
    layout has moved a little still lands in reading order and a number slot
    takes a number.

    Greedy, in the order the source slide lists its shapes. A shape can take a
    placeholder a later one wanted, which a global assignment would avoid; on
    every deck measured this has not come up, because a source item overlaps
    its own slot far more than it overlaps its neighbour's, and the simpler
    thing that can be read in one sitting is worth more here than the last
    fraction of a percent.

    A source placeholder with no counterpart on the new layout returns None
    and is transplanted instead, so extra content is never silently lost.
    """
    family = _family_of(shape)
    if family is None:
        return None
    candidates = [
        index for index, candidate in enumerate(pool)
        if _family_of(candidate) == family
    ]
    if not candidates:
        return None
    best = max(candidates, key=lambda index: _affinity(shape, pool[index], canvas))
    return pool.pop(best)


def _reads_rtl(presentation: Any) -> bool:
    """Whether this written deck reads right to left.

    The same rule as `DeckProfile.rtl` -- a majority of the text-bearing
    shapes -- counted here off the open python-pptx tree, because `fill_runs`
    works on the file PowerPoint has just written and has no DeckProfile for
    it. Reading the deck again to get one costs seconds on a 30MB file for a
    single boolean.

    Never raises: a deck that will not answer reads left to right, which is
    the behaviour this had before the direction was consulted at all.
    """
    from ..script import is_rtl        # noqa: PLC0415 - avoids a cycle

    rtl = latin = 0
    try:
        for slide in presentation.slides:
            for shape in slide.shapes:
                if not _has_copy(shape):
                    continue
                if is_rtl(shape.text_frame.text):
                    rtl += 1
                else:
                    latin += 1
    except Exception:
        log.debug("could not tell which way the deck reads", exc_info=True)
        return False
    return rtl > latin


def claim_runs(
    loose: Sequence[Any], pool: list[Any], everything: Sequence[Any] = (),
    rtl: bool = False,
) -> tuple[dict[int, Any], list[Any], dict[int, tuple]]:
    """Fill the layout's repeated regions from the slide's repeated content.

    THE CASE THIS IS FOR. A messy deck keeps its content in loose boxes, and
    `_claim` above is only ever offered shapes that are already placeholders,
    so on a real deck the master's `Agenda` layout arrived with every one of
    its twenty-four slots empty and the slide's five agenda items transplanted
    at their old inches -- across the bottom of a layout that puts a photograph
    there. The right layout, chosen correctly, and it may as well not have
    been.

    ORDER DECIDES IT, NOT POSITION, and `rebuild.series` explains at length why
    position cannot: the slide's agenda is five cards across the bottom and the
    layout's is a twelve-item list down the right half, so of the five items
    one overlaps no slot at all and the other four collide in pairs. What maps
    them is that both sides are runs, and a designer fills run to run in the
    order they are read.

    Placeholders it claims come out of `pool`, so they are not reported as
    unfilled and cannot be claimed twice. Nothing is filled by halves: a run
    with no layout run long enough to take all of it is left exactly as it was.
    """
    if not loose or not pool:
        return {}, [], {}
    # Empty ones only. In the rebuild the whole pool is empty by construction,
    # but the same work runs over a file PowerPoint has already written, where
    # a region may have been filled by its own placeholder matching -- and
    # writing a run over that would lose the copy that was already home.
    free = [
        p for p in pool
        if _family_of(p) == "content" and not _has_copy(p)
    ]
    if not free:
        return {}, [], {}

    from .series import (  # noqa: PLC0415 - cycle
        belongs_to,
        companion_of,
        find_series,
        fit_into,
        furniture_of,
        pair_up,
    )

    # Loose shapes only, and that restriction is load-bearing. Offered every
    # shape on the slide, `furniture_of` swept up the layout's own empty
    # placeholders -- small, wordless, and sitting squarely in the run's block
    # -- and deleted the regions the run had just been filled into.
    candidates = [shape for shape in everything if not _is_placeholder(shape)]

    # The SAME direction for both sides. Ordering the slide's run one way and
    # the layout's the other is the bug this parameter exists for, not a fix
    # for it. See `series._ordered`.
    regions = find_series(free, shortest=2, rtl=rtl)
    pairs = pair_up(find_series(loose, rtl=rtl), regions)
    claims: dict[int, Any] = {}
    spare: list[Any] = []
    moves: dict[int, tuple[int, int, int, int]] = {}
    for run, slots in pairs:
        for shape, target in zip(run.shapes, slots.shapes):
            claims[id(shape)] = target
            pool.remove(target)

        furniture = furniture_of(run, candidates)
        moved = _move_imagery(
            run, slots, [s for s in furniture if _is_imagery(s)],
            regions, pool, companion_of, belongs_to, fit_into,
        )
        moves.update(moved)
        spare.extend(s for s in furniture if id(s) not in moved)
        log.info(
            "filled %d of %d repeated region(s) from a run of %d loose shape(s)",
            len(run), len(slots), len(run),
        )
    return claims, spare, moves


def _text_of(shape: Any) -> str:
    """A shape's copy, or "" for anything that will not say.

    Here as well as in `pictures` because the two modules read it for
    different questions and neither should import the other for one line.
    """
    try:
        return str(shape.text_frame.text) if shape.has_text_frame else ""
    except Exception:
        return ""


def _leave_alone(
    shape: Any, here: Optional[Any], canvas: tuple[float, float],
    charts: Sequence[Any] = (),
) -> bool:
    """Whether this box must stay exactly where the designer drew it.

    A CHART FIRST, AND ALL OF IT. A native chart, anything drawn inside one's
    frame (`charts`, from `rebuild.charts.chart_boxes`), and anything inside a
    region the model called a chart -- asked by POSITION, because the bars, the
    gridlines and the axis of a chart drawn as shapes carry no copy, and the
    copy test below was the only protection they had. It gave them none, and
    they were deleted as the drawing of the run their data labels made.

    Then the model's word, because it read the picture; the file's own
    reading of small print behind it, because the model is optional. Either
    one saying so is enough -- they answer the same question from different
    evidence and neither is entitled to overrule the other into silence.
    """
    if chartparts.part_of_a_chart(shape, charts, here):
        return True
    if here is not None and here.reviewed:
        from ..ai.roles import normalize_copy  # noqa: PLC0415 - optional dep

        if normalize_copy(_text_of(shape)) in here.protected_texts:
            return True
    return reads_as_a_note(shape, canvas)


def _subtitle_text(here: Optional[Any]) -> Optional[str]:
    """The copy of the one box the model called the subtitle.

    None when there are no roles -- `claim_drawn_over` then behaves as it did,
    filling a subtitle region from whatever is drawn over it. An EMPTY STRING
    when the slide was read and no subtitle was named, which is a different
    answer and has to be: it means somebody looked and there is no standfirst
    on this slide, so the subtitle region is left empty rather than filled
    with a chart's caption.
    """
    if here is None or not here.reviewed:
        return None
    found = here.subtitle
    return found.text if found is not None else ""


def claim_drawn_over(
    loose: Sequence[Any], pool: list[Any], subtitle: Optional[str] = None,
) -> dict[int, Any]:
    """Fill a region from the single box drawn on top of it.

    THE GAP THIS CLOSES. `claim_runs` above fills from RUNS, and a run is three
    or more like-sized boxes, because that is where a run stops being a
    coincidence. A slide's subtitle is one box. So on a real deck the layout's
    subtitle region sat empty, showing its own prompt text, with the deck's
    subtitle drawn across the top of it in a loose box -- the words on the
    slide twice over, once as a prompt and once as content.

    One box is not an arrangement, so there is no order to read. What there is
    instead is geometry, and here it is unambiguous in a way it never was for
    the agenda: the layout's region is 12.28in wide at 0.69in from the edge,
    and the loose box is 12.28in wide at 0.67in. The same rectangle, drawn
    twice. See `series.drawn_over` for why width decides that and height is
    deliberately ignored.

    Runs get first refusal, because a run is the stronger statement: it says
    what a group of boxes IS, where this says only where one of them sits.

    `subtitle` IS THE ONE CASE GEOMETRY GETS WRONG EVERY TIME, and it is the
    case this function was written for. The subtitle region runs the width of
    the page, so EVERY full-width box on the slide is "drawn over" it: the
    standfirst, the chart's caption, the source line at the foot. Which of
    them is the subtitle is not in the file. Handed the copy of the box the
    model called the subtitle, only that box may take a subtitle region.
    Handed an empty string -- the slide was read and there is no standfirst on
    it -- no box may, and the region is left for the designer rather than
    filled with a footnote. Handed None nothing is read and this behaves as it
    always did.
    """
    if not loose or not pool:
        return {}
    free = [
        shape for shape in pool
        if _family_of(shape) in ("content", "subtitle", "title")
        and not _has_copy(shape)
    ]
    if not free:
        return {}

    from .series import pair_over  # noqa: PLC0415 - cycle

    claims: dict[int, Any] = {}
    for shape, region in pair_over(loose, free):
        if not _may_take(shape, region, subtitle):
            continue
        claims[id(shape)] = region
        pool.remove(region)
        log.info(
            "filled %r from the box drawn over it", _name_of(region),
        )
    return claims


def claim_subtitle(
    loose: Sequence[Any], pool: list[Any], subtitle: Optional[str]
) -> dict[int, Any]:
    """Put the box the model called the subtitle into the subtitle region.

    WHY GEOMETRY IS NOT ENOUGH ON ITS OWN. `claim_drawn_over` fills a region
    from the box drawn on top of it, and on a tidy deck the subtitle IS drawn
    over the subtitle region -- that is the case it was written for. On a
    messy one it is not: the old master put its standfirst two inches lower,
    or in a narrower measure, or the slide has no standfirst region at all and
    somebody typed the line into a loose box halfway down the page. The words
    are unmistakably a standfirst to anybody looking at the slide and there is
    nothing in the file that says so.

    So where the model has named one, it is claimed by NAME rather than by
    position, and the region it goes into is the master's subtitle region --
    which is where a standfirst belongs on the page it is being moved onto.

    Runs and drawn-over boxes get first refusal, as they do everywhere else
    here: this only ever fills a region nothing else claimed, and only ever
    from a box nothing else claimed.

    Empty whenever there is nothing to decide -- no roles, no subtitle named,
    no free subtitle region, or no loose box carrying that copy.
    """
    if not subtitle or not loose or not pool:
        return {}
    free = [
        shape for shape in pool
        if _family_of(shape) == "subtitle" and not _has_copy(shape)
    ]
    if not free:
        return {}

    from ..ai.roles import normalize_copy  # noqa: PLC0415 - optional dep

    for shape in loose:
        if normalize_copy(_text_of(shape)) != subtitle:
            continue
        region = free[0]
        pool.remove(region)
        log.info(
            "filled %r from %r, the box that reads as the subtitle",
            _name_of(region), _name_of(shape),
        )
        return {id(shape): region}
    return {}


def _may_take(shape: Any, region: Any, subtitle: Optional[str]) -> bool:
    """Whether this box is allowed to fill this region. See `claim_drawn_over`."""
    if subtitle is None or _family_of(region) != "subtitle":
        return True
    from ..ai.roles import normalize_copy  # noqa: PLC0415 - optional dep

    if not subtitle:
        log.info(
            "leaving %r empty: nothing on this slide reads as a subtitle",
            _name_of(region),
        )
        return False
    if normalize_copy(_text_of(shape)) == subtitle:
        return True
    log.info(
        "not filling %r from %r: it is not the box that reads as the subtitle",
        _name_of(region), _name_of(shape),
    )
    return False


def _is_imagery(shape: Any) -> bool:
    """A picture or a group of them, as against a rule or a blank panel.

    The distinction decides what survives a run moving. An icon beside an
    agenda item is content a designer drew and meant; the rule under it is the
    old layout's way of separating two rows, and the new layout draws its own.
    So one follows the item to its new home and the other goes.
    """
    try:
        name = str(shape.shape_type or "").upper()
    except Exception:
        return False
    return "PICTURE" in name or "GROUP" in name


def _move_imagery(
    run, slots, pictures, regions, pool, companion_of, belongs_to, fit_into
) -> dict[int, tuple[int, int, int, int]]:
    """Put each item's icon in the region the layout pairs with its new slot.

    A designer who draws a list of twelve labels usually draws twelve
    somethings beside them -- an ordinal, a rule, a swatch -- and on the master
    this was built against each agenda label has a 0.48in box at the same row,
    which is where its number goes. It is also exactly where the item's icon
    belongs, now that the item has moved.

    ONE ICON PER ITEM OR NONE AT ALL. If the count does not come out exactly --
    two icons over one card, or icons over only some of them -- the mapping is
    a guess, and a guess scatters a deck's iconography across a list. The whole
    lot then falls back to being treated as the old layout's drawing.
    """
    if not pictures or len(pictures) != len(run):
        return {}
    companion = companion_of(slots, regions, taken=slots.shapes)
    if companion is None:
        return {}

    owner = belongs_to(run, pictures)
    if sorted(owner.values()) != list(range(len(run))):
        return {}          # not one each; see above

    moved: dict[int, tuple[int, int, int, int]] = {}
    for shape in pictures:
        target = companion.shapes[owner[id(shape)]]
        moved[id(shape)] = fit_into(shape, target)
        if target in pool:
            pool.remove(target)
    return moved


def _record_filled(
    result: "RebuildResult",
    filled: dict[int, list[str]],
    shed: dict[int, list[tuple[str, str]]],
) -> None:
    """Say on each slide's record which regions its own runs filled.

    The PowerPoint route reports what PowerPoint did, and this happens after
    it, so without this the report would describe a file that is no longer the
    one on disk -- naming as unfilled the regions a designer can now see copy
    sitting in.
    """
    for record in result.slides:
        names = filled.get(record.number)
        if names:
            record.filled.extend(names)
            record.unfilled = [
                name for name in record.unfilled if name not in names
            ]
        for name, reason in shed.get(record.number, []):
            record.dropped.append(
                DroppedShape(slide=record.number, name=name, reason=reason)
            )


def layout_writes(layout: Any) -> set[str]:
    """The words a layout puts on the page itself, normalised for comparison.

    Its own shapes, not its placeholders: a placeholder is a space for someone
    else's copy, and what is wanted here is the copy the layout ALREADY has.
    """
    words = set()
    try:
        shapes = list(layout.shapes)
    except Exception:
        return words
    for shape in shapes:
        if _is_placeholder(shape):
            continue
        text = _normalised(shape)
        if text:
            words.add(text)
    return words


def _normalised(shape: Any) -> str:
    try:
        if not shape.has_text_frame:
            return ""
        text = shape.text_frame.text
    except Exception:
        return ""
    kept = [c.lower() for c in text if c.isalnum() or c.isspace()]
    return " ".join("".join(kept).split())


def echoes_layout(shape: Any, written: set[str]) -> bool:
    """Whether this shape says only what the layout already says.

    THE SLIDE IT IS FOR. A deck's agenda carries its title in a TITLE
    placeholder, and the master's agenda layout has no title region at all --
    it draws the word AGENDA itself, as artwork. So the title finds no home,
    is transplanted at its old inches, and the rebuilt slide says "Agenda"
    twice: once in the layout's own display type and once in the deck's, over
    the layout's photograph.

    Matched on the whole of the text and not on a part of it, with case,
    punctuation and spacing taken out. A heading that repeats the layout's
    word and then says more is not an echo, it is a heading.
    """
    if not written:
        return False
    return _normalised(shape) in written


# How much of the canvas a layout shape must cover before content sitting on it
# is sitting on the LAYOUT rather than on empty page. A hairline rule and a
# 0.67in logo come to a fraction of a percent; the arcs and panels that swallow
# an icon run from a quarter of the page upward.
_ARTWORK_SHARE = 0.06

# And how large a graphic may be and still be something that labels content
# rather than being content. An icon is under a percent of the page; a
# photograph is a third of it, and moving a photograph is never this function's
# business.
_LABEL_SHARE = 0.02

# Clear air between a lifted graphic and the copy it labels.
_GAP_IN = 0.15

# How much of a graphic has to be inside a piece of artwork before it counts as
# sitting ON it. Touching is not enough and was the original test: a master
# that draws a photograph clipped to the whole page made every shape on the
# slide read as sitting on artwork, and a row of flag roundels was moved out of
# the panel holding it on the strength of that.
_BURIED_SHARE = 0.6

# How far a graphic may travel to reach the copy it labels. A label sits
# directly above its copy; anything that has to cross an inch and a half of
# page to find the nearest copy below it has found the footer, not its label.
_REACH_IN = 1.5


def lift_imagery(slide: Any, canvas: tuple[float, float]) -> list[tuple[str, float]]:
    """Move small graphics off the layout's own artwork, onto what they label.

    THE SLIDE THIS IS FOR. A deck's two-track page carried an icon above each
    column, sitting on the edge of the arc its OLD layout drew. The master's arc
    is bigger, so on the rebuilt slide both icons sat well inside it: dark line
    art on a dark disk, still there and impossible to see. Nothing was wrong
    with the geometry and nothing was wrong with the layout choice. The slide
    was simply drawn against a different arc.

    IT IS NOT ENOUGH TO ASK WHETHER A SHAPE IS ON THE ARTWORK. That layout puts
    its own title and body regions inside the same arc, in reversed type, on
    purpose -- so "sits on the artwork" does not separate wrong from intended.
    What separates them is that these graphics belong to something: each one is
    centred over a column of copy, and the place it belongs is directly above
    that copy. So a graphic is lifted only when it is small enough to be a
    label, is sitting on artwork big enough to swallow it, and has copy beneath
    it that it is centred over. Anything with nothing underneath it is left
    alone, because there is nothing to say where it should go instead.

    Returns what moved and how far, in inches, for the record.
    """
    width, height = canvas
    if width <= 0 or height <= 0:
        return []
    page = width * height

    try:
        shapes = list(slide.shapes)
        artwork = [
            _rect(shape, width, height) for shape in slide.slide_layout.shapes
            if not _is_placeholder(shape)
        ]
    except Exception:
        log.debug("could not read the layout's artwork", exc_info=True)
        return []
    artwork = [box for box in artwork if box and _area(box) >= _ARTWORK_SHARE * page]
    if not artwork:
        return []

    loose = [shape for shape in shapes if not _is_placeholder(shape)]
    labels = [
        shape for shape in loose
        if _is_imagery(shape) and _area(_rect(shape, width, height) or (0, 0, 0, 0))
        <= _LABEL_SHARE * page
    ]
    copy = [shape for shape in loose if _has_copy(shape)]

    moved: list[tuple[str, float]] = []
    for cluster in _clusters(labels, width, height):
        box = _union([_rect(shape, width, height) for shape in cluster])
        if box is None:
            continue
        # BURIED IN the artwork, not merely touching it. Touching was the
        # original test and it is what made this move a row of flag roundels
        # out of the panel holding them: they sat at the bottom of a page
        # whose master draws a photograph clipped to the whole of it, so
        # "overlaps some artwork" was true of every shape on the slide.
        sitting = [art for art in artwork if _buried(box, art)]
        if not sitting:
            continue
        target = _copy_below(box, copy, width, height)
        if target is None:
            continue
        lift = (target[1] - _GAP_IN) - box[3]
        if lift <= 0:
            continue                  # already clear of the copy it labels
        if lift > _REACH_IN:
            # A label sits directly above what it labels. Something that has
            # to travel further than this to reach the nearest copy beneath it
            # is not labelling that copy -- it is a graphic with the footer
            # somewhere below it, which is true of most graphics on most
            # slides.
            continue
        landing = (box[0], box[1] + lift, box[2], box[3] + lift)
        if any(_buried(landing, art) for art in sitting):
            # THE MOVE HAS TO ACHIEVE THE THING THE MOVE IS FOR. This exists
            # to get a graphic off artwork it is invisible against; a
            # destination still buried in that same artwork has not done that,
            # it has just moved the graphic down the page and out of whatever
            # was holding it. That is the whole of the flag-roundel bug: the
            # flags and the footer line they were dragged onto were inside the
            # same full-page picture, so the lift cleared nothing at all.
            continue
        for shape in cluster:
            shape.top = int(shape.top + lift * 914400)
            moved.append((_name_of(shape), lift))
    return moved


def _side_by_side(box, other) -> bool:
    """Two graphics reading as one row: same band, and a gap under their width.

    The gap is measured against the narrower of the two, so a row of equal
    roundels joins and a small icon does not reach across the page to a
    distant one merely because they happen to share a band.
    """
    if other is None:
        return False
    # Sharing a row means overlapping vertically at all: a row of flags is
    # seldom pixel-aligned and a band test with a fixed tolerance would miss
    # the ones a designer nudged.
    if box[3] <= other[1] or other[3] <= box[1]:
        return False
    gap = max(box[0], other[0]) - min(box[2], other[2])
    if gap < 0:
        return True                       # they already touch horizontally
    return gap <= min(box[2] - box[0], other[2] - other[0])


def _buried(box, art) -> bool:
    """Whether this box is sitting ON that artwork rather than beside it.

    By area rather than by touching: an icon resting against the edge of a
    band is not invisible against it, and the case this whole function exists
    for is dark line art well inside a dark disk.
    """
    inside = _area((max(box[0], art[0]), max(box[1], art[1]),
                    min(box[2], art[2]), min(box[3], art[3])))
    own = _area(box)
    return own > 0 and inside / own >= _BURIED_SHARE


def _canvas_in(base: Any) -> tuple[float, float]:
    try:
        return ((base.slide_width or 0) / 914400,
                (base.slide_height or 0) / 914400)
    except Exception:
        return (0.0, 0.0)


def _rect(shape: Any, width: float, height: float):
    """A shape's box in inches, clipped to the canvas, or None if it has none.

    CLIPPED, because a layout's artwork is routinely drawn far off the page: the
    photograph inside one master's arc is a 13.33in-tall picture starting five
    inches above the top edge, and taken at face value its bounding box covers
    everything, so every shape on the slide would read as sitting on it.
    """
    try:
        left = (shape.left or 0) / 914400
        top = (shape.top or 0) / 914400
        right = left + (shape.width or 0) / 914400
        bottom = top + (shape.height or 0) / 914400
    except Exception:
        return None
    left, top = max(0.0, left), max(0.0, top)
    right, bottom = min(width, right), min(height, bottom)
    if right <= left or bottom <= top:
        return None
    return (left, top, right, bottom)


def _area(box) -> float:
    if not box:
        return 0.0
    return (box[2] - box[0]) * (box[3] - box[1])


def _overlaps(a, b) -> bool:
    return not (a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1])


def _union(boxes):
    boxes = [box for box in boxes if box]
    if not boxes:
        return None
    return (min(b[0] for b in boxes), min(b[1] for b in boxes),
            max(b[2] for b in boxes), max(b[3] for b in boxes))


def _clusters(shapes, width: float, height: float) -> list[list[Any]]:
    """Graphics that belong together, grouped, so none is moved on its own.

    TWO WAYS OF BELONGING, and the second was missing. The two icons on the
    slide this was written for are each a ring and a glyph sitting inside it,
    two top-level shapes apiece, so overlapping shapes group -- lifting one and
    not the other would take an icon apart.

    A ROW OF SEPARATE GRAPHICS IS ALSO ONE THING. Six flag roundels in a line
    touch nothing, so each was a cluster of its own and each was judged alone:
    on a real slide three of five moved and two stayed, which is worse than
    either moving the row or leaving it, because a reader cannot tell a row
    that was rearranged from one that was always ragged. So neighbours on the
    same row join, where "neighbour" is a gap narrower than the graphics
    themselves -- close enough to read as a set, and far short of joining two
    icons that head two different columns half a page apart.
    """
    groups: list[list[Any]] = []
    for shape in shapes:
        box = _rect(shape, width, height)
        if box is None:
            continue
        for group in groups:
            if any(
                _overlaps(box, _rect(other, width, height) or box)
                or _side_by_side(box, _rect(other, width, height))
                for other in group
            ):
                group.append(shape)
                break
        else:
            groups.append([shape])
    return groups


def _copy_below(box, copy: Sequence[Any], width: float, height: float):
    """The nearest copy under this graphic that it is centred over, or None."""
    middle = (box[0] + box[2]) / 2
    best = None
    for shape in copy:
        other = _rect(shape, width, height)
        if other is None or other[1] < box[3]:
            continue                          # not below it
        if not (other[0] <= middle <= other[2]):
            continue                          # not what it sits over
        if best is None or other[1] < best[1]:
            best = other
    return best


RUN_DRAWING = "drew a run of content that the layout now draws itself"
LAYOUT_ECHO = "says only what the layout already writes on the page"


def fill_runs(
    path: Path, skip: Collection[int] = (), roles: Optional[Any] = None,
) -> tuple[dict[int, list[str]], dict[int, list[tuple[str, str]]]]:
    """Fill a written deck's repeated regions from its own loose runs.

    THE POWERPOINT ROUTE NEEDS THIS AND THE REBUILD DOES NOT. Assigning a
    CustomLayout runs PowerPoint's own placeholder matching, which is why that
    route is preferred -- but that matching only ever moves a PLACEHOLDER into
    a placeholder. A messy deck's content is in loose boxes, so on a real deck
    PowerPoint put all twenty-four of the agenda layout's slots on the slide,
    empty, and left the five agenda items sitting at their old inches on top of
    the layout's photograph. The XML route does this inline in `_rebuild_slide`;
    here it has to be a second pass over the file PowerPoint wrote.

    `roles` is `ai.roles.RolesResult`: what the model said each shape on the
    original slide IS. WITHOUT IT THIS FILLS BY GEOMETRY ALONE, which is how
    the source line ends up under the title. A footnote is a full-width box,
    so it measures exactly like a subtitle, and `claim_drawn_over` cannot tell
    the two apart from the file -- nothing in a .pptx says "this is small
    print". A chevron banner reads the same way: a box with words in it,
    claimed into a body region, and the chevron deleted with the box its copy
    came out of. Both are visible in a picture and in nothing else.

    So a shape the model called a source, a chart part or decoration is never
    a candidate here, and the subtitle region is filled from the one box the
    model called the subtitle or from nothing at all. Handed no roles, the
    deterministic reading in `reads_as_a_note` holds the worst of it back and
    everything else behaves as it did.

    The file IS re-serialised by this, which the PowerPoint route otherwise
    avoids on purpose. It is a narrow re-serialisation -- text set into
    placeholders that are already on the slide, and the consumed boxes removed
    -- and nothing is copied between parts, which is where that route's
    corruption came from. It also only happens at all on a slide where a run
    was actually matched, so a deck this cannot help is written by PowerPoint
    and not touched again.
    """
    from pptx import Presentation  # noqa: PLC0415 - lazy heavy dependency

    presentation = Presentation(str(path))
    canvas = (
        (presentation.slide_width or 0) / 914400,
        (presentation.slide_height or 0) / 914400,
    )
    filled: dict[int, list[str]] = {}
    shed: dict[int, list[tuple[str, str]]] = {}
    touched: set[int] = set()
    # Decided once for the file, not per slide. Which way a run is read is a
    # property of the deck, and a slide of mostly numbers in an Arabic deck
    # would otherwise be filled in the other direction from the one before it.
    rtl = _reads_rtl(presentation)
    if rtl:
        log.info("%s reads right to left; its runs are filled that way", path.name)
    for number, slide in enumerate(presentation.slides, start=1):
        if number in skip:
            # Handed back to a designer untouched. See `rebuild.quarantine`.
            continue
        try:
            # One list, for the reason `_rebuild_slide` gives: a fresh proxy
            # per access makes `id()` meaningless across two walks.
            shapes = list(slide.shapes)
            charts_here = chartparts.chart_boxes(shapes)
            pool = [
                p for p in slide.placeholders if not chartparts.is_chart(p)
            ]
            here = roles.for_slide(number) if roles is not None else None
            loose = [
                shape for shape in shapes
                if not _is_placeholder(shape) and _has_copy(shape)
                and not _leave_alone(shape, here, canvas, charts_here)
            ]
            claims, furniture, moves = claim_runs(loose, pool, shapes, rtl=rtl)
            subtitle = _subtitle_text(here)
            claims.update(claim_drawn_over(
                [s for s in loose if id(s) not in claims], pool,
                subtitle=subtitle,
            ))
            claims.update(claim_subtitle(
                [s for s in loose if id(s) not in claims], pool, subtitle,
            ))
            written = layout_writes(slide.slide_layout)
        except Exception:
            log.debug("could not settle slide %d", number, exc_info=True)
            continue

        drawing = {id(shape) for shape in furniture} - set(moves)
        names: list[str] = []
        gone: list[tuple[str, str]] = []

        for shape in shapes:
            target = claims.get(id(shape))
            if target is not None:
                # The box the copy came out of goes with the copy, or the slide
                # says it twice: once where the layout puts it and once where
                # it used to be, which is worse than either on its own.
                #
                # ONLY IF THE COPY ACTUALLY LANDED. `_copy_text` returns
                # silently when the target has no text frame, and this deleted
                # the source regardless -- so a run claimed by a table, chart
                # or media region (all of them "content" in `FAMILIES`, none of
                # them able to hold a paragraph) had its words copied nowhere
                # and its boxes removed, and the slide was reported as filled.
                # Losing copy is the one outcome this whole module is not
                # allowed to produce.
                if _copy_text(shape, target):
                    names.append(_name_of(target))
                    _remove(shape)
                else:
                    log.warning(
                        "slide %d: %r was claimed by %r, which cannot hold "
                        "text; the copy was left where it is rather than "
                        "deleted",
                        number, _name_of(shape), _name_of(target),
                    )
            elif _leave_alone(shape, here, canvas, charts_here):
                # NOTHING PROTECTED IS DELETED BY ANY ROUTE. Keeping it out of
                # `loose` stops it being claimed into a region; it does not
                # stop `furniture_of` sweeping it up as the run's own drawing,
                # or `echoes_layout` reading it as a title the layout already
                # writes. Both delete, and a chevron banner deleted as a run's
                # drawing is exactly as gone as one deleted with its copy.
                pass
            elif id(shape) in drawing:
                # The run's own drawing goes for the same reason: the new
                # layout draws its own. See `series.furniture_of`.
                _remove(shape)
                gone.append((_name_of(shape), RUN_DRAWING))
            elif echoes_layout(shape, written):
                # Separate defect, same slide: a title the layout already
                # writes as artwork, carried over because the layout offers no
                # title region for it to land in. Reported as itself, not as
                # the run's drawing -- a reader checking the work needs to know
                # which of the two happened.
                _remove(shape)
                gone.append((_name_of(shape), LAYOUT_ECHO))
            elif id(shape) in moves:
                shape.left, shape.top, shape.width, shape.height = moves[id(shape)]

        lifted = lift_imagery(slide, canvas)
        if lifted:
            log.info(
                "slide %d: lifted %d graphic(s) off the layout's artwork onto "
                "the copy they label", number, len(lifted),
            )

        if names:
            filled[number] = names
        if gone:
            shed[number] = gone
        if names or gone or moves or lifted:
            touched.add(number)

    if touched:
        presentation.save(str(path))
        log.info(
            "settled %d slide(s): %d region(s) filled from the deck's own runs, "
            "%d shape(s) removed as the layout's own work",
            len(touched),
            sum(len(names) for names in filled.values()),
            sum(len(names) for names in shed.values()),
        )
    return filled, shed


def _reposition(new_slide: Any, where) -> None:
    """Move the shape just transplanted into the box a run's companion gives it.

    Addressed as the last shape on the new slide, because that is what
    `_transplant` appends and the source shape itself is in the other file.
    """
    if where is None:
        return
    try:
        shape = list(new_slide.shapes)[-1]
        shape.left, shape.top, shape.width, shape.height = where
    except Exception:
        log.debug("could not reposition a run's imagery", exc_info=True)


def _remove(shape: Any) -> None:
    parent = shape._element.getparent()
    if parent is not None:
        parent.remove(shape._element)


def _has_copy(shape: Any) -> bool:
    """Whether this shape carries copy, rather than being furniture.

    A run is matched on the boxes, and an empty box is the same size as a full
    one. The separator rules and blank panels a designer draws behind a row of
    items are exactly that shape, and a layout region filled from one of them
    is a region filled with nothing -- while the item that should have gone
    there keeps its old inches.
    """
    try:
        return bool(shape.has_text_frame and shape.text_frame.text.strip())
    except Exception:
        return False


def _affinity(source: Any, target: Any, canvas: tuple[int, int]) -> float:
    """How much these two boxes look like the same place on the slide.

    Overlap first, because two boxes that share ground are the same slot
    however differently they are sized. Where nothing overlaps -- a messy
    deck's item sitting where the new layout puts nothing -- the nearer of the
    free slots is the better guess, so the distance between centres decides,
    negated so that closer scores higher and always below any real overlap.
    """
    a, b = _box_of(source), _box_of(target)
    if a is None or b is None:
        return -1e9
    width, height = canvas
    if width > 0 and height > 0:
        a = (a[0] / width, a[1] / height, a[2] / width, a[3] / height)
        b = (b[0] / width, b[1] / height, b[2] / width, b[3] / height)

    overlap = (
        max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
        * max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    )
    if overlap > 0:
        union = (
            (a[2] - a[0]) * (a[3] - a[1])
            + (b[2] - b[0]) * (b[3] - b[1])
            - overlap
        )
        return overlap / union if union > 0 else 0.0

    centre_a = ((a[0] + a[2]) / 2, (a[1] + a[3]) / 2)
    centre_b = ((b[0] + b[2]) / 2, (b[1] + b[3]) / 2)
    distance = (
        (centre_a[0] - centre_b[0]) ** 2 + (centre_a[1] - centre_b[1]) ** 2
    ) ** 0.5
    return -distance


def _box_of(shape: Any) -> Optional[tuple[float, float, float, float]]:
    """(left, top, right, bottom), or None where the shape will not say.

    A placeholder that states no position of its own answers with the value it
    inherits from its layout, which is the position it will be drawn at and so
    the one to match on.
    """
    try:
        left, top = shape.left, shape.top
        width, height = shape.width, shape.height
    except Exception:
        return None
    if None in (left, top, width, height):
        return None
    return (float(left), float(top), float(left + width), float(top + height))


def _canvas_of(presentation: Any) -> tuple[int, int]:
    """The output canvas, for measuring both decks in the same units.

    A deck built at 4:3 and a master at 16:9 state the same place with
    different numbers, so boxes are compared as fractions of their own canvas
    rather than in EMU.
    """
    try:
        return int(presentation.slide_width or 0), int(presentation.slide_height or 0)
    except Exception:
        return (0, 0)


def _family_of(shape: Any) -> Optional[str]:
    try:
        raw = str(shape.placeholder_format.type or "")
    except Exception:
        return None
    token = raw.split("(")[0].strip().upper()
    if not token or token in LATENT:
        return None
    return FAMILIES.get(token)


def _copy_text(source: Any, target: Any) -> bool:
    """Move the copy across, keeping emphasis and dropping brand formatting.

    Bold, italic and underline are part of what the writer meant, so they
    survive. Typeface, size and colour are the layout's business, and carrying
    them over would reproduce exactly the drift this is meant to remove.

    Returns whether the copy landed. The caller deletes the box it came out
    of, so "it did not land" and "it landed" must be tellable apart: this used
    to return None either way and the source was deleted regardless, which
    turned a region that cannot hold text into lost copy.
    """
    if not _has_text(source) or not _has_text(target):
        return False

    frame = target.text_frame
    frame.clear()

    for index, src_para in enumerate(source.text_frame.paragraphs):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.level = src_para.level
        for src_run in src_para.runs:
            run = para.add_run()
            run.text = src_run.text
            for attr in ("bold", "italic", "underline"):
                value = getattr(src_run.font, attr, None)
                if value is not None:
                    setattr(run.font, attr, value)
            _copy_hyperlink(src_run, run)
    return True


def _copy_hyperlink(src_run: Any, run: Any) -> None:
    try:
        address = src_run.hyperlink.address
    except Exception:
        return
    if not address:
        return
    try:
        run.hyperlink.address = address
    except Exception:
        log.debug("could not carry the hyperlink %r across", address)


def _has_text(shape: Any) -> bool:
    try:
        return bool(shape.has_text_frame)
    except Exception:
        return False


def _copy_notes(src_slide: Any, new_slide: Any) -> None:
    """Carry the speaker notes across as plain text.

    Notes are written for the presenter, not the audience, so their formatting
    is not a brand concern and only the words are worth moving.
    """
    try:
        if not src_slide.has_notes_slide:
            return
        text = src_slide.notes_slide.notes_text_frame.text
    except Exception:
        return
    if not text.strip():
        return
    try:
        new_slide.notes_slide.notes_text_frame.text = text
    except Exception:
        log.debug("could not carry notes across for slide")


# --------------------------------------------------------------------------- #
# Loose shapes
# --------------------------------------------------------------------------- #

def _transplant(
    shape: Any, src_slide: Any, new_slide: Any, record: SlideRecord
) -> None:
    """Copy a shape's XML across, rebuilding the parts it points at.

    A shape referencing a part this cannot rebuild is left behind rather than
    copied broken: a chart pointing at a chart part that is not in the new file
    opens as an error in PowerPoint, which is harder to notice and to fix than
    an absence the report names.
    """
    name = _name_of(shape)
    element = copy.deepcopy(shape._element)
    if chartparts.is_chart(element):
        # Floating on the new slide at its own frame, not a slot of a layout
        # that never drew it. See `charts.unpin`.
        ph = element.find(f".//{_P_NS}nvPr/{_P_NS}ph")
        if ph is not None:
            ph.getparent().remove(ph)

    unsupported = _remap_relationships(element, src_slide.part, new_slide.part)
    if unsupported:
        record.dropped.append(
            DroppedShape(
                slide=record.number,
                name=name,
                reason=f"references {', '.join(sorted(unsupported))}",
            )
        )
        return

    new_slide.shapes._spTree.append(element)
    record.transplanted.append(name)


def _remap_relationships(element: Any, src_part: Any, tgt_part: Any) -> set[str]:
    """Repoint every r: reference in the copied XML at the new package.

    Returns the relationship kinds that could not be rebuilt. Images are
    re-related to the same image part, which the target package pulls in when
    it is saved. External references (hyperlinks, linked images) are recreated
    from their URL. Anything else points at a part with its own internal
    structure -- a diagram, an embedded object -- and is reported rather than
    copied. A chart is the exception: its parts are cloned as a set, so the
    chart crosses as the one element it is.
    """
    unsupported: set[str] = set()

    for node in list(element.iter()):
        for attr, rid in list(node.attrib.items()):
            if not attr.startswith(_R_NS):
                continue

            # An empty r:id is the idiom for "no hyperlink here", left behind
            # when one is removed. It is not a reference and there is nothing
            # to rebuild; treating it as a broken one cost fifteen shapes on a
            # real deck, including a cover title.
            if not rid:
                continue

            if _is_tag(node):
                # Customer-data tags carry no content. Dropping the shape over
                # metadata loses the shape; dropping the metadata loses
                # nothing anybody can see.
                _detach(node)
                break

            rel = _relationship(src_part, rid)
            if rel is None:
                unsupported.add("a missing relationship")
                continue
            try:
                if rel.is_external:
                    new_rid = tgt_part.relate_to(
                        rel.target_ref, rel.reltype, is_external=True
                    )
                elif _kind(rel.reltype) in _IMAGE_KINDS:
                    new_rid = _reimport_image(rel, tgt_part)
                elif _kind(rel.reltype) in chartparts.CHART_KINDS:
                    # Copied whole, workbook and all -- see `charts.clone_chart`.
                    new_rid = chartparts.clone_chart(rel, tgt_part)
                else:
                    unsupported.add(_kind(rel.reltype))
                    continue
            except Exception:
                if _is_alternate_image(node):
                    # A high-definition or vector copy of an image the shape
                    # already carries in `a:blip`. Losing the enhancement is
                    # invisible; losing the shape is two missing portraits.
                    log.debug("dropping an alternate image reference on %s", node.tag)
                    _detach(node)
                    break
                unsupported.add(_kind(rel.reltype))
                continue
            node.set(attr, new_rid)

    return unsupported


# Parts that are a picture and nothing else, so copying the bytes copies
# everything. `hdphoto` is the Windows Media Photo alternate PowerPoint writes
# alongside a JPEG; `svg` is the vector original beside its raster fallback.
_IMAGE_KINDS = frozenset({"image", "hdphoto"})

_ALTERNATE_TAGS = ("imgLayer", "svgBlip")


def _is_tag(node: Any) -> bool:
    return node.tag.rsplit("}", 1)[-1] == "tags"


def _is_alternate_image(node: Any) -> bool:
    return node.tag.rsplit("}", 1)[-1] in _ALTERNATE_TAGS


def _detach(node: Any) -> None:
    """Remove a node, and any container left behind with nothing in it.

    Dropping the node alone is not enough. A reference like the saturation
    layer on a picture lives at the bottom of a chain that exists only to hold
    it:

        a:extLst / a:ext / a14:imgProps / a14:imgLayer r:embed="rId10"

    Take out the imgLayer and `a14:imgProps` is left empty, which its schema
    does not allow, and PowerPoint rejects the whole file for it -- "could not
    open the file", no part named. One of these on one slide cost all five
    slides of a real deck.

    So a reference inside an extension takes the whole `a:ext` with it. That
    is the right unit and the only safe one: an extension is a self-contained
    opaque payload, and both `a:ext` and `a:extLst` are permitted to be absent
    entirely, so removing one can never leave a hole the schema minds.

    Outside an extension the node alone goes. Walking up removing emptied
    parents there does more harm than good -- on the same deck it emptied a
    fill on slide 1 and moved the corruption rather than fixing it -- because
    those parents are ordinary required content, not opaque payload.
    """
    parent = node.getparent()
    if parent is None:
        return
    extension = _extension_of(node)
    if extension is not None:
        holder = extension.getparent()
        if holder is not None:
            holder.remove(extension)
            return
    parent.remove(node)


def _extension_of(node: Any) -> Optional[Any]:
    """The `a:ext` this node sits inside, if any."""
    current = node.getparent()
    while current is not None:
        if current.tag.rsplit("}", 1)[-1] == "ext":
            return current
        current = current.getparent()
    return None


def _reimport_image(rel: Any, tgt_part: Any) -> str:
    """Copy the image bytes into the target package and return the new rId.

    Relating the source part directly looks like it works and quietly corrupts
    the file: the foreign part keeps its own partname, and the master almost
    always already has a `ppt/media/image1.png` of its own. Both then write to
    the same name in the zip and one image silently replaces the other.

    Going through the bytes lets the target package allocate a free partname,
    and it deduplicates by content, so a logo transplanted onto twenty slides
    is stored once.

    It also re-sniffs the bytes, which is where this got dangerous. python-pptx
    asks Pillow what the blob is, and Pillow's WMF reader accepts EMF as well
    and answers "WMF" for it. An EMF picture therefore came back typed
    image/x-wmf and was written as image8.wmf. PowerPoint reads the extension,
    tries to parse EMF bytes as WMF, and refuses the entire file:
    "PowerPoint could not open the file", naming no part. One such picture on
    one slide of a real deck cost all five slides.

    So the type is checked before anything is added. A blob the target would
    label differently than the source did is refused, and the caller drops the
    shape and reports it, which is this module's answer to everything it cannot
    rebuild faithfully. Checked first because get_or_add_image_part creates the
    relationship as well as the part: raising afterwards would leave the
    mislabelled part in the package, and content types are global.
    """
    source_type = _content_type(rel.target_part)
    blob = rel.target_part.blob
    target_type = _sniffed_type(blob)
    if target_type is None:
        # Nothing recognised the bytes, which is not the same as bad bytes.
        # An icon out of PowerPoint's own library is an SVG with an EMF
        # fallback, and Pillow identifies neither, so `get_or_add_image_part`
        # raised on both halves of every one of them: the SVG was dropped as
        # an "invisible enhancement" and the EMF was left holding the source's
        # relationship id, which means nothing on the new slide. The icon
        # arrived as a broken-image box -- "The picture can't be displayed" --
        # and only for library icons, which is why some survived and some did
        # not.
        #
        # These do not need identifying. They are valid parts already; what
        # they need is a partname in the target package and the content type
        # the source declared. See `_carry_media`.
        return _carry_media(rel, tgt_part, source_type)
    if source_type and source_type != target_type:
        raise MediaTypeMismatch(
            f"the source calls this {source_type} and re-importing it would "
            f"write {target_type}"
        )
    _image_part, rid = tgt_part.get_or_add_image_part(BytesIO(blob))
    return rid


def _carry_media(rel: Any, tgt_part: Any, source_type: Optional[str]) -> str:
    """Copy media bytes across verbatim, keeping the type the source declared.

    Used for the parts python-pptx will not sniff. It does the two things
    `get_or_add_image_part` was doing that actually matter -- a free partname
    in the target package, and dedup so an icon on twenty slides is stored
    once -- and skips the part that was in the way, asking Pillow what the
    bytes are.

    The type is taken from the source rather than guessed, which is the whole
    reason this is safe: the source package declared it, PowerPoint wrote it,
    and copying a declaration is not a decision. Where the source declares
    nothing there is nothing to copy faithfully, so it raises and the caller
    reports the shape as it would any other media it cannot rebuild.
    """
    from pptx.opc.package import Part  # noqa: PLC0415 - lazy, opc internals

    if not source_type:
        raise MediaTypeMismatch(
            "these bytes are of no type anything here can name, so copying "
            "them would be a guess"
        )

    blob = rel.target_part.blob
    package = tgt_part.package
    # Dedup by content, like the image path it stands in for. Held on the
    # package because that is what the parts belong to, and a rebuild makes
    # one of those per run.
    carried = getattr(package, "_carried_media", None)
    if carried is None:
        carried = {}
        package._carried_media = carried

    key = (source_type, hashlib.sha1(blob).hexdigest())
    part = carried.get(key)
    if part is None:
        extension = PurePosixPath(str(rel.target_part.partname)).suffix or ".bin"
        partname = package.next_partname(f"/ppt/media/image%d{extension}")
        part = Part(partname, source_type, package, blob)
        carried[key] = part
    return tgt_part.relate_to(part, rel.reltype)


class MediaTypeMismatch(Exception):
    """Re-importing a media part would relabel its bytes as another type."""


def _content_type(part: Any) -> Optional[str]:
    try:
        declared = part.content_type
    except Exception:
        return None
    # Not str(): a part with no declared type would come back as "None", which
    # is truthy and matches nothing, so every such blob read as a mismatch.
    return str(declared) if declared else None


def _sniffed_type(blob: bytes) -> Optional[str]:
    """The content type the target package would give these bytes.

    None when it cannot tell, which is not a mismatch: the re-import will
    raise on its own and the caller already handles that.
    """
    from pptx.parts.image import Image  # noqa: PLC0415 - lazy heavy dependency

    try:
        return str(Image.from_blob(blob).content_type)
    except Exception:
        return None


def _kind(reltype: str) -> str:
    return reltype.rsplit("/", 1)[-1] or reltype


def _relationship(part: Any, rid: str) -> Optional[Any]:
    try:
        return part.rels[rid]
    except Exception:
        return None


def _name_of(shape: Any) -> str:
    try:
        return str(shape.name)
    except Exception:
        return "unnamed shape"


# --------------------------------------------------------------------------- #
# Deck level
# --------------------------------------------------------------------------- #

def _drop_slides(base: Any) -> int:
    """Remove the master file's own sample slides.

    The master is being used as a template here, not as content. Its example
    slides are what a designer built to show the layouts off, and shipping them
    inside the rebuilt deck is never wanted.
    """
    slide_ids = base.slides._sldIdLst
    removed = 0
    for slide_id in list(slide_ids):
        base.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
        removed += 1
    return removed


def _size_note(base: Any, source: Any) -> Optional[str]:
    """Flag a canvas size change, which moves every transplanted shape.

    The master's size wins, because it is the approved one. Loose shapes keep
    their position in inches, so on a different canvas they land somewhere else
    relative to the edges and a person has to look at them.
    """
    if (
        base.slide_width == source.slide_width
        and base.slide_height == source.slide_height
    ):
        return None
    return (
        f"canvas changes from {source.slide_width / 914400:.2f} x "
        f"{source.slide_height / 914400:.2f}in to {base.slide_width / 914400:.2f} x "
        f"{base.slide_height / 914400:.2f}in; transplanted shapes keep their "
        f"position in inches and will need repositioning"
    )


# --------------------------------------------------------------------------- #
# Reset
# --------------------------------------------------------------------------- #

def reset_layouts(path: Path, skip: Collection[int] = ()) -> dict[int, int]:
    """Put every placeholder back on the geometry its layout gives it.

    THIS IS "RESET SLIDE", DONE TO THE FILE. A layout supplies a placeholder's
    position and size only while the placeholder does not state its own, and a
    messy deck states its own everywhere -- an explicit `a:xfrm` on every
    shape, written against whatever layout it used to sit on. Re-pointing the
    slide at a new layout does not clear those, so the copy lands in the new
    design still wearing the old one's boxes: a title that was 4.4in wide on
    its old layout stays 4.4in wide on a layout that gives titles the full
    width, and a title that was narrow stays narrow and reflows to one word a
    line down the side of the page.

    Only where the new layout actually has something to say. A placeholder with
    no counterpart on the layout it now points at inherits nothing, so
    stripping its `a:xfrm` would not reset it, it would send it to the origin
    at a default size. Those keep what they have and are left for the rest of
    the rebuild to deal with.

    Returns the count of placeholders reset per slide.
    """
    from pptx import Presentation  # noqa: PLC0415 - lazy heavy dependency

    presentation = Presentation(str(path))
    reset: dict[int, int] = {}
    for number, slide in enumerate(presentation.slides, start=1):
        if number in skip:
            # Still on its own layout, and still the designer's own geometry.
            # See `rebuild.quarantine`.
            continue
        try:
            offered = _layout_slots(slide.slide_layout)
            count = sum(
                1 for shape in list(slide.shapes)
                if _is_placeholder(shape)
                and _slot_of(shape) in offered
                and _drop_xfrm(shape)
            )
        except Exception:
            log.debug("could not reset slide %d", number, exc_info=True)
            continue
        if count:
            reset[number] = count

    if reset:
        presentation.save(str(path))
        log.info(
            "reset %d placeholder(s) across %d slide(s) onto the geometry "
            "their layout gives them",
            sum(reset.values()), len(reset),
        )
    return reset


def _layout_slots(layout: Any) -> set:
    """Which placeholder slots this layout defines, by idx and by type."""
    slots = set()
    try:
        shapes = list(layout.placeholders)
    except Exception:
        return slots
    for shape in shapes:
        slot = _slot_of(shape)
        if slot is not None:
            slots.add(slot)
    return slots


def _slot_of(shape: Any):
    """A placeholder's identity on its layout: its idx, or its type.

    Both, because PowerPoint matches on both. A title has no meaningful idx
    and is found by type; a body region is one of several and is found by idx.
    """
    try:
        fmt = shape.placeholder_format
        return (str(fmt.type).split()[0].upper(), fmt.idx)
    except Exception:
        return None


def _drop_xfrm(shape: Any) -> bool:
    """Remove a shape's own position and size, so it inherits again."""
    try:
        spPr = shape._element.spPr
    except Exception:
        return False
    if spPr is None:
        return False
    removed = False
    for xfrm in spPr.findall(f"{_A_NS}xfrm"):
        spPr.remove(xfrm)
        removed = True
    return removed


_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
