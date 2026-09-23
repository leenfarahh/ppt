"""Keep a photograph's frame and its cropped-to shape across a layout swap.

The bug, from a real deck: a portrait cropped to a circle came out square and
hard against the left margin, while the decorative arcs drawn around it stayed
exactly where they were. Nothing had gone wrong with the copy. The circle and
the position were never on that slide. They were on the OLD layout's picture
placeholder, and the slide's `p:pic` carried only the image, the crop and
`p:ph idx="N"`. Point that at another layout and both are simply gone:
geometry falls back to the implicit `rect`, and the frame to whatever the new
layout puts at that idx -- the origin, when it puts nothing there. The arcs
came through untouched because they are ordinary autoshapes carrying their own
`a:xfrm`, which is the whole difference.

So the inherited values are written onto the shape while the old layout is
still reachable, and the `p:ph` is then removed. Removing it matters as much as
baking the values in: left in place it is a standing invitation to inherit
again, from a layout that knows nothing about this shape.

This lives in a module of its own because both rebuild routes need it and
neither can reach the other: `builder` imports `master_apply`, so
`master_apply` cannot import `builder`.

WHY IT IS DONE ON THE FILE FOR THE POWERPOINT ROUTE, and not through COM.
Automation can read the designed geometry -- ask about the layout placeholder
and COM answers `msoShapeOval` where it answers `msoShapeNotPrimitive` for the
slide's own shape -- and it can write a preset back through `AutoShapeType`.
What it cannot do is carry a `a:custGeom` mask, which COM has no vocabulary
for: it reports `msoShapeNotPrimitive` for a drawn shape and for an inherited
one alike, so the two are indistinguishable and a custom mask is silently
lost. Baking the geometry into the file first is the only form of this fix
that does not depend on how the circle happens to be authored.
"""

from __future__ import annotations

import copy
import hashlib
import logging
from pathlib import Path
from typing import Any, Optional

from lxml import etree

log = logging.getLogger(__name__)

_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
# Every attribute in this namespace names a relationship, whatever the tag and
# whatever the attribute is called. See `_relationships`.
_R_NAMESPACE = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_R_LINK = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/"
    "relationships}link"
)
_R_EMBED = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/"
    "relationships}embed"
)
_IMAGE_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
)


def freeze_file(path: Path, roles: Optional[Any] = None) -> int:
    """Freeze every slide in the file at `path`, in place. Best effort.

    Returns the number of shapes changed, and 0 both when there was nothing to
    do and when the file could not be handled -- the caller carries on either
    way, because the fallbacks that remain are the old behaviour rather than a
    broken file.

    The caller must hand this a disposable copy. It rewrites what it is given.

    Saving nothing when nothing was frozen is deliberate, not an optimisation.
    A deck with no picture placeholders reaches PowerPoint byte for byte as
    the designer saved it, which is the property `master_apply` is built
    around and worth keeping for every deck that does not need this.
    """
    from pptx import Presentation  # noqa: PLC0415 - lazy heavy dependency

    try:
        presentation = Presentation(str(path))
        frozen = [name for slide in presentation.slides for name in freeze_slide(slide)]
        # Same pass and same save. Both exist for one reason -- what a slide
        # inherits from the layout it is about to stop pointing at -- and a
        # deck needing neither still reaches PowerPoint byte for byte.
        carried = inherit_artwork(presentation)
        pinned = pin_by_role(presentation, roles)
        if not frozen and not carried and not pinned:
            return 0
        presentation.save(str(path))
    except Exception:
        log.warning(
            "could not freeze the picture geometry in %s; a photograph cropped "
            "to a shape by its layout may come out square",
            path.name, exc_info=True,
        )
        return 0
    if frozen:
        log.info("froze the inherited frame of %d picture(s): %s", len(frozen),
                 ", ".join(frozen))
    if carried:
        log.info(
            "carried %d picture(s) off the old layouts onto the slides that "
            "inherit them, so the rebuild does not leave them behind: %s",
            len(carried), ", ".join(carried),
        )
    if pinned:
        log.info(
            "pinned %d placeholder(s) where they were drawn, so applying the "
            "master does not move them into a region meant for something "
            "else: %s", len(pinned), ", ".join(pinned),
        )
    return len(frozen) + len(carried) + len(pinned)


# --------------------------------------------------------------------------- #
# Placeholders that must not be re-placed
# --------------------------------------------------------------------------- #

def pin_by_role(presentation: Any, roles: Optional[Any] = None) -> list[str]:
    """Un-placeholder the boxes the new layout has no right to move.

    THE DEFECT. Assigning a `CustomLayout` runs PowerPoint's own placeholder
    matching, and that matching reads a `p:ph` and nothing else: it knows a
    shape is "body placeholder 3" and has no idea whether the words in it are
    a paragraph or the source line under a chart. On a real deck it took
    "Source: Oxford Economics, Team analysis" out of the foot of the page and
    printed it under the title, on top of the subtitle, because the old
    master had tagged the footnote as a placeholder and the new one puts its
    third slot at the top.

    A shape with no `p:ph` is not a placeholder, so nothing matches it and it
    stays exactly where the designer drew it -- which is the whole correction.
    The frame and the look are baked in first, on the same terms and for the
    same reason as `freeze_slide`: while the `p:ph` is there the shape may be
    inheriting its position and its fill from a layout it is about to stop
    pointing at, and stripping the tag without baking those would drop it at
    the origin with no fill.

    WHICH SHAPES. The model's reading first -- a source, a chart's parts, a
    piece of drawn furniture -- because it looked at the picture and that is
    the only place the answer is written. The file's own reading of small
    print behind it, so a run with no model still holds back the case a
    designer notices first. Never a title and never body copy: moving those
    into the master's regions is what applying a master IS.

    Returns a line per shape pinned. Mutates the presentation.
    """
    pinned: list[str] = []
    canvas = (
        (presentation.slide_width or 0) / 914400,
        (presentation.slide_height or 0) / 914400,
    )
    from .charts import unpin  # noqa: PLC0415 - sibling, kept off import time

    for index, slide in enumerate(presentation.slides, start=1):
        protected = _protected_texts(roles, index)
        for shape in list(slide.shapes):
            if not _is_placeholder(shape):
                continue
            # A chart keeps its own frame, whatever slot it was dropped into.
            # See `charts.unpin`.
            if unpin(shape):
                pinned.append(f"slide {index}: {_name_of(shape)}")
                continue
            if not _text_of(shape).strip():
                # An empty placeholder is a slot, and a slot is the new
                # layout's business. Pinning one would leave the deck showing
                # the OLD master's empty prompt box forever.
                continue
            if not _pin_worthy(shape, protected, canvas):
                continue
            if _freeze(shape, _layout_placeholder(shape, slide)):
                pinned.append(f"slide {index}: {_name_of(shape)}")
    return pinned


def _protected_texts(roles: Optional[Any], number: int) -> set[str]:
    """The copy on this slide the model said must not be re-placed."""
    if roles is None:
        return set()
    try:
        here = roles.for_slide(number)
    except Exception:
        return set()
    return set(here.protected_texts) if here.reviewed else set()


def _pin_worthy(
    shape: Any, protected: set[str], canvas: tuple[float, float]
) -> bool:
    if protected:
        from ..ai.roles import normalize_copy  # noqa: PLC0415 - optional dep

        if normalize_copy(_text_of(shape)) in protected:
            return True
    return reads_as_a_note(shape, canvas)


def _name_of(shape: Any) -> str:
    try:
        return str(shape.name)
    except Exception:
        return "unnamed shape"


def _is_placeholder(shape: Any) -> bool:
    try:
        return bool(shape.is_placeholder)
    except Exception:
        return False



# Words that open a footnote, in the two languages this tool's decks are
# written in. THE FALLBACK, NOT THE RULE. `ai.roles` reads the picture and
# answers for every kind of small print, including the methodology line that
# opens with no marker at all; this is what holds back the commonest half of
# it on a run with no model, and it is deliberately a short list of things
# that are never anything else.
_NOTE_OPENERS = (
    "source:", "sources:", "note:", "notes:", "footnote:", "n.b.",
    "disclaimer:", "methodology:",
    "المصدر",      # al-masdar, "source"
    "ملاحظة",      # mulahaza, "note"
)

# The strip at the foot of the page where small print lives, as a share of the
# canvas, and the size above which type is not small print however low it
# sits. A 9pt line in the bottom eighth is a footnote; a 20pt line there is a
# closing statement somebody set large on purpose.
_FOOT_BAND = 0.86
_SMALL_PT = 13.0


def reads_as_a_note(shape: Any, canvas: tuple[float, float]) -> bool:
    """Whether this box is small print, read off the file alone.

    THE DETERMINISTIC HALF OF WHAT `ai.roles` DOES BETTER. It exists because
    the model is optional -- no key, no renderer, a spent quota -- and the
    defect it prevents is one a designer sees at a glance: a source line
    filled into the subtitle region and printed under the title.

    Two ways to qualify, and both are deliberately narrow, because a false
    positive here costs a region left empty and a false negative costs the
    footnote at the top of the page.

      - it opens with a word that opens nothing else. "Source:" is not a
        sentence anybody starts a body paragraph with.
      - it is set small AND sits in the strip at the foot of the page. Either
        on its own is ordinary: a caption is small and a closing line is low.

    A shape that will not answer is not a note, which keeps this from
    swallowing content on a file it cannot read.
    """
    text = _text_of(shape).strip()
    if not text:
        return False
    if text.casefold().lstrip("*†‡ ").startswith(_NOTE_OPENERS):
        return True

    _width_in, height_in = canvas
    if height_in <= 0:
        return False
    try:
        top = float(shape.top or 0) / 914400
    except Exception:
        return False
    if top / height_in < _FOOT_BAND:
        return False
    size = _largest_size(shape)
    return size is not None and size <= _SMALL_PT


def _largest_size(shape: Any) -> Optional[float]:
    """The biggest type on the shape, in points, or None if it states none.

    The biggest rather than the smallest: a footnote with a superscript marker
    in it states two sizes, and the marker is not what the box is set at.
    """
    found: list[float] = []
    try:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size is not None:
                    found.append(run.font.size.pt)
    except Exception:
        return None
    return max(found) if found else None


def freeze_slide(src_slide: Any) -> list[str]:
    """Make a picture placeholder's inherited look explicit, then un-placeholder it.

    The bug, from a real deck: a portrait cropped to a circle came out square
    and hard against the left margin, while the decorative arcs drawn around
    it stayed exactly where they were. Nothing had gone wrong with the copy.
    The circle and the position were never on that slide. They were on the OLD
    layout's picture placeholder, and the slide's `p:pic` carried only the
    image, the crop and `p:ph idx="N"`. Point that at another layout and both
    are simply gone: geometry falls back to the implicit `rect`, and the frame
    to whatever the new layout puts at that idx -- the origin, when it puts
    nothing there. The arcs came through untouched because they are ordinary
    autoshapes carrying their own `a:xfrm`, which is the whole difference.

    So the inherited values are written onto the shape while the old layout is
    still reachable, and the `p:ph` is then removed. Removing it matters as
    much as baking the values in: left in place it is a standing invitation to
    inherit again, from a layout that knows nothing about this shape.

    Only placeholders holding an image are touched, never one holding words. A
    text placeholder taking the new layout's type and position is the entire
    point of the rebuild, and freezing one would reproduce exactly the drift
    this exists to remove.

    Returns the names of the shapes it changed. Mutates `src_slide`, which is
    a throwaway in-memory read of the source file and is never written back.
    """
    from .charts import unpin  # noqa: PLC0415 - sibling, kept off import time

    frozen: list[str] = []
    for shape in list(src_slide.shapes):
        if _is_placeholder(shape) and unpin(shape):
            # A chart in a content placeholder: its own `p:xfrm` is its frame,
            # and it crosses as one element. See `rebuild.charts`.
            frozen.append(_name_of(shape))
            continue
        if not _is_placeholder(shape) or not _carries_picture(shape):
            continue
        if _freeze(shape, _layout_placeholder(shape, src_slide)):
            frozen.append(_name_of(shape))
    return frozen


def _carries_picture(shape: Any) -> bool:
    """A shape whose content is an image rather than words.

    Both ways of authoring one count: `p:pic`, and the autoshape with a
    picture fill that PowerPoint writes when a photo is dropped onto a shape.

    Any text disqualifies it. A titled panel that happens to have a photo fill
    is a text placeholder, and its words moving into the new layout's type is
    what has to happen to it.
    """
    element = getattr(shape, "_element", None)
    if element is None:
        return False
    if _text_of(shape).strip():
        return False
    return element.find(f".//{_A_NS}blip") is not None


def _text_of(shape: Any) -> str:
    try:
        return str(shape.text_frame.text) if shape.has_text_frame else ""
    except Exception:
        return ""


def _freeze(shape: Any, layout_ph: Optional[Any]) -> bool:
    """Bake the frame, then the look, then the geometry, then drop the `p:ph`."""
    spPr = getattr(shape._element, "spPr", None)
    if spPr is None:
        return False
    _bake_frame(spPr, _frame_of(shape))
    if layout_ph is not None:
        _bake_look(spPr, getattr(layout_ph._element, "spPr", None))
    _ensure_geometry(spPr)
    return _strip_ph(shape._element)


def _ensure_geometry(spPr: Any) -> None:
    """Leave every frozen picture with a geometry of its own. Always.

    THE PHOTOGRAPH THAT RENDERED AS WHITE. On a real deck, a full-bleed image
    down the right half of a section divider came out of the rebuild as a white
    panel, and every check said the file was fine: the `p:pic` was there, on
    top, with its `a:xfrm` baked in and its `r:embed` resolving to a 2.4MB
    JPEG that was present in the package. PowerPoint opened it without
    complaint and listed the shape as visible at the right size. Then it
    exported the slide with nothing in that half.

    What the shape had lost was its GEOMETRY. A picture placeholder states no
    `a:prstGeom`, because it takes one from the layout placeholder behind it --
    and where the layout states none either, both are relying on the implicit
    rectangle the schema gives a shape that is still a placeholder. Strip the
    `p:ph`, as this module must, and that inheritance is over: the shape is now
    an ordinary picture with an explicit frame and no geometry at all, and a
    shape with no geometry has nothing to fill.

    `_bake_look` is not the answer to this, though it looks like it should be.
    It copies what the layout SUPPLIES, which is right for a circle or a drawn
    mask and useless here, where the layout supplies nothing and there is no
    layout placeholder at all in the case a template has renamed its indices.

    So the implicit rectangle is written out, and only when the shape has said
    nothing of its own. It is not a new decision about how the picture looks --
    it is the value the file already meant, made explicit before the thing that
    was carrying it is taken away.
    """
    if spPr is None:
        return
    if spPr.find(f"{_A_NS}prstGeom") is not None:
        return
    if spPr.find(f"{_A_NS}custGeom") is not None:
        return
    geometry = etree.SubElement(spPr, f"{_A_NS}prstGeom")
    geometry.set("prst", "rect")
    etree.SubElement(geometry, f"{_A_NS}avLst")
    # Order matters to the schema: a:xfrm comes first, then the geometry, and
    # everything else after. `SubElement` appends, so a shape that already
    # carries a fill or an outline would have the geometry land behind them.
    _put_geometry_after_xfrm(spPr, geometry)


def _put_geometry_after_xfrm(spPr: Any, geometry: Any) -> None:
    """`a:xfrm`, then geometry, then the rest -- which is what `a:CT_ShapeProperties` says."""
    xfrm = spPr.find(f"{_A_NS}xfrm")
    index = list(spPr).index(xfrm) + 1 if xfrm is not None else 0
    spPr.remove(geometry)
    spPr.insert(index, geometry)


def _frame_of(shape: Any) -> Optional[tuple[int, int, int, int]]:
    """Where the shape actually sits, resolving what it inherits.

    python-pptx reports the effective value for a placeholder -- its own where
    it has one, otherwise the layout's, otherwise the master's -- and that
    chain is the whole reason this works, because a picture placeholder in a
    designed template carries none of the four itself.
    """
    try:
        frame = (shape.left, shape.top, shape.width, shape.height)
    except Exception:
        # `LayoutPlaceholder._base_placeholder` looks its type up in a table
        # and raises KeyError on one the table does not list, which a
        # hand-built template can carry.
        return None
    if any(value is None for value in frame):
        return None
    return (int(frame[0]), int(frame[1]), int(frame[2]), int(frame[3]))


def _bake_frame(spPr: Any, frame: Optional[tuple[int, int, int, int]]) -> None:
    """Write an explicit `a:xfrm`, unless the shape already had one."""
    if frame is None or spPr.xfrm is not None:
        return
    xfrm = spPr.get_or_add_xfrm()
    xfrm.x, xfrm.y, xfrm.cx, xfrm.cy = frame


def _bake_look(spPr: Any, source: Optional[Any]) -> None:
    """Take the geometry, outline and shadow the layout was supplying.

    Geometry is one or the other: `a:custGeom` and `a:prstGeom` are a schema
    choice, so copying both writes a file PowerPoint refuses. Custom wins
    where a layout has one, because a shape that needed a drawn outline was
    never going to be describable as a preset.

    Anything the slide states for itself is left alone throughout. It was
    overriding the layout before the move and it still should after.
    """
    if source is None:
        return
    if spPr.prstGeom is None and spPr.custGeom is None:
        for name in ("custGeom", "prstGeom"):
            geometry = source.find(f"{_A_NS}{name}")
            if geometry is not None:
                getattr(spPr, f"_insert_{name}")(copy.deepcopy(geometry))
                break
    for name in ("ln", "effectLst"):
        if spPr.find(f"{_A_NS}{name}") is not None:
            continue
        inherited = source.find(f"{_A_NS}{name}")
        if inherited is not None:
            getattr(spPr, f"_insert_{name}")(copy.deepcopy(inherited))


def _layout_placeholder(shape: Any, src_slide: Any) -> Optional[Any]:
    """The source layout's placeholder this shape has been inheriting from."""
    try:
        return src_slide.slide_layout.placeholders.get(idx=shape._element.ph_idx)
    except Exception:
        return None


def _strip_ph(element: Any) -> bool:
    """Turn the shape into an ordinary floating one.

    Not cosmetic. While the `p:ph` is there the shape is a placeholder on a
    slide whose layout has changed, so PowerPoint resolves anything it does
    not state against the wrong layout, and `_claim` offers it a target that
    an image cannot be copied into -- which is its own silent loss, because
    `_copy_text` finds no text frame on either side and returns having done
    nothing, while the shape is reported as filled.
    """
    ph = element.ph
    if ph is None:
        return False
    parent = ph.getparent()
    if parent is None:
        return False
    parent.remove(ph)
    return True


# --------------------------------------------------------------------------- #
# Artwork the slide inherits from its layout
# --------------------------------------------------------------------------- #
#
# What a designed slide shows is often not on the slide. It is on the layout,
# and the slide inherits it: a section divider carries a full-bleed photograph,
# a pair of connector lines and the grey panels behind its copy while its own
# `p:cSld` holds nothing but a title. Point that slide at another layout and
# all of it is gone -- not dropped, never there, so nothing reports it and no
# shape was lost.
#
# Measured on a deck built for it:
#
#     before   layout carries  Rounded Rectangle, Connector, Oval, Picture
#     after    slide shows     nothing of them
#
# So the artwork is copied onto the slide while the old layout is still
# reachable, which is the same move `freeze_slide` makes for a placeholder's
# frame and for the same reason.
#
# EVERY non-placeholder shape, not only pictures. The first version of this
# carried pictures alone and a real deck came back with its photographs intact
# and its connector lines and panel fills missing, which is the same defect
# wearing a different tag name. A placeholder is excluded because it is a slot
# rather than artwork, and taking the new layout's position and type for those
# is the entire point of a rebuild.
#
# WHAT IS NOT CARRIED, and why there has to be a rule at all. The old layouts
# hold two kinds of shape and only one should travel. A section photograph and
# the panels drawn for that slide belong to the slide; the old brand's logo and
# its header band belong to the design being replaced, and carrying those would
# stamp the previous client's mark on every slide of a rebranded deck -- worse
# than the missing artwork, and harder to spot.
#
# They are told apart by repetition rather than by size, which is the signal
# the file actually carries: furniture is on many layouts because it is on
# every slide, and a section's own artwork is on the one layout drawn for it.

def inherit_artwork(presentation: Any) -> list[str]:
    """Copy each slide's inherited layout artwork onto the slide itself.

    Returns a line per shape copied. Mutates the presentation.
    """
    furniture = _repeated_shapes(presentation)
    carried: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        if _is_cover_layout(slide.slide_layout):
            # A cover's layout artwork IS the old brand's cover design. It is
            # on the one layout the one cover uses, so the repetition test
            # below reads it as that slide's own content -- and on a real deck
            # the old cover's freeforms, accent rule and logo were stamped onto
            # the slide, on top of the new master's cover. The master being
            # applied always brings a cover of its own.
            continue
        artwork = _artwork_of(slide.slide_layout, furniture)
        # Kept in the layout's own order, and inserted at the front of the
        # z-order where a layout draws: under the slide's own shapes, which is
        # where these were. A panel arriving on top of the copy it used to sit
        # behind is a different defect from the one being fixed.
        for offset, shape in enumerate(artwork):
            if _copy_onto(slide, shape, _Z_FRONT + offset):
                carried.append(f"slide {index}: {_name_of(shape)}")
    return carried


def _is_cover_layout(layout: Any) -> bool:
    """Whether a layout is a title slide: PowerPoint's own `title` layout type,
    or a layout with a centre title, which is what a cover is built on."""
    element = getattr(layout, "_element", None)
    if element is not None and str(element.get("type") or "") == "title":
        return True
    for placeholder in _safe(lambda: list(layout.placeholders)) or []:
        kind = str(_safe(lambda p=placeholder: p.placeholder_format.type) or "")
        if kind.startswith("CENTER_TITLE"):
            return True
    return False


# `p:spTree` opens with `p:nvGrpSpPr` and `p:grpSpPr`; shapes follow.
_Z_FRONT = 2


def _repeated_shapes(presentation: Any) -> set[str]:
    """The old brand's furniture, which does not travel.

    Counted over SLIDES, not over layouts, and that is the whole correction. A
    layout count says furniture is what several layouts share -- true of a
    deck with many layouts, and useless for the common one where every slide
    sits on the same layout. On a real 12-slide deck that meant its
    background picture, its logo, its slide-number box and two footer text
    boxes were each on exactly one layout, counted as content, and carried
    onto all twelve slides. The master applied underneath them and the output
    looked like the deck that came in.

    What a slide inherits is what a slide inherits: chrome reaches nearly
    every slide because that is what chrome is for, and a section's own
    artwork reaches the handful of slides in that section. So a shape most of
    the deck inherits is the brand's, and one a few slides inherit is theirs.

    The layout count is kept as well. A logo on nine layouts is furniture even
    when the slides using them are spread unevenly, and that reading costs
    nothing to keep. So is anything on a slide master, which is drawn once for
    the whole deck and is therefore the brand's by construction.

    The limit worth knowing: chrome sitting on a layout only one slide uses
    reaches one slide, and nothing here can tell that from that slide's own
    artwork. It errs toward carrying, which is the safer way round -- one
    stray shape on one slide is a smaller wrong than a section image
    stripped out of every slide that needed it.
    """
    by_layout: dict[str, set[int]] = {}
    on_master: set[str] = set()
    for master in _safe(lambda: list(presentation.slide_masters)) or []:
        # Whatever sits on a slide master is the brand's by construction: it
        # is drawn once for the whole deck. No counting needed.
        for shape in _loose_shapes(master):
            key = _content_key(shape)
            if key:
                on_master.add(key)
        for layout in _safe(lambda: list(master.slide_layouts)) or []:
            for shape in _loose_shapes(layout):
                key = _content_key(shape)
                if key:
                    by_layout.setdefault(key, set()).add(id(layout))

    slides = _safe(lambda: list(presentation.slides)) or []
    reach: dict[str, int] = {}
    for slide in slides:
        layout = _safe(lambda s=slide: s.slide_layout)
        if layout is None:
            continue
        for shape in _loose_shapes(layout):
            key = _content_key(shape)
            if key:
                reach[key] = reach.get(key, 0) + 1

    ceiling = max(1, int(len(slides) * _CHROME_SHARE))
    return {
        key
        for key, layouts in by_layout.items()
        if key in on_master
        or len(layouts) > 1
        or reach.get(key, 0) > ceiling
    } | on_master


# The share of a deck a shape has to reach before it is the brand's furniture
# rather than a section's artwork. Above half: a divider image belongs to the
# handful of slides in its section, and a logo belongs to all of them.
_CHROME_SHARE = 0.5


def _artwork_of(layout: Any, furniture: set[str]) -> list[Any]:
    """The shapes on a layout that belong to the slide rather than to the brand."""
    return [
        shape for shape in _loose_shapes(layout)
        if _content_key(shape) not in furniture
    ]


# A chart, a SmartArt diagram, an embedded OLE object: everything PowerPoint
# writes as `p:graphicFrame`. What it draws is not in the shape, it is in parts
# the shape points at, and those parts have their own internal structure.
_GRAPHIC_FRAME = (
    "{http://schemas.openxmlformats.org/presentationml/2006/main}graphicFrame"
)


def _loose_shapes(container: Any) -> list[Any]:
    """Non-placeholder shapes directly on a layout, that can actually be copied.

    Placeholders are excluded on purpose: one is a slot, not artwork, and the
    slide's own shape is what fills it.

    AND SO IS EVERY `p:graphicFrame`, WHICH IS NOT A REFINEMENT. A real deck's
    layouts carried embedded OLE objects -- `Objekt 4`, `Objekt 8`, `Objekt 11`
    in a template authored in German -- and copying one onto a slide produced a
    file PowerPoint would not open at all: `0x80070570`, "the file or directory
    is corrupted and unreadable". Not a shape that looked wrong. A file that
    did not open.

    The damage went a long way from here. `master_apply` freezes the working
    copy before handing it to PowerPoint, so the corruption happened first and
    the apply failed; the rebuild fell back to recreating the file, and that
    route drops the charts it cannot rebuild and reports them. A designer saw a
    deck come back with its graphs missing and its titles reflowed, and none of
    that was anywhere near the code that caused it.

    Re-pointing every relationship in the subtree -- which `_copy_onto` now
    does, and is right for its own reasons -- is not enough for these. What a
    graphic frame draws lives in parts with their own internal structure and
    their own expectations about the shape that hosts them, and duplicating one
    by copying its XML is not something this module can do correctly. The rest
    of the tool already says so about the same class of shape: the rebuild
    leaves charts, SmartArt and embedded objects behind and names them, because
    a silently broken chart is worse than a missing one. A file that will not
    open is worse than both.
    """
    try:
        shapes = [s for s in container.shapes if not _is_placeholder(s)]
    except Exception:
        return []

    keep, skipped = [], []
    for shape in shapes:
        if getattr(shape, "_element", None) is not None \
                and shape._element.tag == _GRAPHIC_FRAME:
            skipped.append(_name_of(shape))
            continue
        keep.append(shape)
    if skipped:
        # Debug, not info: this is asked once per slide per layout, so a deck
        # whose every layout carries one says the same thing fifty times.
        log.debug(
            "not carrying %s off the layout: a chart, diagram or embedded "
            "object cannot be copied onto a slide without breaking the file",
            ", ".join(skipped),
        )
    return keep


def _content_key(shape: Any) -> Optional[str]:
    """What a shape draws, identified by content so copies collapse together.

    A picture is its image. Anything else is its markup with the identifiers
    stripped out -- the shape id, the name and the creation GUID all differ
    between two copies of one band and say nothing about what it looks like.
    """
    try:
        sha1 = _safe(lambda: shape.image.sha1)
        if sha1:
            return f"image:{sha1}"
    except Exception:
        pass
    try:
        element = copy.deepcopy(shape._element)
        for node in element.iter():
            for attribute in ("id", "name"):
                if attribute in node.attrib:
                    del node.attrib[attribute]
        for ext in element.iter(f"{_A_NS}extLst"):
            ext.getparent().remove(ext)
        return "xml:" + hashlib.sha1(
            etree.tostring(element)
        ).hexdigest()
    except Exception:
        return None


def _copy_onto(slide: Any, shape: Any, position: int) -> bool:
    """Put a copy of a layout's shape on the slide at `position` in the z-order.

    Any part it points at is related to the slide rather than copied, so a
    photograph on nine slides is still one image in the file.

    EVERY relationship in the copied subtree, not just the one on `a:blip`.
    An icon from PowerPoint's library carries two: the `a:blip` and, inside
    its extension list, an `asvg:svgBlip` pointing at the SVG it actually
    draws from. Re-pointing only the first left the second holding a
    relationship id that means something else on the slide, or nothing at all,
    and the icon arrived as a broken-image box. The tag is not worth matching
    on -- what matters is that the attribute is there.
    """
    try:
        element = copy.deepcopy(shape._element)
        for node in element.iter():
            for attribute in _relationships(node):
                _repoint(node, attribute, shape.part, slide.part)
        slide.shapes._spTree.insert(position, element)
        return True
    except Exception:
        log.debug("could not carry a layout shape onto a slide", exc_info=True)
        return False


def _relationships(node: Any) -> list[str]:
    """Every attribute on this node that names a relationship.

    EVERY ONE, and the list is not worth enumerating: an attribute in the
    relationships namespace IS a relationship reference, whatever it is called
    and whatever tag carries it.

    Naming them was the bug, and it took a deck down. The two that were named
    were `r:embed` and `r:link`, which is a picture's whole vocabulary and not
    much else's. A layout carrying an embedded OLE object states three more:

        <p14:tags  r:id="rId1"/>          the shape's tag store
        <p:oleObj  r:id="rId3"/>          the embedded object itself
        <a:blip    r:embed="rId4"/>       the picture it falls back to

    Only the last was re-pointed. The other two were copied onto the slide
    holding ids that mean something else there, or nothing, and a dangling
    relationship is not a shape that looks wrong -- it is a file PowerPoint
    refuses to open at all: `0x80070570`, "the file or directory is corrupted
    and unreadable". Which meant the apply failed, the rebuild fell back to
    recreating the file, and that route drops the charts it cannot rebuild. A
    deck arrived with its graphs missing, and the cause was three characters in
    a tuple two modules away.
    """
    return [name for name in node.attrib if name.startswith(_R_NAMESPACE)]


def _repoint(node: Any, attribute: str, source: Any, target: Any) -> None:
    """Point one relationship at the same thing, from the slide's part.

    Raises if it cannot, and the caller then carries nothing. That is the point:
    a shape whose relationships cannot all be brought across must not be
    inserted, because the file it lands in stops opening.
    """
    rid = node.get(attribute)
    if not rid:
        return
    rel = source.rels[rid]
    if rel.is_external:
        # An image or object linked rather than embedded. The target is a URL,
        # not a part, so it is related as one -- resolving it as a part raises.
        node.set(
            attribute,
            target.relate_to(rel.target_ref, rel.reltype, is_external=True),
        )
        return
    node.set(attribute, target.relate_to(rel.target_part, rel.reltype))


def _rel_type_of(part: Any, rid: str) -> str:
    """The relationship type the source used, so the copy declares the same.

    An SVG and its raster fallback are both images, but reading the type back
    rather than assuming it means a relationship this does not know about is
    carried across intact instead of being relabelled.
    """
    try:
        return str(part.rels[rid].reltype)
    except Exception:
        return _IMAGE_REL


def _safe(call):
    try:
        return call()
    except Exception:
        return None
