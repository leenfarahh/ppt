"""Corrections the AI layer proposes, and the checks that make them safe.

An AI finding carried prose and nothing else, so none of them could ever be
applied: "the caption should sit closer to the chart" has no target to parse.
`Issue.fix` is where the model names an action from a closed set and a target
a machine can act on.

The whole safety of it is that a proposal is not an instruction. Every value
is checked against the master before the file is touched, so the worst a bad
proposal can do is cost itself -- and that check is what most of these are
about, because a colour the model liked the look of, written into a client
deck wearing a brand label, is the failure that matters.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.apply import FixBrand, apply_fixes, fixer_for
from formatting_tool.apply.fixers import fix_ai_action
from formatting_tool.models import (
    FIX_OPS,
    Category,
    FixAction,
    Issue,
    Severity,
    Source,
)

NAVY = "1D252D"
NOT_BRAND = "FF00FF"


def _brand() -> FixBrand:
    return FixBrand(
        palette={"brand navy": NAVY, "theme:lt1": "FFFFFF"},
        allowed_fonts=["Aptos Display", "Aptos"],
        roles={"title": (32.0, 40.0), "subtitle": (18.0, 22.0)},
    )


def _ai(op: str | None = None, **fix) -> Issue:
    issue = Issue(
        category=Category.TYPOGRAPHY,
        severity=Severity.WARNING,
        message=fix.pop("message", "the model's wording"),
        source=Source.AI,
        slide=1,
        shape="Body 1",
        deck="messy.pptx",
        confidence=0.8,
        fix=FixAction(op=op, shape="Body 1", **fix) if op else None,
    )
    issue.id = issue.fingerprint()
    return issue


def _deck(tmp_path: Path, text: str = "Health Systems", blanks: int = 0) -> Path:
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.name = "Body 1"
    box.text_frame.text = text
    for _ in range(blanks):
        box.text_frame.add_paragraph()
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))
    return deck


# --------------------------------------------------------------------------- #
# Which findings get one at all
# --------------------------------------------------------------------------- #

def test_an_ai_finding_without_an_action_still_needs_a_designer() -> None:
    """Most of them, and that is the point. The AI layer exists for the
    judgement the rules cannot make, and a judgement has no op."""
    assert fixer_for(_ai()) is None


def test_an_ai_finding_with_an_action_becomes_fixable() -> None:
    assert fixer_for(_ai("set_font", font="Aptos")) is fix_ai_action


def test_an_op_nothing_implements_is_not_offered() -> None:
    """A fix that cannot run is worse than a finding saying it needs a
    designer, because it reads as an offer."""
    assert fixer_for(_ai("teleport")) is None
    assert fixer_for(_ai("delete_slide")) is None


def test_geometry_is_a_proposal_like_any_other() -> None:
    """Refused outright at first, on the grounds that the model is told not to
    measure off a rendered image. That was the wrong line: the payload carries
    every box in inches, the slide size and the safe margins, all exact, and
    reasoning from those is arithmetic on numbers it was given rather than an
    impression of a picture. What makes it safe is the checking, not the
    trusting."""
    from formatting_tool.models import GEOMETRIC_OPS

    assert GEOMETRIC_OPS <= FIX_OPS
    assert _ai("move", left_in=1.0, top_in=1.0).fix.geometric
    assert not _ai("recolor_text", hex=NAVY).fix.geometric


# --------------------------------------------------------------------------- #
# The checks
# --------------------------------------------------------------------------- #

def test_a_colour_the_master_does_not_declare_is_refused(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    issue = _ai("recolor_text", hex=NOT_BRAND)

    result = apply_fixes(
        deck, [issue], tmp_path / "out.pptx", selected=[issue.id],
        spec=_spec(),
    )

    assert not result.applied
    assert "not in the brand palette" in result.skipped[0].detail


def test_a_palette_entry_is_matched_exactly_not_nearly(tmp_path: Path) -> None:
    """"Nearest" is how a colour the model liked gets into a client deck
    wearing a brand label. One digit off is not the brand colour."""
    deck = _deck(tmp_path)
    almost = _ai("recolor_text", hex="1D252E")      # NAVY, last digit moved

    result = apply_fixes(
        deck, [almost], tmp_path / "out.pptx", selected=[almost.id], spec=_spec()
    )

    assert not result.applied


def test_an_unapproved_typeface_is_refused(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    issue = _ai("set_font", font="Comic Sans MS")

    result = apply_fixes(
        deck, [issue], tmp_path / "out.pptx", selected=[issue.id], spec=_spec()
    )

    assert not result.applied
    assert "approved typefaces" in result.skipped[0].detail


def test_nothing_is_applied_without_a_master_to_check_against(
    tmp_path: Path,
) -> None:
    """An unchecked proposal is a guess with a brand label on it."""
    deck = _deck(tmp_path)
    issue = _ai("set_font", font="Aptos")

    result = apply_fixes(deck, [issue], tmp_path / "out.pptx", selected=[issue.id])

    assert not result.applied
    assert "no brand reference" in result.skipped[0].detail


def test_a_size_outside_the_roles_range_is_refused(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    issue = _ai("set_font_size", size_pt=48.0)
    issue.shape = "Title 1"
    issue.fix.shape = "Title 1"
    issue.id = issue.fingerprint()

    from pptx import Presentation

    prs = Presentation(str(deck))
    prs.slides[0].shapes[0].name = "Title 1"
    prs.save(str(deck))

    result = apply_fixes(
        deck, [issue], tmp_path / "out.pptx", selected=[issue.id], spec=_spec()
    )

    assert not result.applied
    assert "outside the 32.0-40.0pt range" in result.skipped[0].detail


def test_a_role_the_brand_is_silent_about_does_not_veto_a_size() -> None:
    """Refusing every proposal for a role nobody wrote a rule about would be
    inventing a rule."""
    assert _brand().size_range("body") == (None, None)
    assert _brand().size_range("title") == (32.0, 40.0)


# --------------------------------------------------------------------------- #
# The actions that carry
# --------------------------------------------------------------------------- #

def test_an_approved_typeface_is_set(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    issue = _ai("set_font", font="Aptos")
    out = tmp_path / "out.pptx"

    result = apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    assert len(result.applied) == 1
    from pptx import Presentation

    run = Presentation(str(out)).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Aptos"


def test_trailing_empty_paragraphs_come_off(tmp_path: Path) -> None:
    """`typography.whitespace` works inside a paragraph and cannot see these:
    an empty paragraph carries no runs to strip. It is what stops a vertically
    centred box looking centred."""
    deck = _deck(tmp_path, blanks=3)
    issue = _ai("delete_empty_paragraphs")
    out = tmp_path / "out.pptx"

    result = apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    assert len(result.applied) == 1
    from pptx import Presentation

    frame = Presentation(str(out)).slides[0].shapes[0].text_frame
    assert len(frame.paragraphs) == 1
    assert frame.paragraphs[0].text == "Health Systems"


def test_the_last_paragraph_is_never_removed(tmp_path: Path) -> None:
    """An empty text frame is a different shape from a text frame with an
    empty paragraph in it, and removing the last one is not a formatting fix."""
    deck = _deck(tmp_path, text="", blanks=2)
    issue = _ai("delete_empty_paragraphs")
    out = tmp_path / "out.pptx"

    apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    from pptx import Presentation

    frame = Presentation(str(out)).slides[0].shapes[0].text_frame
    assert len(frame.paragraphs) >= 1


def _spec():
    """A stand-in for MasterSpec: FixBrand.from_spec reads it by attribute."""

    class _Role:
        def __init__(self, low, high):
            self.min_size_pt, self.max_size_pt = low, high

    class _Spec:
        palette = {"brand navy": NAVY, "theme:lt1": "FFFFFF"}
        allowed_fonts = ["Aptos Display", "Aptos"]
        roles = {"title": _Role(32.0, 40.0), "subtitle": _Role(18.0, 22.0)}

    return _Spec()
