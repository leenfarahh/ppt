"""A band across a column that does not reach the column's edges.

Off a real slide: a headline band reading "A trillion-dollar industry is
passing the inflection point" over three rows of content, 0.06in narrower than
the rows beneath it. Small enough to survive a review, and obvious once seen,
because the eye reads a band and the block under it as one object and a short
edge as a mistake.

Three rules looked like they should have caught it and each had a reason not
to, which is why this is a rule of its own rather than a widened tolerance
somewhere:

- `space.alignment_grid` reads the LEADING edge only, against a deck-wide
  grid. The band's left edge is correct; it is the trailing edge, and so the
  width, that is wrong.
- `space.repeat_out_of_line` and `space.row_out_of_line` group through
  `repeats._series`, which buckets on identical size and type. A band is not
  the size of the boxes it heads, so it is never grouped with them.
- The overlap and crowding rules are about shapes touching. These do not
  touch.

What makes it detectable is comparing the band against the BOUNDING BOX of
the cluster it caps rather than against any shape in it: the band spans a dark
label column and a light body column and matches neither alone.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.extract import read_deck
from formatting_tool.extract.master_spec import derive_master_spec
from formatting_tool.models import BrandGuidelines
from formatting_tool.rules import RuleContext
from formatting_tool.rules.space import BandWidthRule

# The column runs 0.15 to 6.25in in every fixture here, so a band's numbers
# can be read against those two without going back to the builder.
COLUMN_LEFT, COLUMN_RIGHT = 0.15, 6.25


def _slide(tmp_path: Path, *, band=(0.15, 6.04), rows: int = 3,
           band_top: float = 0.40, band_height: float = 0.45,
           extra=None) -> Path:
    """A band over `rows` rows of [dark label | light body].

    `band` is (left, width) so a fixture can put an edge anywhere without
    arithmetic at the call site.
    """
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def box(name, left, top, width, height, fill):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
            Inches(width), Inches(height),
        )
        shape.name = name
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.fill.background()
        shape.text_frame.text = name
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
        return shape

    box("Headline band", band[0], band_top, band[1], band_height, "BFB295")
    for row in range(rows):
        top = 1.00 + row * 0.95
        box(f"Label {row + 1}", COLUMN_LEFT, top, 1.55, 0.85, "3F3628")
        box(f"Body {row + 1}", 1.78, top, 4.47, 0.85, "EFEDE8")
    for name, left, top, width, height in extra or ():
        box(name, left, top, width, height, "FFFFFF")

    path = tmp_path / "band.pptx"
    prs.save(str(path))
    return path


def _found(deck: Path) -> list:
    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    ctx = RuleContext(deck=read_deck(deck), spec=spec)
    issues = list(BandWidthRule().check(ctx))
    for issue in issues:
        issue.id = issue.fingerprint()
    return issues


# --------------------------------------------------------------------------- #
# The finding
# --------------------------------------------------------------------------- #

def test_the_short_band_is_reported(tmp_path: Path) -> None:
    """The slide this came from: 0.06in short on the right."""
    found = _found(_slide(tmp_path))

    assert len(found) == 1
    assert found[0].shape == "Headline band"
    assert "0.06in short of the right edge" in found[0].message
    assert "6 shapes" in found[0].message      # the whole cluster, not one row


def test_it_carries_the_column_span_to_fix_to(tmp_path: Path) -> None:
    found = _found(_slide(tmp_path))

    # left, top, width x height -- the band's own top and height, the
    # column's left and width.
    assert "at 0.15, 0.40, 6.10 x 0.45in" in found[0].expected


def test_a_band_that_overhangs_is_reported_too(tmp_path: Path) -> None:
    """Too wide is the same defect as too narrow, and reads the same way."""
    found = _found(_slide(tmp_path, band=(COLUMN_LEFT, 6.18)))

    assert len(found) == 1
    assert "past the right edge" in found[0].message


def test_a_footer_band_is_read_the_same_way(tmp_path: Path) -> None:
    """A band under a table is the same object and the same defect."""
    found = _found(_slide(tmp_path, band=(COLUMN_LEFT, 6.04), band_top=3.90))

    assert len(found) == 1
    assert found[0].shape == "Headline band"


def test_a_flush_band_says_nothing(tmp_path: Path) -> None:
    found = _found(_slide(tmp_path, band=(COLUMN_LEFT, COLUMN_RIGHT - COLUMN_LEFT)))

    assert found == []


# --------------------------------------------------------------------------- #
# What it must not report
# --------------------------------------------------------------------------- #

def test_a_band_inset_at_both_ends_is_left_alone(tmp_path: Path) -> None:
    """The discriminator between drift and design.

    A band dragged into place snaps one edge and misses the other. A band
    inset from its column on purpose is inset at both ends, and straightening
    that would be the tool overruling a designer.
    """
    found = _found(_slide(tmp_path, band=(0.55, 5.30)))

    assert found == []


def test_a_band_far_out_is_a_decision_not_drift(tmp_path: Path) -> None:
    """Out by a hair is carelessness; out by a lot is a device."""
    found = _found(_slide(tmp_path, band=(COLUMN_LEFT, 5.40)))

    assert found == []


def test_a_hair_inside_tolerance_says_nothing(tmp_path: Path) -> None:
    """0.01in is EMU rounding, not a defect anybody can see."""
    found = _found(
        _slide(tmp_path, band=(COLUMN_LEFT, COLUMN_RIGHT - COLUMN_LEFT - 0.01))
    )

    assert found == []


def test_a_tall_shape_is_not_a_band(tmp_path: Path) -> None:
    """A band is wide for its height. Without that a label above a column is
    read as capping it."""
    found = _found(_slide(tmp_path, band=(COLUMN_LEFT, 1.10), band_height=0.90))

    assert found == []


def test_a_column_of_two_is_not_enough(tmp_path: Path) -> None:
    """Same minimum the repeat rules use: below it a set is a coincidence."""
    found = _found(_slide(tmp_path, rows=1))

    assert found == []


def test_a_band_over_nothing_says_nothing(tmp_path: Path) -> None:
    """The cluster is grown from the band outwards, row by row, so a band with
    a gap under it caps nothing rather than reaching down the slide."""
    found = _found(_slide(tmp_path, band_top=0.10, band_height=0.20))

    # The first row starts 0.70in below the band's bottom, past the reach.
    assert found == []


# --------------------------------------------------------------------------- #
# Applying it
# --------------------------------------------------------------------------- #

def _edges(path: Path):
    from pptx import Presentation

    boxes = {
        s.name: (s.left / 914400, (s.left + s.width) / 914400)
        for s in Presentation(str(path)).slides[0].shapes
    }
    band = boxes.pop("Headline band")
    return band, (
        min(v[0] for v in boxes.values()),
        max(v[1] for v in boxes.values()),
    )


def test_the_band_is_taken_to_the_column_edges(tmp_path: Path) -> None:
    from formatting_tool.apply import apply_fixes

    deck = _slide(tmp_path)
    found = _found(deck)
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, found, out, selected=[i.id for i in found])

    assert len(result.applied) == 1
    assert "widened 0.06in" in result.applied[0].detail
    band, column = _edges(out)
    assert band == pytest.approx(column, abs=0.001)


def test_the_top_and_height_are_left_alone(tmp_path: Path) -> None:
    """A band that is the wrong width is no evidence it is at the wrong
    height."""
    from pptx import Presentation

    from formatting_tool.apply import apply_fixes

    deck = _slide(tmp_path)
    was = [
        (s.top, s.height) for s in Presentation(str(deck)).slides[0].shapes
        if s.name == "Headline band"
    ][0]
    found = _found(deck)
    out = tmp_path / "fixed.pptx"

    apply_fixes(deck, found, out, selected=[i.id for i in found])

    now = [
        (s.top, s.height) for s in Presentation(str(out)).slides[0].shapes
        if s.name == "Headline band"
    ][0]
    assert now == was


def test_applying_it_twice_changes_nothing_the_second_time(
    tmp_path: Path,
) -> None:
    """Idempotent, which is what writing BOTH edges buys: the fix reports the
    edge that was out and sets the whole span, so the next run has nothing to
    say."""
    from formatting_tool.apply import apply_fixes

    deck = _slide(tmp_path)
    found = _found(deck)
    out = tmp_path / "fixed.pptx"
    apply_fixes(deck, found, out, selected=[i.id for i in found])

    assert _found(out) == []


def test_a_band_that_cannot_grow_without_covering_something_is_refused(
    tmp_path: Path,
) -> None:
    """The guard this fix needs and cannot borrow.

    The applier reverts a bad geometric move by restoring left and top, which
    cannot undo a width, so a resize would sail through a guard unable to put
    it back. This one asks first.
    """
    from pptx import Presentation

    from formatting_tool.apply import apply_fixes

    # A caption pinned in the gap the band would have to grow into.
    deck = _slide(
        tmp_path, extra=[("Sidebar 1", COLUMN_RIGHT - 0.04, 0.40, 0.30, 0.45)]
    )
    found = _found(deck)
    assert len(found) == 1
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, found, out, selected=[i.id for i in found])

    assert result.applied == []
    assert "Sidebar 1" in result.skipped[0].detail
    band = [
        s for s in Presentation(str(out)).slides[0].shapes
        if s.name == "Headline band"
    ][0]
    assert band.width / 914400 == pytest.approx(6.04, abs=0.001)


# --------------------------------------------------------------------------- #
# What a real deck taught it
# --------------------------------------------------------------------------- #
#
# Run over three real decks, the first version of this rule reported six
# findings and every one was wrong. Two lessons, and both are now conditions.


def test_a_lone_overhanging_element_does_not_define_the_span(
    tmp_path: Path,
) -> None:
    """The false positive off a real slide.

    A grey band read as 0.23in short on its left, and it was not: the
    circular badge icon beside it deliberately overhangs the component, and
    one overhanging shape was enough to move the cluster's bounding edge. A
    component's outer edge is a SHARED edge -- three labels starting at the
    same left -- so an edge only one shape reaches is not the component's.
    """
    deck = _slide(
        tmp_path,
        band=(COLUMN_LEFT, COLUMN_RIGHT - COLUMN_LEFT),
        # A badge overhanging 0.3in to the left of the column, as the slide
        # had it.
        extra=[("Badge 1", COLUMN_LEFT - 0.30, 1.05, 0.42, 0.42)],
    )

    assert _found(deck) == []


def test_the_drift_ceiling_separates_the_two_real_cases(tmp_path: Path) -> None:
    """The numbers are set from real decks, not picked.

    The genuine defect was 0.06in out on a 6.04in band. The false positive
    was 0.23in out on a 5.83in band. The ceiling has to sit between them with
    room either side, which is what an eighth of an inch and 2% gives.
    """
    from formatting_tool.rules.space import _band_ceiling

    genuine, spurious = 0.06, 0.23

    assert genuine <= _band_ceiling(6.04)
    assert spurious > _band_ceiling(5.83)


def test_a_title_is_not_a_band(tmp_path: Path) -> None:
    """Four titles, a subtitle and a footer placeholder were reported on real
    decks, all wide and short, all compared against most of the slide because
    the cluster under a title IS most of the slide.

    A band is filled and is not a placeholder: a coloured bar drawn behind
    copy, where a title is a text frame whose width comes from the layout.
    """
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt

    from formatting_tool.rules.space import _is_band

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    # Layout 5 is Title Only, so this is a real title placeholder.
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "A trillion-dollar industry"
    unfilled = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6),
                                        Inches(0.4))
    unfilled.name = "Loose caption"
    unfilled.text_frame.text = "wide, short and see-through"
    unfilled.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    path = tmp_path / "titled.pptx"
    prs.save(str(path))

    shapes = {s.name: s for s in read_deck(path).slides[0].shapes}
    assert not _is_band(shapes[slide.shapes.title.name])
    assert not _is_band(shapes["Loose caption"])
