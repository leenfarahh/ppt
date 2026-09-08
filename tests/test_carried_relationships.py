"""Relationships that travel with a shape carried off a layout.

The bug, from a real deck: an icon arrived on the rebuilt slide as a
broken-image box. An icon from PowerPoint's library carries TWO relationships
-- the `a:blip` and, inside its extension list, an `asvg:svgBlip` pointing at
the SVG it actually draws from. Re-pointing only the first left the second
holding a relationship id that means something else on the slide, or nothing
at all.

So every relationship in the copied subtree is re-pointed, whatever tag
carries it. The tag is not worth matching on; what matters is that the
attribute is there.
"""

from __future__ import annotations

import copy
from pathlib import Path

import pytest

R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
IMAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
SVG_BLIP = "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip"


def _svg_icon_on_a_layout(tmp_path: Path):
    """A deck whose layout carries an SVG icon, as a designed template does."""
    pytest.importorskip("pptx")
    from pptx import Presentation

    source = Path("strategy&mini.pptx")
    if not source.exists():                     # pragma: no cover - fixture guard
        pytest.skip("needs the sample deck that carries an SVG icon")

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == 6:           # GROUP: the icon is inside one
                yield from walk(shape.shapes)
            else:
                yield shape

    src = Presentation(str(source))
    icon = None
    for slide in src.slides:
        for shape in walk(slide.shapes):
            element = getattr(shape, "_element", None)
            if element is not None and element.find(f".//{SVG_BLIP}") is not None:
                icon = shape
                break
        if icon is not None:
            break
    if icon is None:                            # pragma: no cover - fixture guard
        pytest.skip("no SVG icon in the sample deck")

    prs = Presentation()
    layout = prs.slide_layouts[5]
    element = copy.deepcopy(icon._element)
    for node in element.iter():
        rid = node.get(R_EMBED)
        if rid:
            node.set(
                R_EMBED,
                layout.part.relate_to(icon.part.related_part(rid), IMAGE_REL),
            )
    layout.shapes._spTree.append(element)

    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Energy"
    path = tmp_path / "icon-on-layout.pptx"
    prs.save(str(path))
    return Presentation(str(path))


def test_an_svg_icons_second_relationship_travels_with_it(tmp_path: Path) -> None:
    """The one that was left behind. Without it the icon is a broken-image box
    on the rebuilt slide, and nothing reports it."""
    from formatting_tool.rebuild.pictures import inherit_artwork
    from formatting_tool.svgicon import svg_part

    prs = _svg_icon_on_a_layout(tmp_path)
    slide = prs.slides[-1]

    carried = inherit_artwork(prs)

    assert carried, "the icon should have been carried onto the slide"
    pictures = [s for s in slide.shapes if s.shape_type == 13]
    assert pictures, "no picture landed on the slide"
    # The SVG the icon draws from resolves from the SLIDE's relationships now,
    # not from the layout's.
    assert svg_part(pictures[0]) is not None


def test_the_carried_copy_survives_a_save(tmp_path: Path) -> None:
    """A dangling relationship is not visible until the file is written and
    read back, which is where the broken-image box came from."""
    from pptx import Presentation

    from formatting_tool.rebuild.pictures import inherit_artwork
    from formatting_tool.svgicon import colors_of

    prs = _svg_icon_on_a_layout(tmp_path)
    inherit_artwork(prs)
    out = tmp_path / "carried.pptx"
    prs.save(str(out))

    reopened = Presentation(str(out))
    pictures = [s for s in reopened.slides[-1].shapes if s.shape_type == 13]

    assert pictures
    assert colors_of(pictures[0]), "the icon's drawing could not be read back"
