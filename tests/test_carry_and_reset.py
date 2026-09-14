"""Carrying a layout's artwork onto a slide, and handing placeholders back.

Both of these were found on one deck, and the first is the worst defect this
tool has produced. A template authored in German carried embedded OLE objects
on its layouts -- `Objekt 4`, `Objekt 8`, `Objekt 11` -- and copying one onto a
slide wrote a file PowerPoint would not open at all: 0x80070570, "the file or
directory is corrupted and unreadable".

The damage travelled a long way from the cause. `master_apply` freezes its
working copy before handing it to PowerPoint, so the corruption happened first
and the apply failed. The rebuild fell back to recreating the file, and that
route drops the charts it cannot rebuild and transplants placeholders with
their old geometry. What a designer saw was a deck with its graphs missing and
its titles reflowed one word to a line, and none of it was anywhere near the
three characters that caused it.
"""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches

from formatting_tool.rebuild.builder import _layout_slots, _slot_of, reset_layouts
from formatting_tool.rebuild.pictures import _loose_shapes, _relationships

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


class _Element:
    def __init__(self, tag, attrib=None):
        self.tag = tag
        self.attrib = attrib or {}

    def get(self, name):
        return self.attrib.get(name)


class _Shape:
    def __init__(self, name, tag, placeholder=False):
        self.name = name
        self._element = _Element(tag)
        self.is_placeholder = placeholder


class _Container:
    def __init__(self, shapes):
        self.shapes = shapes


# --------------------------------------------------------------------------- #
# What may be carried off a layout
# --------------------------------------------------------------------------- #

def test_an_embedded_object_is_never_carried_onto_a_slide() -> None:
    """The whole cascade starts here. A graphic frame is a chart, a diagram or
    an embedded object: what it draws is not in the shape, it is in parts the
    shape points at, and copying its XML onto a slide does not bring those."""
    layout = _Container([
        _Shape("Objekt 8", P_NS + "graphicFrame"),
        _Shape("Picture 2", P_NS + "pic"),
        _Shape("Rectangle 7", P_NS + "sp"),
    ])

    carried = [shape.name for shape in _loose_shapes(layout)]

    assert carried == ["Picture 2", "Rectangle 7"]


def test_a_placeholder_is_not_artwork() -> None:
    """A placeholder is a slot, not artwork, and the slide's own shape fills
    it."""
    layout = _Container([
        _Shape("Title 1", P_NS + "sp", placeholder=True),
        _Shape("Rectangle 7", P_NS + "sp"),
    ])

    assert [s.name for s in _loose_shapes(layout)] == ["Rectangle 7"]


# --------------------------------------------------------------------------- #
# Every relationship, not the two that were named
# --------------------------------------------------------------------------- #

def test_every_relationship_attribute_is_found_not_just_the_picture_ones() -> None:
    """`r:embed` and `r:link` are a picture's whole vocabulary and not much
    else's. The OLE object that broke the deck states three more, and only the
    last of them was being re-pointed:

        <p14:tags r:id="rId1"/>     <p:oleObj r:id="rId3"/>     <a:blip r:embed="rId4"/>
    """
    assert _relationships(_Element("tags", {R_NS + "id": "rId1"})) == [R_NS + "id"]
    assert _relationships(_Element("blip", {R_NS + "embed": "rId4"})) == [R_NS + "embed"]
    assert _relationships(_Element("blip", {R_NS + "link": "rId5"})) == [R_NS + "link"]
    # And nothing that is not a relationship.
    assert _relationships(_Element("sp", {"id": "7", "name": "Objekt 8"})) == []


# --------------------------------------------------------------------------- #
# Handing placeholders back to their layout
# --------------------------------------------------------------------------- #

def _deck(tmp_path, move_title: bool):
    """A one-slide deck whose title may or may not state its own box."""
    presentation = Presentation()
    layout = presentation.slide_layouts[1]          # Title and Content
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "A title"
    if move_title:
        title = slide.shapes.title
        title.left, title.top = Inches(9.0), Inches(5.0)
        title.width, title.height = Inches(1.2), Inches(0.4)
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    return path


def _title_states_its_own_box(path) -> bool:
    slide = Presentation(str(path)).slides[0]
    spPr = slide.shapes.title._element.spPr
    return spPr is not None and spPr.find(A_NS + "xfrm") is not None


def test_a_placeholder_that_states_its_own_box_is_handed_back(tmp_path) -> None:
    """The defect this fixes: a title 1.2in wide on the layout it used to sit
    on stays 1.2in wide on a layout that gives titles the full width, and
    reflows to one word a line down the side of the page."""
    path = _deck(tmp_path, move_title=True)
    assert _title_states_its_own_box(path)

    reset = reset_layouts(path)

    assert reset == {1: 1}
    assert not _title_states_its_own_box(path)


def test_a_placeholder_already_inheriting_is_left_alone(tmp_path) -> None:
    """Which is why this is silent after the PowerPoint route: assigning a
    CustomLayout already hands every placeholder back, so there is nothing
    here to do and the file is not rewritten."""
    path = _deck(tmp_path, move_title=False)
    before = path.read_bytes()

    assert reset_layouts(path) == {}
    assert path.read_bytes() == before


class _Slot:
    def __init__(self, kind, idx):
        self.type = kind
        self.idx = idx


class _Placeholder:
    def __init__(self, kind, idx):
        self.placeholder_format = _Slot(kind, idx)


class _Layout:
    def __init__(self, slots):
        self.placeholders = [_Placeholder(*slot) for slot in slots]


def test_a_slot_is_identified_by_its_type_and_its_index() -> None:
    """Both, because PowerPoint matches on both: a title has no meaningful idx
    and is found by type, a body region is one of several and is found by
    idx."""
    assert _slot_of(_Placeholder("TITLE (1)", 0)) == ("TITLE", 0)
    assert _slot_of(_Placeholder("BODY (2)", 14)) == ("BODY", 14)
    assert _slot_of(object()) is None


def test_only_the_slots_a_layout_defines_are_reset() -> None:
    """A placeholder with no counterpart on the layout it now points at
    inherits nothing, so stripping its box would not reset it -- it would send
    it to the origin at a default size. Those keep what they have."""
    layout = _Layout([("TITLE (1)", 0), ("BODY (2)", 14)])

    offered = _layout_slots(layout)

    assert offered == {("TITLE", 0), ("BODY", 14)}
    assert _slot_of(_Placeholder("TITLE (1)", 0)) in offered
    # The orphan: a subtitle the new layout has no region for.
    assert _slot_of(_Placeholder("SUBTITLE (4)", 1)) not in offered
