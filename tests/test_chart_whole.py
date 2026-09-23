"""A chart is one element and crosses the restyle exactly as it arrived.

Found on real decks: charts came back with bars and lines missing. The bars of
a chart drawn as shapes carry no copy, so the copy-matched protection never
covered them and the run cleanup deleted them as drawing; a native chart was
left behind by the XML route because it points at a chart part; and the
palette sweep rewrote every series colour inside the chart part.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches

from formatting_tool.ai.roles import ShapeRole, SlideRoles
from formatting_tool.apply.palette import sweep_to_palette
from formatting_tool.rebuild import charts
from formatting_tool.rebuild.builder import _leave_alone, rebuild

CANVAS = (13.333, 7.5)


def _chart_data() -> CategoryChartData:
    data = CategoryChartData()
    data.categories = ["Q1", "Q2", "Q3", "Q4"]
    data.add_series("Revenue", (12.0, 15.0, 9.0, 21.0))
    data.add_series("Cost", (8.0, 7.0, 6.0, 11.0))
    return data


def _master(tmp_path: Path) -> Path:
    path = tmp_path / "master.pptx"
    Presentation().save(str(path))
    return path


def _series_values(chart) -> list[tuple]:
    return [tuple(series.values) for series in chart.plots[0].series]


def _charts_on(path: Path) -> list:
    return [
        shape for slide in Presentation(str(path)).slides
        for shape in slide.shapes if getattr(shape, "has_chart", False)
    ]


def test_a_native_chart_crosses_the_xml_rebuild_whole(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1),
        Inches(6), Inches(4), _chart_data(),
    )
    source = tmp_path / "deck.pptx"
    deck.save(str(source))

    out = tmp_path / "out.pptx"
    result = rebuild(_master(tmp_path), source, out, route="xml")

    assert not result.dropped
    found = _charts_on(out)
    assert len(found) == 1
    chart = found[0]
    assert (chart.left, chart.top, chart.width, chart.height) == (
        Inches(1), Inches(1), Inches(6), Inches(4),
    )
    assert _series_values(chart.chart) == [
        (12.0, 15.0, 9.0, 21.0), (8.0, 7.0, 6.0, 11.0),
    ]
    # "Edit Data" still has a workbook to open.
    assert chart.chart.part.chart_workbook.xlsx_part is not None


def test_two_charts_get_two_parts(tmp_path: Path) -> None:
    """Named one at a time, so the second cannot overwrite the first."""
    deck = Presentation()
    for _ in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_chart(
            XL_CHART_TYPE.LINE, Inches(1), Inches(1), Inches(6), Inches(4),
            _chart_data(),
        )
    source = tmp_path / "deck.pptx"
    deck.save(str(source))

    out = tmp_path / "out.pptx"
    rebuild(_master(tmp_path), source, out, route="xml")

    found = _charts_on(out)
    assert len(found) == 2
    assert len({str(c.chart.part.partname) for c in found}) == 2


def test_a_chart_in_a_placeholder_keeps_its_own_frame(tmp_path: Path) -> None:
    """Written the way PowerPoint writes a chart dropped into a content slot:
    a graphic frame carrying `p:ph idx`, its own `p:xfrm` and the chart."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    body = slide.placeholders[1]
    body._element.getparent().remove(body._element)
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.7), Inches(2.2),
        Inches(5), Inches(3.5), _chart_data(),
    )
    nvPr = frame._element.find(".//" + qn("p:nvPr"))
    nvPr.append(parse_xml(f'<p:ph {nsdecls("p")} idx="1"/>'))
    assert frame.is_placeholder
    where = (Inches(0.7), Inches(2.2), Inches(5), Inches(3.5))
    source = tmp_path / "deck.pptx"
    deck.save(str(source))

    out = tmp_path / "out.pptx"
    result = rebuild(_master(tmp_path), source, out, route="xml")

    assert not result.dropped
    found = _charts_on(out)
    assert len(found) == 1
    assert not found[0].is_placeholder
    assert (found[0].left, found[0].top, found[0].width, found[0].height) == where


def test_every_piece_of_a_drawn_chart_is_left_alone() -> None:
    """A bar has no copy. It is protected by where it sits, not what it says."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    bar = slide.shapes.add_shape(1, Inches(2), Inches(3), Inches(0.5), Inches(2))
    beside = slide.shapes.add_shape(1, Inches(9), Inches(3), Inches(0.5), Inches(2))
    here = SlideRoles(
        slide=1,
        shapes=(ShapeRole(ref="s1", role="chart", box=(1.0, 1.0, 7.0, 6.0)),),
        reviewed=True,
    )

    assert _leave_alone(bar, here, CANVAS)
    assert not _leave_alone(beside, here, CANVAS)


def test_a_shape_drawn_over_a_native_chart_is_part_of_it() -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1),
        Inches(6), Inches(4), _chart_data(),
    )
    callout = slide.shapes.add_shape(1, Inches(3), Inches(2), Inches(1), Inches(0.3))
    boxes = charts.chart_boxes(slide.shapes)

    assert _leave_alone(frame, None, CANVAS, boxes)
    assert _leave_alone(callout, None, CANVAS, boxes)


def test_the_palette_sweep_leaves_a_chart_as_it_is(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1),
        Inches(6), Inches(4), _chart_data(),
    )
    for series, colour in zip(frame.chart.plots[0].series, ("3A7BD5", "3E80D0")):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(colour)
    path = tmp_path / "deck.pptx"
    deck.save(str(path))
    before = Presentation(str(path)).slides[0].shapes[0].chart.part.blob

    sweep_to_palette(path, {"brand": "#1F4E79", "accent": "#C00000"}, 2.0)

    after = Presentation(str(path)).slides[0].shapes[0].chart.part.blob
    assert after == before


def test_a_partname_template_numbers_the_right_place() -> None:
    assert charts._template("/ppt/charts/chart12.xml") == "/ppt/charts/chart%d.xml"
    assert (
        charts._template("/ppt/embeddings/Microsoft_Excel_Worksheet.xlsx")
        == "/ppt/embeddings/Microsoft_Excel_Worksheet%d.xlsx"
    )
