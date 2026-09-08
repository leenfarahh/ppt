"""Moving a set of shapes rather than one member of it.

The refusal this answers, from a real deck: 37 of 50 safe-margin fixes came
back "moving it would break its alignment with ...", and not one of them was
wrong. Slide 1 carried a column of shapes all sitting at left 0.48in against a
0.92in margin, so every one was 0.44in outside it, and moving any single one
in would have broken the column.

The set is what is out of place, so the set is what moves -- shapes no finding
named included, because a column half corrected is worse than one uncorrected.
After this, that deck refused none of the fifty for alignment and applied ten
instead of one; what still refuses, refuses because the set would land on
something, which is a different and much better answer.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.apply import apply_fixes
from formatting_tool.apply.fixers import COHORT
from formatting_tool.models import Category, Issue, Severity, Source

MARGINS = "top 0.4in, right 0.92in, bottom 0.73in, left 0.92in safe margin"


def _issue(rule_id: str, shape: str, shape_id: int, **kwargs) -> Issue:
    issue = Issue(
        category=Category.SPACE,
        severity=Severity.WARNING,
        message=kwargs.pop("message", f"{rule_id} finding"),
        source=Source.RULE,
        rule_id=rule_id,
        slide=1,
        shape=shape,
        shape_id=shape_id,
        deck="messy.pptx",
        expected=kwargs.pop("expected", MARGINS),
        **kwargs,
    )
    issue.id = issue.fingerprint()
    return issue


def _column(tmp_path: Path, lefts: list[float], extra=None):
    """A column of boxes at the given left edges, all outside the margin."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    made = []
    for index, left in enumerate(lefts):
        box = slide.shapes.add_textbox(
            Inches(left), Inches(1.0 + index * 0.8), Inches(4.0), Inches(0.6)
        )
        box.name = f"Row {index}"
        made.append(box)
    if extra is not None:
        left, top, width, height, name = extra
        block = slide.shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        block.name = name
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))
    return deck, made


def _lefts(path: Path) -> dict[str, float]:
    from pptx import Presentation

    return {
        s.name: round(s.left / 914400, 2)
        for s in Presentation(str(path)).slides[0].shapes
    }


# --------------------------------------------------------------------------- #
# The set moves
# --------------------------------------------------------------------------- #

def test_a_column_outside_the_margin_moves_together(tmp_path: Path) -> None:
    deck, made = _column(tmp_path, [0.48, 0.48, 0.48])
    issue = _issue("space.safe_margin", "Row 0", made[0].shape_id)
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out, selected=[issue.id])

    assert len(result.applied) == 1
    assert "the 2 shape(s) aligned with it" in result.applied[0].detail
    assert _lefts(out) == {"Row 0": 0.92, "Row 1": 0.92, "Row 2": 0.92}


def test_a_ragged_set_is_aligned_before_it_is_moved(tmp_path: Path) -> None:
    """Alignment within tolerance is not alignment, it is drift a tolerance
    forgave. Carrying it along would preserve it at the new position forever,
    so the set is squared up first and moved after."""
    deck, made = _column(tmp_path, [0.48, 0.50, 0.47])
    issue = _issue("space.safe_margin", "Row 0", made[0].shape_id)
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out, selected=[issue.id])

    assert len(result.applied) == 1
    assert "onto the edge first" in result.applied[0].detail
    # One edge, not three within a tolerance of each other.
    assert set(_lefts(out).values()) == {0.92}


def test_a_shape_alone_still_moves_alone(tmp_path: Path) -> None:
    """No set, nothing to bring along, and the old behaviour is unchanged."""
    deck, made = _column(tmp_path, [0.48])
    issue = _issue("space.safe_margin", "Row 0", made[0].shape_id)
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out, selected=[issue.id])

    assert len(result.applied) == 1
    assert "aligned with it" not in result.applied[0].detail
    assert _lefts(out)["Row 0"] == 0.92


# --------------------------------------------------------------------------- #
# And what stops it
# --------------------------------------------------------------------------- #

def test_the_whole_set_is_put_back_when_it_would_land_on_something(
    tmp_path: Path,
) -> None:
    """Every shape that moves is checked against its own neighbours, not only
    the one the finding named: a shape dragged along by its column can land on
    something just as easily."""
    # To the right of the column and clear of it, until the column moves
    # right by 0.44in and runs into it. Aligned with nothing, so it is an
    # obstacle rather than a member.
    deck, made = _column(
        tmp_path, [0.48, 0.48], extra=(4.6, 1.0, 2.0, 0.6, "In the way"),
    )
    issue = _issue("space.safe_margin", "Row 0", made[0].shape_id)
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out, selected=[issue.id])

    assert not result.applied
    assert "the set it belongs to" in result.skipped[0].detail
    kept = _lefts(out)
    assert kept["Row 0"] == 0.48 and kept["Row 1"] == 0.48


def test_only_the_hard_constraints_may_move_a_set() -> None:
    """A deck cannot ship with content outside the frame, so a column outside
    it together moves together. A grid snap is a preference, and dragging a
    neighbour to satisfy one is the failure that got `space.alignment_grid`
    disabled once already -- a section label pulled off the table it
    captioned, and with a set move it would take the table with it."""
    assert COHORT == {"space.safe_margin", "space.off_canvas"}
    assert "space.alignment_grid" not in COHORT
    assert "space.repeat_out_of_line" not in COHORT
    assert "space.satellite_offset" not in COHORT


def test_only_the_moving_axis_drags_anything(tmp_path: Path) -> None:
    """A horizontal nudge cannot break a shared top edge, so a shape sharing
    only that has no business being dragged sideways."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    subject = slide.shapes.add_textbox(
        Inches(0.48), Inches(2.0), Inches(3.0), Inches(0.6)
    )
    subject.name = "Subject"
    # Same top, nowhere near the same left: a row, not a column.
    neighbour = slide.shapes.add_textbox(
        Inches(8.0), Inches(2.0), Inches(3.0), Inches(0.6)
    )
    neighbour.name = "Same row"
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    issue = _issue("space.safe_margin", "Subject", subject.shape_id)
    out = tmp_path / "fixed.pptx"

    apply_fixes(deck, [issue], out, selected=[issue.id])

    kept = _lefts(out)
    assert kept["Subject"] == 0.92
    assert kept["Same row"] == 8.0        # it shares a top edge, not a column
