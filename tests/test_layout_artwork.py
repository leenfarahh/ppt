"""Artwork a slide inherits from the layout it is about to stop pointing at.

A photograph on a designed slide is often not on the slide at all. It is on
the layout, and the slide inherits it -- which is how a section divider shows
a full-bleed image while its own `p:cSld` holds nothing but a title. Point
that slide at another layout and the image is gone, and nothing reports it,
because no shape was lost: there was never a shape.

Measured on a deck built for it, before this existed:

    before   slide 2   layout 'Title Only'   0 on the slide, 1 on the layout
    after    slide 2   layout 'Cover'        0 on the slide, 0 on the layout

The other half is what does NOT travel. The old layouts hold two kinds of
picture: a section photograph, which belongs to the slide, and the old brand's
logo, which belongs to the design being replaced. Carrying the second would
stamp the previous client's mark on every slide of a rebranded deck, which is
worse than the missing photo and harder to notice.
"""

from __future__ import annotations

import copy
from pathlib import Path

import pytest

IMAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


def _png(path: Path, colour: tuple[int, int, int]) -> Path:
    Image = pytest.importorskip("PIL.Image", reason="Pillow builds the fixtures")
    Image.new("RGB", (240, 180), colour).save(path)
    return path


def _put_on_layout(layout, source) -> None:
    """Move a picture's XML onto a layout, the way a designed template has it."""
    part = source.part.related_part(source._element.blipFill.blip.rEmbed)
    rid = layout.part.relate_to(part, IMAGE_REL)
    element = copy.deepcopy(source._element)
    element.blipFill.blip.rEmbed = rid
    layout.shapes._spTree.append(element)


def _deck(tmp_path: Path, logo_on: list[int], photo_on: int):
    """A deck whose artwork lives on its layouts, not on its slides."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    scratch = prs.slides.add_slide(prs.slide_layouts[6])
    logo = scratch.shapes.add_picture(
        str(_png(tmp_path / "logo.png", (10, 10, 10))),
        Inches(11), Inches(0.2), Inches(1), Inches(0.6),
    )
    photo = scratch.shapes.add_picture(
        str(_png(tmp_path / "photo.png", (20, 90, 140))),
        Inches(6), Inches(1), Inches(3), Inches(3),
    )
    for index in logo_on:
        _put_on_layout(prs.slide_layouts[index], logo)
    _put_on_layout(prs.slide_layouts[photo_on], photo)

    slide = prs.slides.add_slide(prs.slide_layouts[photo_on])
    slide.shapes.title.text = "Our Understanding"

    for shape in (logo, photo):
        shape._element.getparent().remove(shape._element)

    path = tmp_path / "inherits.pptx"
    prs.save(str(path))
    return Presentation(str(path))


def _pictures_on(slide) -> list[str]:
    return [s.name for s in slide.shapes if s.shape_type == 13]


def test_the_photograph_is_copied_onto_the_slide(tmp_path: Path) -> None:
    from formatting_tool.rebuild.pictures import inherit_artwork

    prs = _deck(tmp_path, logo_on=[], photo_on=5)
    slide = prs.slides[-1]
    assert _pictures_on(slide) == []          # it is on the layout, not here

    carried = inherit_artwork(prs)

    assert len(carried) == 1
    assert len(_pictures_on(slide)) == 1


def test_the_old_brands_furniture_stays_behind(tmp_path: Path) -> None:
    """The logo is on three layouts because it is on every slide. The
    photograph is on the one layout drawn for it. Repetition is the signal the
    file actually carries -- size would be a guess."""
    from formatting_tool.rebuild.pictures import inherit_artwork

    prs = _deck(tmp_path, logo_on=[3, 4, 5], photo_on=5)
    slide = prs.slides[-1]
    on_layout = [s for s in slide.slide_layout.shapes if s.shape_type == 13]
    assert len(on_layout) == 2                # the logo and the photograph

    inherit_artwork(prs)

    assert len(_pictures_on(slide)) == 1      # only one of them travelled


def test_artwork_lands_behind_the_slides_own_shapes(tmp_path: Path) -> None:
    """A layout draws under the slide. A photograph arriving on top of the copy
    it used to sit behind is a different defect from the one being fixed."""
    from formatting_tool.rebuild.pictures import inherit_artwork

    prs = _deck(tmp_path, logo_on=[], photo_on=5)
    slide = prs.slides[-1]

    inherit_artwork(prs)

    assert slide.shapes[0].shape_type == 13   # first in the tree is furthest back


def test_a_deck_inheriting_nothing_is_left_alone(tmp_path: Path) -> None:
    """The property `master_apply` is built around: a deck that needs none of
    this reaches PowerPoint byte for byte as the designer saved it."""
    pytest.importorskip("pptx")
    from pptx import Presentation

    from formatting_tool.rebuild.pictures import inherit_artwork

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5])

    assert inherit_artwork(prs) == []


def test_a_picture_placeholder_on_the_layout_is_not_artwork(tmp_path: Path) -> None:
    """An empty picture placeholder is a slot, not a photograph. The slide's
    own shape is what fills it, and `freeze_slide` is what protects that."""
    pytest.importorskip("pptx")
    from pptx import Presentation

    from formatting_tool.rebuild.pictures import inherit_artwork

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[8])   # Picture with Caption

    assert inherit_artwork(prs) == []


# --------------------------------------------------------------------------- #
# Not only pictures
# --------------------------------------------------------------------------- #
#
# The first version of this carried pictures alone, and a real deck came back
# with its photographs intact and its connector lines and panel fills missing.
# Same defect, different tag name: a rounded rectangle behind the copy and the
# line joining an icon to it are as inherited as the photograph is.

def _drawn(tmp_path: Path, on_layouts: dict[int, list[str]]):
    """A deck whose layouts carry drawn shapes rather than pictures.

    `on_layouts` maps a layout index to the shapes to put on it, so one shape
    can be placed on several layouts to stand for brand furniture.
    """
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    scratch = prs.slides.add_slide(prs.slide_layouts[6])

    made = {}
    panel = scratch.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(3), Inches(4), Inches(2)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(0xF0, 0xF0, 0xF0)
    made["panel"] = panel
    made["line"] = scratch.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(2), Inches(1), Inches(2), Inches(3)
    )
    made["band"] = scratch.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13), Inches(0.4)
    )

    for index, names in on_layouts.items():
        for name in names:
            layout = prs.slide_layouts[index]
            layout.shapes._spTree.append(copy.deepcopy(made[name]._element))

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Dual-Track. One Direction"

    for shape in made.values():
        shape._element.getparent().remove(shape._element)

    path = tmp_path / "drawn.pptx"
    prs.save(str(path))
    return Presentation(str(path))


def _kinds_on(slide) -> list[str]:
    return [str(s.shape_type) for s in slide.shapes]


def test_lines_and_panels_travel_too(tmp_path: Path) -> None:
    from formatting_tool.rebuild.pictures import inherit_artwork

    prs = _drawn(tmp_path, {5: ["panel", "line"]})
    slide = prs.slides[-1]
    assert len(slide.shapes) == 1              # the title, and nothing else

    carried = inherit_artwork(prs)

    assert len(carried) == 2
    assert len(slide.shapes) == 3


def test_a_band_on_every_layout_stays_behind(tmp_path: Path) -> None:
    """Furniture is on many layouts because it is on every slide. A panel drawn
    for one slide is on the one layout drawn for it."""
    from formatting_tool.rebuild.pictures import inherit_artwork

    prs = _drawn(tmp_path, {3: ["band"], 4: ["band"], 5: ["band", "panel"]})
    slide = prs.slides[-1]

    inherit_artwork(prs)

    # The panel came; the band did not, though both sat on the same layout.
    assert len(slide.shapes) == 2


def test_the_layouts_own_order_is_kept(tmp_path: Path) -> None:
    """They are inserted at the front of the z-order, where a layout draws, and
    in the order the layout drew them -- not reversed by inserting each one at
    the same index."""
    from formatting_tool.rebuild.pictures import inherit_artwork

    prs = _drawn(tmp_path, {5: ["panel", "line"]})
    slide = prs.slides[-1]

    inherit_artwork(prs)

    kinds = _kinds_on(slide)
    assert "AUTO_SHAPE" in kinds[0]            # the panel, drawn first
    assert "LINE" in kinds[1]                  # then the line, over it
    assert "PLACEHOLDER" in kinds[2]           # the slide's own copy, on top
