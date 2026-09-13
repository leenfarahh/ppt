"""Placeholder text takes the master's colour, not the nearest one to it.

Snapping an off-palette title to the closest brand colour is a guess made in
front of an answer. The master has already said what colour a title on that
layout is -- and it is not one colour per deck: off a real master, a title is
white on the cover, the accent green on a section divider, and the dark text
colour on a content layout.

The cost of guessing is not theoretical. On that master, a red title on the
COVER measured nearest to the theme's dark colour, so the fix was black text on
a dark navy slide: unreadable, and a change the report described as putting the
deck on brand.

Three pieces, tested in the three places they live: reading the colour out of
the master, reporting it as the expected value, and applying it without the
deck-wide colour plan having an opinion.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from lxml import etree

from formatting_tool.models import (
    BrandGuidelines,
    Category,
    DeckProfile,
    Geometry,
    Issue,
    LayoutProfile,
    MasterSpec,
    ParagraphProfile,
    RunProfile,
    Severity,
    ShapeProfile,
    SlideProfile,
    Source,
)
from formatting_tool.rules import RuleContext
from formatting_tool.rules.colors import MASTER_SETS_IT, OffPaletteTextRule

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
RED = "D42A2A"


# --------------------------------------------------------------------------- #
# Reading it off the master
# --------------------------------------------------------------------------- #

def test_a_scheme_colour_on_a_layout_placeholder_resolves_through_the_theme(
    tmp_path: Path,
) -> None:
    """The value lives in default run properties, which is why nothing found
    it before: a layout placeholder holds no text, so it carries no run to
    read a colour from."""
    pytest.importorskip("pptx")
    from pptx import Presentation

    from formatting_tool.extract import read_deck

    prs = Presentation()
    layout = prs.slide_layouts[5]           # Title Only
    title = layout.placeholders[0]
    # What a designed master states: the title is the theme's second accent.
    # A layout placeholder's text body is `p:txBody`, not `a:txBody` -- the
    # one place in this file where the namespace is the presentation one. And
    # it already carries an empty `a:lstStyle`, which is the one to fill in:
    # adding a second would be a file PowerPoint never writes, and the reader
    # would take the first.
    body = title._element.find(f"{P}txBody")
    lst = body.find(f"{A}lstStyle")
    if lst is None:
        lst = etree.SubElement(body, f"{A}lstStyle")
    level = etree.SubElement(lst, f"{A}lvl1pPr")
    default = etree.SubElement(level, f"{A}defRPr")
    fill = etree.SubElement(default, f"{A}solidFill")
    scheme = etree.SubElement(fill, f"{A}schemeClr")
    scheme.set("val", "accent2")
    deck = tmp_path / "master.pptx"
    prs.save(str(deck))

    profile = read_deck(deck)
    read = next(
        layout for layout in profile.layouts if layout.name == "Title Only"
    )
    placeholder = next(
        shape for shape in read.placeholders
        if shape.placeholder_token == "TITLE"
    )

    assert placeholder.text_color_theme == "accent2"
    assert placeholder.text_color_hex == profile.theme_colors["accent2"].upper()


def test_a_placeholder_that_states_nothing_inherits_from_the_master(
    tmp_path: Path,
) -> None:
    """The rest of the chain, which is most of the value: a designed master
    states a colour on a few layouts and lets the others fall through to its
    own text styles. PowerPoint's default template does exactly that -- no
    layout states a title colour, and `p:titleStyle` says `tx1`."""
    pytest.importorskip("pptx")
    from pptx import Presentation

    from formatting_tool.extract import read_deck

    deck = tmp_path / "bare.pptx"
    Presentation().save(str(deck))

    profile = read_deck(deck)
    titles = [
        placeholder
        for layout in profile.layouts
        for placeholder in layout.placeholders
        if placeholder.placeholder_token in ("TITLE", "CENTER_TITLE")
    ]

    assert titles
    assert all(t.text_color_theme == "dk1" for t in titles)
    assert all(t.text_color_hex == profile.theme_colors["dk1"].upper() for t in titles)


def test_a_chain_that_states_nothing_anywhere_gives_no_answer() -> None:
    """Black is not a brand decision. Where the master, its placeholders and
    its text styles all say nothing, this says nothing and the nearest-palette
    fallback stands."""
    from formatting_tool.extract.textstyle import placeholder_color

    class _Bare:
        _element = etree.Element(f"{P}sp")
        placeholders: list = []

    assert placeholder_color(_Bare(), _Bare(), {}, "TITLE") == (None, None)


# --------------------------------------------------------------------------- #
# Reporting it
# --------------------------------------------------------------------------- #

def _ctx(*, layout_name: str, colour: str | None) -> RuleContext:
    """A slide whose title is red, on a master layout that states `colour`."""
    title = ShapeProfile(
        shape_id=7,
        name="Title 1",
        shape_type="PLACEHOLDER (14)",
        geometry=Geometry(left_in=0.9, top_in=0.4, width_in=11.5, height_in=1.2),
        placeholder_type="TITLE (1)",
        placeholder_idx=0,
        text="Approach",
        paragraphs=[ParagraphProfile(
            text="Approach",
            runs=[RunProfile(text="Approach", color_hex=RED)],
        )],
    )
    layout_ph = ShapeProfile(
        shape_id=2,
        name="Title Placeholder 1",
        shape_type="PLACEHOLDER (14)",
        geometry=Geometry(left_in=0.9, top_in=0.4, width_in=11.5, height_in=1.2),
        placeholder_type="TITLE (1)",
        placeholder_idx=0,
    )
    layout_ph.text_color_hex = colour
    deck = DeckProfile(
        path="messy.pptx", width_in=13.333, height_in=7.5,
        slides=[SlideProfile(number=1, layout_name=layout_name, shapes=[title])],
    )
    spec = MasterSpec(
        source="master.pptx", width_in=13.333, height_in=7.5,
        guidelines=BrandGuidelines(),
        palette={"navy": "1F2A44", "ink": "000000"},
        layouts=[LayoutProfile(name="Section Divider", index=0, shapes=[layout_ph])],
    )
    return RuleContext(deck=deck, spec=spec)


def test_the_expected_colour_is_the_one_the_master_states() -> None:
    found = list(OffPaletteTextRule().check(
        _ctx(layout_name="Section Divider", colour="004F71")
    ))

    assert len(found) == 1
    assert found[0].expected == f"{MASTER_SETS_IT} in #004F71"
    assert "Section Divider" in found[0].message


def test_a_slide_on_a_layout_the_master_does_not_have_falls_back() -> None:
    """A deck still sitting on its previous master names layouts this one has
    never heard of. Matching them by name would be reading another brand's
    file, so the nearest palette entry stands -- which is what this did before
    the master had anything to say."""
    found = list(OffPaletteTextRule().check(
        _ctx(layout_name="Some Other Master's Layout", colour="004F71")
    ))

    assert len(found) == 1
    assert MASTER_SETS_IT not in (found[0].expected or "")


def test_a_placeholder_the_master_states_no_colour_for_falls_back() -> None:
    found = list(OffPaletteTextRule().check(
        _ctx(layout_name="Section Divider", colour=None)
    ))

    assert len(found) == 1
    assert MASTER_SETS_IT not in (found[0].expected or "")


# --------------------------------------------------------------------------- #
# Applying it
# --------------------------------------------------------------------------- #

def test_the_colour_plan_does_not_get_a_say_in_the_master_s_own_value() -> None:
    """The plan exists to keep colours that MEAN something distinct from each
    other -- three pills in a legend, a chart's series -- by mapping them one
    to one onto the palette. A title is not an encoding: every title on a
    layout is the same colour on purpose, so two headings both becoming the
    master's heading colour is the right answer rather than a collision."""
    from formatting_tool.apply.applier import FixContext
    from formatting_tool.apply.fixers import _planned_target

    issue = Issue(
        category=Category.COLOR,
        severity=Severity.ERROR,
        message="off-palette",
        source=Source.RULE,
        rule_id="color.text.off_palette",
        slide=1,
        shape="Title 1",
        expected=f"{MASTER_SETS_IT} in #004F71",
        found=f"#{RED}",
    )

    class _Plan:
        def choice_for(self, _hex):
            raise AssertionError("the plan must not be consulted")

    context = FixContext(width_emu=12192000, height_emu=6858000)
    context.color_plan = _Plan()

    target, note = _planned_target(issue, context)

    assert target == "004F71"
    assert "master" in note
