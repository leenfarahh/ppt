"""Snapping an off-palette colour to the nearest palette entry.

The rule already measured the distance in CIEDE2000 to decide the colour was
off-palette, so the nearest entry is a number it has in hand and names on the
finding. The fixer reads it back rather than re-deriving it.

What these pin down is the narrowness of it: only the runs that carry the
colour the finding measured, only the fill or the outline the finding named,
and nothing at all when the finding names no colour.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.apply import apply_fixes, fixer_for
from formatting_tool.apply.fixers import _hex_of
from formatting_tool.models import Category, Issue, Severity, Source

NAVY = "1D252D"
BLACK = "000000"
GREY = "808080"


def _issue(rule_id: str, **kwargs) -> Issue:
    issue = Issue(
        category=kwargs.pop("category", Category.COLOR),
        severity=kwargs.pop("severity", Severity.ERROR),
        message=kwargs.pop("message", f"{rule_id} finding"),
        source=Source.RULE,
        rule_id=rule_id,
        deck="messy.pptx",
        **kwargs,
    )
    issue.id = issue.fingerprint()
    return issue


# --------------------------------------------------------------------------- #
# Reading the target off the finding
# --------------------------------------------------------------------------- #

def test_the_target_is_read_out_of_the_finding() -> None:
    """The rule writes "theme:dk1 #1D252D"; the label is not a colour."""
    assert _hex_of("theme:dk1 #1D252D") == NAVY
    assert _hex_of("#000000") == BLACK


def test_a_finding_naming_no_colour_yields_nothing() -> None:
    assert _hex_of("brand palette") is None
    assert _hex_of(None) is None


def test_seven_hex_digits_is_not_a_colour() -> None:
    """A longer run of hex would otherwise be truncated into a plausible one."""
    assert _hex_of("#1D252DE") is None


def test_both_palette_rules_have_a_fixer() -> None:
    assert fixer_for(_issue("color.text.off_palette")) is not None
    assert fixer_for(_issue("color.shape.off_palette")) is not None


# --------------------------------------------------------------------------- #
# Text
# --------------------------------------------------------------------------- #

def test_only_the_runs_carrying_that_colour_are_recoloured(tmp_path: Path) -> None:
    """A shape routinely holds one run in the brand colour and one somebody
    typed over. The finding is about the second."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.name = "Body"
    frame = box.text_frame
    frame.text = "on brand"
    frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(NAVY)
    stray = frame.add_paragraph()
    stray.text = "typed over"
    stray.runs[0].font.color.rgb = RGBColor.from_string(BLACK)
    untouched = frame.add_paragraph()
    untouched.text = "a third colour"
    untouched.runs[0].font.color.rgb = RGBColor.from_string(GREY)
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    issue = _issue(
        "color.text.off_palette", slide=1, shape="Body", shape_id=box.shape_id,
        found=f"#{BLACK}", expected=f"theme:dk1 #{NAVY}",
    )
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out)

    assert len(result.applied) == 1
    assert "1 run(s)" in result.applied[0].detail
    colours = [
        str(run.font.color.rgb)
        for paragraph in Presentation(str(out)).slides[0].shapes[0]
        .text_frame.paragraphs
        for run in paragraph.runs
    ]
    assert colours == [NAVY, NAVY, GREY]      # the stray moved, the third did not


def test_a_shape_whose_runs_do_not_carry_the_colour_is_left_alone(
    tmp_path: Path,
) -> None:
    """The report can be older than the deck. Recolouring on the strength of a
    shape name alone would repaint a run nobody measured."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.name = "Body"
    box.text_frame.text = "already fixed"
    box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(NAVY)
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    issue = _issue(
        "color.text.off_palette", slide=1, shape="Body", shape_id=box.shape_id,
        found=f"#{BLACK}", expected=f"theme:dk1 #{NAVY}",
    )
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out)

    assert result.applied == []
    assert "already correct" in result.skipped[0].detail


# --------------------------------------------------------------------------- #
# Shapes
# --------------------------------------------------------------------------- #

def test_a_fill_is_recoloured(tmp_path: Path) -> None:
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1),
                                   Inches(4), Inches(2))
    panel.name = "Panel"
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string(BLACK)
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    issue = _issue(
        "color.shape.off_palette", slide=1, shape="Panel",
        shape_id=panel.shape_id,
        message=f"Shape fill #{BLACK} is off-palette (nearest: theme:dk1).",
        found=f"#{BLACK}", expected=f"theme:dk1 #{NAVY}",
    )
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out)

    assert len(result.applied) == 1
    assert "fill" in result.applied[0].detail
    fixed = Presentation(str(out)).slides[0].shapes[0]
    assert str(fixed.fill.fore_color.rgb) == NAVY


def test_an_outline_finding_recolours_the_outline_not_the_fill(
    tmp_path: Path,
) -> None:
    """The finding says which of the two it measured, and repainting the wrong
    one would be a visible change nobody asked for."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1),
                                   Inches(4), Inches(2))
    panel.name = "Panel"
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string(GREY)
    panel.line.color.rgb = RGBColor.from_string(BLACK)
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    issue = _issue(
        "color.shape.off_palette", slide=1, shape="Panel",
        shape_id=panel.shape_id,
        message=f"Shape outline #{BLACK} is off-palette (nearest: theme:dk1).",
        found=f"#{BLACK}", expected=f"theme:dk1 #{NAVY}",
    )
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out)

    assert len(result.applied) == 1
    assert "outline" in result.applied[0].detail
    fixed = Presentation(str(out)).slides[0].shapes[0]
    assert str(fixed.line.color.rgb) == NAVY
    assert str(fixed.fill.fore_color.rgb) == GREY     # the fill did not move


# --------------------------------------------------------------------------- #
# The fallback target
# --------------------------------------------------------------------------- #
#
# A colour the rule found no intended entry for is snapped to the nearest one
# anyway, because a finding that corrects nothing is the more common
# complaint. The fixer has to apply it and has to say what it did, since this
# is the one recolour that changes a colour's hue.

ORANGE = "E97132"
RED = "A32020"
FALLBACK = f"nearest theme:accent2 #{ORANGE}"


def test_a_marked_fallback_is_recognised() -> None:
    from formatting_tool.apply.fixers import _is_fallback

    assert _is_fallback(_issue("color.text.off_palette", expected=FALLBACK))
    assert not _is_fallback(
        _issue("color.text.off_palette", expected=f"theme:dk1 #{NAVY}")
    )


def test_the_hex_is_still_read_out_of_a_marked_target() -> None:
    """The mark is a prefix, not a replacement: the fixer reads the target the
    same way it reads any other."""
    assert _hex_of(FALLBACK) == ORANGE


def test_a_red_is_recoloured_to_the_nearest_entry(tmp_path: Path) -> None:
    """The case this exists for: a red title on a palette holding no red."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.name = "Title"
    box.text_frame.text = "a red heading"
    box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(RED)
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    issue = _issue(
        "color.text.off_palette", slide=1, shape="Title", shape_id=box.shape_id,
        found=f"#{RED}", expected=FALLBACK,
    )
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out)

    assert len(result.applied) == 1
    detail = result.applied[0].detail
    # The warning is the whole reason this is allowed to apply.
    assert "nearest" in detail
    assert "reads as a different colour" in detail
    run = (Presentation(str(out)).slides[0].shapes[0]
           .text_frame.paragraphs[0].runs[0])
    assert str(run.font.color.rgb) == ORANGE


def test_a_confident_target_is_not_hedged(tmp_path: Path) -> None:
    """The warning has to distinguish, so it must be absent when the rule
    named the entry the colour was meant to be."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.name = "Body"
    box.text_frame.text = "off by a little"
    box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(BLACK)
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    issue = _issue(
        "color.text.off_palette", slide=1, shape="Body", shape_id=box.shape_id,
        found=f"#{BLACK}", expected=f"theme:dk1 #{NAVY}",
    )
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out)

    assert len(result.applied) == 1
    assert "reads as a different colour" not in result.applied[0].detail


def test_a_finding_with_no_colour_at_all_is_still_left_alone(
    tmp_path: Path,
) -> None:
    """The empty-palette case. Nothing to be nearest to, so nothing to apply,
    and the fixer must not invent one."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.name = "Title"
    box.text_frame.text = "a red heading"
    box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(RED)
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    issue = _issue(
        "color.text.off_palette", slide=1, shape="Title", shape_id=box.shape_id,
        found=f"#{RED}", expected="brand palette",
    )
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, [issue], out)

    assert result.applied == []
    assert "design call" in result.skipped[0].detail
    unchanged = (Presentation(str(out)).slides[0].shapes[0]
                 .text_frame.paragraphs[0].runs[0])
    assert str(unchanged.font.color.rgb) == RED
