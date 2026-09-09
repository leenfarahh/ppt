"""Messages to whoever is making the deck, taken off before it goes out.

A deck being worked on collects them: "Design - can you redo the map and make
the colour contrast stronger", "TBC with legal", a coloured comment box parked
in a corner. They are addressed to the production process and they must not
reach a client.

Only the AI layer can tell one from a caption -- the question is who the text
is talking to, which is not in the geometry -- so it reports them and proposes
`remove_note`. This is the only op that takes something off a slide, and
removal is the one change a designer cannot check by looking at the result:
everything else leaves evidence, this leaves a gap. So most of what follows is
about what it refuses to do.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.apply import apply_fixes
from formatting_tool.models import (
    FIX_OPS,
    Category,
    FixAction,
    Issue,
    Severity,
    Source,
)

NOTE = "Design - can you redo the map and make the color contrast stronger"
CAPTION = "Select a plot size to continue"


def _issue(shape: str, shape_id: int, confidence: float = 0.9, text: str = NOTE):
    issue = Issue(
        category=Category.PRODUCTION_NOTE,
        severity=Severity.WARNING,
        message="This is a message to the designer, not deck content.",
        source=Source.AI,
        slide=1,
        shape=shape,
        shape_id=shape_id,
        deck="messy.pptx",
        found=text,
        confidence=confidence,
        fix=FixAction(op="remove_note", shape=shape, shape_id=shape_id),
    )
    issue.id = issue.fingerprint()
    return issue


def _spec():
    from formatting_tool.extract import read_deck
    from formatting_tool.extract.master_spec import derive_master_spec
    from formatting_tool.models import BrandGuidelines

    return derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())


def _deck(tmp_path: Path, *, in_placeholder: bool = False, text: str = NOTE):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])       # Title Only
    slide.shapes.title.text = "Land area"
    if in_placeholder:
        slide.shapes.title.text = text
        target = slide.shapes.title
    else:
        target = slide.shapes.add_textbox(
            Inches(9.0), Inches(0.4), Inches(3.0), Inches(0.9)
        )
        target.name = "Comment box"
        target.text_frame.text = text
    path = tmp_path / "messy.pptx"
    prs.save(str(path))
    return path, target.name, target.shape_id


def _names(path: Path) -> list[str]:
    from pptx import Presentation

    return [s.name for s in Presentation(str(path)).slides[0].shapes]


# --------------------------------------------------------------------------- #
# It works
# --------------------------------------------------------------------------- #

def test_remove_note_is_an_op() -> None:
    assert "remove_note" in FIX_OPS


def test_the_note_comes_off_the_slide(tmp_path: Path) -> None:
    deck, name, shape_id = _deck(tmp_path)
    out = tmp_path / "clean.pptx"
    issue = _issue(name, shape_id)

    result = apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    assert len(result.applied) == 1
    assert name not in _names(out)


def test_what_it_said_is_quoted_back(tmp_path: Path) -> None:
    """A report saying "removed a shape" cannot be checked. One that quotes
    what the shape said can."""
    deck, name, shape_id = _deck(tmp_path)
    issue = _issue(name, shape_id)

    result = apply_fixes(
        deck, [issue], tmp_path / "clean.pptx", selected=[issue.id], spec=_spec()
    )

    assert "redo the map" in result.applied[0].detail


def test_removals_are_listed_apart_from_everything_else(tmp_path: Path) -> None:
    """Everything else leaves something on the slide to look at."""
    deck, name, shape_id = _deck(tmp_path)
    issue = _issue(name, shape_id)

    result = apply_fixes(
        deck, [issue], tmp_path / "clean.pptx", selected=[issue.id], spec=_spec()
    )

    assert len(result.removed) == 1
    assert result.removed[0].issue.slide == 1


def test_the_rest_of_the_slide_is_untouched(tmp_path: Path) -> None:
    deck, name, shape_id = _deck(tmp_path)
    out = tmp_path / "clean.pptx"
    issue = _issue(name, shape_id)

    apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    from pptx import Presentation

    assert Presentation(str(out)).slides[0].shapes.title.text == "Land area"


# --------------------------------------------------------------------------- #
# And what it refuses
# --------------------------------------------------------------------------- #

def test_a_placeholder_is_never_removed(tmp_path: Path) -> None:
    """It is part of the layout's structure, it comes back empty on the next
    rebuild, and a title box holding a note is a title box with a note typed
    into it: the copy is the problem, not the shape."""
    deck, name, shape_id = _deck(tmp_path, in_placeholder=True)
    issue = _issue(name, shape_id)

    result = apply_fixes(
        deck, [issue], tmp_path / "clean.pptx", selected=[issue.id], spec=_spec()
    )

    assert not result.applied
    assert "layout placeholder" in result.skipped[0].detail


def test_an_unsure_model_removes_nothing(tmp_path: Path) -> None:
    """Leaving a note in costs a designer ten seconds. Taking a caption out of
    a client deck is a defect nobody sees until the client does."""
    deck, name, shape_id = _deck(tmp_path)
    out = tmp_path / "clean.pptx"
    issue = _issue(name, shape_id, confidence=0.6)

    result = apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    assert not result.applied
    assert "more certainty" in result.skipped[0].detail
    assert name in _names(out)


def test_the_bar_for_removing_is_higher_than_for_reporting() -> None:
    """0.5 separates a judgement call from a defect. Removal is not worth the
    same as everything else, so it does not use the same number."""
    from formatting_tool.apply.fixers import _SURE_ENOUGH

    assert _SURE_ENOUGH > 0.5


def test_a_shape_with_no_text_is_not_a_note(tmp_path: Path) -> None:
    """A shape this names and cannot read is not one to delete."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    blob = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
    )
    blob.name = "Decoration"
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))
    issue = _issue("Decoration", blob.shape_id)

    result = apply_fixes(
        deck, [issue], tmp_path / "clean.pptx", selected=[issue.id], spec=_spec()
    )

    assert not result.applied
    assert "no text here" in result.skipped[0].detail


def test_nothing_is_removed_without_a_master(tmp_path: Path) -> None:
    """Every AI proposal is checked against the brand system before it is
    applied, and removal is not the exception."""
    deck, name, shape_id = _deck(tmp_path)
    out = tmp_path / "clean.pptx"
    issue = _issue(name, shape_id)

    result = apply_fixes(deck, [issue], out, selected=[issue.id])

    assert not result.applied
    assert name in _names(out)


# --------------------------------------------------------------------------- #
# The note is moved, not destroyed
# --------------------------------------------------------------------------- #
#
# Deleting a production note is correct about the slide and wrong about the
# note: somebody wrote it on purpose, it is often the only record that the
# thing it asks for is outstanding, and a deletion is the one edit nobody can
# check by looking at the result. So the note is copied out FIRST and the shape
# comes off only once the copy exists.

def _comment_parts(path: Path) -> list[str]:
    import zipfile

    with zipfile.ZipFile(path) as archive:
        return [n for n in archive.namelist() if "comment" in n.lower()]


def _comment_text(path: Path) -> str:
    import zipfile

    with zipfile.ZipFile(path) as archive:
        return "".join(
            archive.read(n).decode("utf-8", "ignore") for n in _comment_parts(path)
        )


def _notes_text(path: Path) -> str:
    from pptx import Presentation

    slide = Presentation(str(path)).slides[0]
    if not slide.has_notes_slide:
        return ""
    return slide.notes_slide.notes_text_frame.text


def test_the_note_becomes_a_real_powerpoint_comment(tmp_path: Path) -> None:
    """The destination a designer already looks in for work assigned to them.

    Written by PowerPoint's own API rather than assembled here: a current
    build writes a modern comment, which is a set of parts and author GUIDs
    not worth hand-rolling.
    """
    from formatting_tool import powerpoint

    if not powerpoint.available():
        pytest.skip("desktop PowerPoint is needed to make a comment")

    deck, name, shape_id = _deck(tmp_path)
    issue = _issue(name, shape_id)
    out = tmp_path / "clean.pptx"

    result = apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    assert [r.where for r in result.notes_lifted] == ["comment"]
    assert NOTE in _comment_text(out)          # the note survived, verbatim
    assert "Comment box" not in _names(out)    # and the shape did not


def test_without_powerpoint_the_note_goes_to_the_notes_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The fallback. Not the Comments pane and nobody is notified by it, but
    off the slide, beside the right slide, and still there tomorrow."""
    from formatting_tool import powerpoint

    monkeypatch.setattr(powerpoint, "available", lambda: False)

    deck, name, shape_id = _deck(tmp_path)
    issue = _issue(name, shape_id)
    out = tmp_path / "clean.pptx"

    result = apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    assert [r.where for r in result.notes_lifted] == ["notes"]
    assert NOTE in _notes_text(out)
    assert "Comment box" not in _names(out)
    assert "notes page" in result.applied[0].detail


def test_a_note_that_cannot_be_copied_out_stays_on_the_slide(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The failure that matters, and the reason for the order of the two halves.

    With both destinations broken there is no copy, so the shape must not come
    off. A note left on a slide is a defect somebody notices; a note deleted
    with no copy anywhere is one nobody can.
    """
    from formatting_tool import powerpoint
    from formatting_tool.apply import notes as notes_module

    monkeypatch.setattr(powerpoint, "available", lambda: False)

    def broken(deck, lifts):
        raise RuntimeError("no writable destination")

    monkeypatch.setattr(notes_module, "_through_notes_pages", broken)

    deck, name, shape_id = _deck(tmp_path)
    issue = _issue(name, shape_id)
    out = tmp_path / "clean.pptx"

    result = apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    assert [r.where for r in result.notes_lifted] == ["kept"]
    assert "Comment box" in _names(out)        # still there, which is the point
    assert result.applied == []                # and not reported as corrected
    assert any("still on the slide" in o.detail for o in result.skipped)


def test_the_note_is_quoted_in_the_report_whatever_became_of_it(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Every path has to leave the designer able to re-check it, and that means
    the report quotes the note rather than saying a shape was dealt with."""
    from formatting_tool import powerpoint

    monkeypatch.setattr(powerpoint, "available", lambda: False)

    deck, name, shape_id = _deck(tmp_path)
    issue = _issue(name, shape_id)
    out = tmp_path / "clean.pptx"

    result = apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    assert result.notes_lifted[0].lift.text == NOTE
    lines = [o.detail for o in result.applied + result.skipped]
    assert any("redo the map" in line for line in lines)


def test_the_moved_note_says_where_it_came_from(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A bare instruction appearing in a comment reads as coming from nobody.
    It is prefixed so it reads as a record of something moved."""
    from formatting_tool import powerpoint
    from formatting_tool.apply.notes import PREFIX

    monkeypatch.setattr(powerpoint, "available", lambda: False)

    deck, name, shape_id = _deck(tmp_path)
    issue = _issue(name, shape_id)
    out = tmp_path / "clean.pptx"

    apply_fixes(deck, [issue], out, selected=[issue.id], spec=_spec())

    assert PREFIX in _notes_text(out)
