"""Two corrections the design check makes to a set of siblings, in PowerPoint.

- `match_format`: the body copy of two cards set at two sizes. The model says
  which card is right; its text is copied onto the other, and where the copy
  does not fit at that size a PAIR comes down together to the largest size
  that does.
- `widen` on a box whose words are whole: one label down a column drawn at a
  third of the width the others share, wrapping to four lines where they take
  one. Widening used to refuse it, because nothing was breaking mid-word.

Driven through PowerPoint, as the corrections are, so skipped without it.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool import powerpoint
from formatting_tool.apply.qafix import Step, apply_steps

pytestmark = pytest.mark.skipif(
    not powerpoint.available(), reason="desktop PowerPoint is needed"
)

COPY = ("At ADFD, our people are the driving force behind our success, the "
        "heart of the institution and the foundation of its achievements.")


def _deck(tmp_path: Path, build) -> tuple[Path, dict]:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    build(slide)
    path = tmp_path / "siblings.pptx"
    prs.save(str(path))
    where = {s.name: (i, s.shape_id) for i, s in enumerate(slide.shapes, 1)}
    return path, where


def _box(slide, name, left, top, width, height, text, size):
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width),
                                   Inches(height))
    from pptx.enum.text import MSO_AUTO_SIZE

    box.name = name
    box.text_frame.auto_size = MSO_AUTO_SIZE.NONE   # a fixed card, as drawn
    box.text_frame.word_wrap = True
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    return box


def _step(where, op, shape, reference=None, **kwargs) -> Step:
    index, shape_id = where[shape]
    extra = {}
    if reference is not None:
        r_index, r_id = where[reference]
        extra = dict(parent_id=r_id, parent=reference, parent_path=(r_index,))
    return Step(op=op, slide=1, shape_id=shape_id, shape=shape,
                path=(index,), **extra, **kwargs)


def _sizes(path: Path) -> dict:
    from pptx import Presentation
    from pptx.util import Pt

    default = Pt(18)

    # PowerPoint drops `sz` where it equals the default, 18pt for these boxes.
    return {
        s.name: (s.text_frame.paragraphs[0].runs[0].font.size or default).pt
        for s in Presentation(str(path)).slides[0].shapes
    }



def test_the_other_card_takes_the_reference_size(tmp_path: Path) -> None:
    deck, where = _deck(tmp_path, lambda slide: (
        _box(slide, "Card 1", 1.0, 2.0, 4.2, 3.0, COPY, 15),
        _box(slide, "Card 2", 6.0, 2.0, 4.2, 3.0, COPY, 10),
    ))
    out = tmp_path / "out.pptx"

    result = apply_steps(deck, out, [
        _step(where, "match_format", "Card 2", reference="Card 1"),
    ])

    assert len(result.applied) == 1, [s.reason for s in result.skipped]
    assert _sizes(out) == {"Card 1": 15.0, "Card 2": 15.0}


def test_a_pair_comes_down_together_to_the_size_that_fits(tmp_path: Path) -> None:
    deck, where = _deck(tmp_path, lambda slide: (
        _box(slide, "Card 1", 1.0, 2.0, 4.2, 3.0, COPY, 20),
        _box(slide, "Card 2", 6.0, 2.0, 4.2, 1.2, COPY, 10),
    ))
    out = tmp_path / "out.pptx"

    result = apply_steps(deck, out, [
        _step(where, "match_format", "Card 2", reference="Card 1",
              shared_fit=True),
    ])

    assert len(result.applied) == 1, [s.reason for s in result.skipped]
    sizes = _sizes(out)
    assert sizes["Card 1"] == sizes["Card 2"]
    assert 9.0 <= sizes["Card 2"] < 20.0


def test_without_a_shared_fit_a_card_that_would_spill_is_left(
    tmp_path: Path,
) -> None:
    deck, where = _deck(tmp_path, lambda slide: (
        _box(slide, "Card 1", 1.0, 2.0, 4.2, 3.0, COPY, 20),
        _box(slide, "Card 2", 6.0, 2.0, 4.2, 1.2, COPY, 10),
    ))
    out = tmp_path / "out.pptx"

    result = apply_steps(deck, out, [
        _step(where, "match_format", "Card 2", reference="Card 1"),
    ])

    assert result.applied == []
    assert "no longer fits" in result.skipped[0].reason
    assert _sizes(out) == {"Card 1": 20.0, "Card 2": 10.0}


def test_a_narrow_box_takes_its_columns_width(tmp_path: Path) -> None:
    from pptx import Presentation

    deck, where = _deck(tmp_path, lambda slide: (
        _box(slide, "Item 1", 6.8, 2.2, 5.8, 0.75, "Promoting growth globally", 12),
        _box(slide, "Item 2", 6.8, 3.4, 2.0, 0.85,
             "Delivering aid and investments with tangible development impact", 12),
        _box(slide, "Item 3", 6.8, 4.6, 5.8, 0.75, "Facilitating export financing", 12),
    ))
    out = tmp_path / "out.pptx"

    result = apply_steps(deck, out, [_step(where, "widen", "Item 2")])

    assert len(result.applied) == 1, [s.reason for s in result.skipped]
    assert "the boxes it lines up with" in result.applied[0].detail
    item = [s for s in Presentation(str(out)).slides[0].shapes
            if s.name == "Item 2"][0]
    assert item.width / 914400 == pytest.approx(5.8, abs=0.02)


def test_a_box_already_as_wide_as_its_column_is_left(tmp_path: Path) -> None:
    deck, where = _deck(tmp_path, lambda slide: (
        _box(slide, "Item 1", 6.8, 2.2, 5.8, 0.75, "Promoting growth globally", 12),
        _box(slide, "Item 2", 6.8, 3.4, 5.8, 0.75, "Delivering aid", 12),
        _box(slide, "Item 3", 6.8, 4.6, 5.8, 0.75, "Facilitating export", 12),
    ))

    result = apply_steps(deck, tmp_path / "out.pptx",
                         [_step(where, "widen", "Item 2")])

    assert result.applied == []
    assert "already as wide" in result.skipped[0].reason
