"""A circular portrait came out square and against the left margin.

Found on a real deck, in the before/after preview: slide 15 had a headshot
cropped to a circle inside a ring of decorative arcs. After the rebuild the
arcs were exactly where they had been and the photo was a bare rectangle at
x=0. That the arcs survived is the diagnosis -- they are ordinary autoshapes
carrying their own `a:xfrm`, and the photo was a picture PLACEHOLDER whose
circle and position lived on the OLD layout. Re-point the slide and both are
gone: geometry falls back to the implicit `rect` and the frame to whatever the
new layout puts at that idx, which is the origin when it puts nothing.

The XML helpers are tested directly, and then the whole rebuild is run over a
fixture built the way PowerPoint builds one: `insert_picture` deliberately
writes a `p:pic` with no `a:xfrm` so its extents inherit, which is exactly the
shape of the bug.
"""

from __future__ import annotations

import struct
import zlib
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from formatting_tool.rebuild.builder import (
    _bake_frame,
    _bake_look,
    _carries_picture,
    _strip_ph,
    rebuild,
)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = f'xmlns:a="{A}" xmlns:p="{P}"'

# The frame the fixture's layout gives its picture placeholder, in EMU.
FRAME = (1792288, 612775, 5486400, 4114800)


def _spPr(inner: str = "") -> etree._Element:
    from pptx.oxml import parse_xml

    return parse_xml(f"<p:spPr {NS}>{inner}</p:spPr>")


def _names(element) -> list[str]:
    return [child.tag.rsplit("}", 1)[-1] for child in element]


# --------------------------------------------------------------------------- #
# Baking the frame
# --------------------------------------------------------------------------- #

def test_an_inherited_frame_is_written_onto_the_shape() -> None:
    spPr = _spPr()

    _bake_frame(spPr, FRAME)

    xfrm = spPr.find(f"{{{A}}}xfrm")
    assert (xfrm.x, xfrm.y, xfrm.cx, xfrm.cy) == FRAME


def test_a_frame_the_shape_states_itself_is_left_alone() -> None:
    """The slide was overriding the layout before the move and still should."""
    spPr = _spPr(
        '<a:xfrm><a:off x="1" y="2"/><a:ext cx="3" cy="4"/></a:xfrm>'
    )

    _bake_frame(spPr, FRAME)

    xfrm = spPr.find(f"{{{A}}}xfrm")
    assert (xfrm.x, xfrm.y, xfrm.cx, xfrm.cy) == (1, 2, 3, 4)


def test_an_unreadable_frame_writes_nothing() -> None:
    """`_frame_of` returns None where the inheritance chain breaks, and a
    partial `a:xfrm` would be worse than none: PowerPoint reads a missing
    `a:ext` as zero and the shape vanishes."""
    spPr = _spPr()

    _bake_frame(spPr, None)

    assert _names(spPr) == []


# --------------------------------------------------------------------------- #
# Baking the look
# --------------------------------------------------------------------------- #

def test_the_layouts_geometry_and_outline_come_down() -> None:
    spPr = _spPr()
    source = _spPr(
        '<a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom>'
        '<a:ln w="38100"><a:solidFill><a:srgbClr val="00B0A0"/></a:solidFill></a:ln>'
        '<a:effectLst><a:outerShdw blurRad="50800"/></a:effectLst>'
    )

    _bake_look(spPr, source)

    assert _names(spPr) == ["prstGeom", "ln", "effectLst"]
    assert spPr.find(f"{{{A}}}prstGeom").get("prst") == "ellipse"


def test_only_one_geometry_is_copied() -> None:
    """`a:custGeom` and `a:prstGeom` are a schema choice. A file carrying both
    is one PowerPoint refuses, and the drawn outline is the one that could not
    have been expressed as a preset, so it wins."""
    source = _spPr(
        '<a:custGeom><a:pathLst/></a:custGeom>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
    )
    spPr = _spPr()

    _bake_look(spPr, source)

    assert _names(spPr) == ["custGeom"]


def test_geometry_the_shape_states_itself_is_left_alone() -> None:
    spPr = _spPr('<a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>')
    source = _spPr('<a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom>')

    _bake_look(spPr, source)

    assert spPr.find(f"{{{A}}}prstGeom").get("prst") == "roundRect"


def test_a_layout_with_nothing_to_give_leaves_the_shape_untouched() -> None:
    spPr = _spPr('<a:solidFill><a:srgbClr val="FF0000"/></a:solidFill>')

    _bake_look(spPr, _spPr())
    _bake_look(spPr, None)

    assert _names(spPr) == ["solidFill"]


# --------------------------------------------------------------------------- #
# Dropping the p:ph
# --------------------------------------------------------------------------- #

PIC = f"""<p:pic {NS}>
  <p:nvPicPr>
    <p:cNvPr id="4" name="Portrait"/>
    <p:cNvPicPr/>
    <p:nvPr><p:ph type="pic" idx="1"/></p:nvPr>
  </p:nvPicPr>
  <p:blipFill><a:blip r:embed="rId2" xmlns:r="x"/></p:blipFill>
  <p:spPr/>
</p:pic>"""


def _pic():
    from pptx.oxml import parse_xml

    return parse_xml(PIC)


def test_the_placeholder_marker_is_removed() -> None:
    """Left in place it is a standing invitation to inherit again, from a
    layout that knows nothing about this shape."""
    pic = _pic()

    assert _strip_ph(pic) is True

    assert pic.ph is None
    assert pic.has_ph_elm is False
    # The chain above it is required content and stays.
    assert pic.find(f"{{{P}}}nvPicPr/{{{P}}}nvPr") is not None


def test_a_shape_that_was_never_a_placeholder_is_not_disturbed() -> None:
    pic = _pic()
    _strip_ph(pic)

    assert _strip_ph(pic) is False


# --------------------------------------------------------------------------- #
# Which shapes qualify
# --------------------------------------------------------------------------- #

class _FakeShape:
    def __init__(self, element, text=None):
        self._element = element
        self.has_text_frame = text is not None
        self.text_frame = type("F", (), {"text": text or ""})()


def test_a_picture_qualifies() -> None:
    assert _carries_picture(_FakeShape(_pic())) is True


def test_a_shape_with_a_photo_fill_qualifies() -> None:
    """The other way PowerPoint authors one: an autoshape whose fill is the
    image. Same inheritance, same failure."""
    from pptx.oxml import parse_xml

    sp = parse_xml(
        f'<p:sp {NS}><p:spPr><a:blipFill>'
        f'<a:blip r:embed="rId2" xmlns:r="x"/></a:blipFill></p:spPr></p:sp>'
    )
    assert _carries_picture(_FakeShape(sp, text="")) is True


def test_a_titled_panel_with_a_photo_fill_does_not_qualify() -> None:
    """Its words are what has to move into the new layout's type. Freezing it
    would reproduce exactly the drift the rebuild exists to remove."""
    from pptx.oxml import parse_xml

    sp = parse_xml(
        f'<p:sp {NS}><p:spPr><a:blipFill>'
        f'<a:blip r:embed="rId2" xmlns:r="x"/></a:blipFill></p:spPr></p:sp>'
    )
    assert _carries_picture(_FakeShape(sp, text="Team & References")) is False


def test_an_empty_picture_placeholder_does_not_qualify() -> None:
    """A prompt with no photo in it. There is nothing to preserve, and it
    should still be offered the new layout's picture frame."""
    from pptx.oxml import parse_xml

    sp = parse_xml(
        f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="4" name="Picture Placeholder"/>'
        f'<p:cNvSpPr/><p:nvPr><p:ph type="pic" idx="1"/></p:nvPr></p:nvSpPr>'
        f'<p:spPr/></p:sp>'
    )
    assert _carries_picture(_FakeShape(sp, text="")) is False


# --------------------------------------------------------------------------- #
# End to end
# --------------------------------------------------------------------------- #

def _png() -> bytes:
    """A 1x1 opaque PNG, hand-built so the test needs no fixture file."""
    def chunk(kind: bytes, body: bytes) -> bytes:
        return (
            struct.pack(">I", len(body))
            + kind
            + body
            + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)
        )

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(b"\x00\x00\x00\x00"))
        + chunk(b"IEND", b"")
    )


def _fixture(tmp_path: Path) -> tuple[Path, Path]:
    """A deck whose photo is a circle only because its layout says so.

    Built the way a designed template is: the geometry and the ring go on the
    LAYOUT's picture placeholder, and the slide carries just the image. That
    `insert_picture` writes no `a:xfrm` is not a convenience here, it is the
    documented behaviour that produces the bug.
    """
    image = tmp_path / "portrait.png"
    image.write_bytes(_png())

    source = Presentation()
    layout = source.slide_masters[0].slide_layouts[8]     # Picture with Caption
    placeholder = layout.placeholders[1]                  # PICTURE, idx 1
    spPr = placeholder._element.spPr
    spPr._insert_prstGeom(
        etree.fromstring(f'<a:prstGeom xmlns:a="{A}" prst="ellipse"><a:avLst/></a:prstGeom>')
    )
    spPr._insert_ln(
        etree.fromstring(
            f'<a:ln xmlns:a="{A}" w="38100"><a:solidFill>'
            f'<a:srgbClr val="00B0A0"/></a:solidFill></a:ln>'
        )
    )

    slide = source.slides.add_slide(layout)
    slide.placeholders[0].text_frame.text = "Team & References"
    slide.placeholders[1].insert_picture(str(image))

    deck = tmp_path / "deck.pptx"
    source.save(str(deck))

    master = tmp_path / "master.pptx"
    Presentation().save(str(master))
    return master, deck


def _pictures(path: Path) -> list:
    return [
        shape._element
        for slide in Presentation(str(path)).slides
        for shape in slide.shapes
        if shape._element.tag == qn("p:pic")
    ]


def test_the_source_deck_really_does_carry_nothing(tmp_path: Path) -> None:
    """Guards the fixture itself. If `insert_picture` ever starts writing an
    `a:xfrm`, every assertion below would pass for the wrong reason."""
    _master, deck = _fixture(tmp_path)

    pic = _pictures(deck)[0]

    assert pic.spPr.xfrm is None
    assert pic.spPr.prstGeom is None
    assert pic.has_ph_elm is True


def test_the_photo_keeps_its_circle_and_its_place(tmp_path: Path) -> None:
    master, deck = _fixture(tmp_path)
    out = tmp_path / "out.pptx"

    rebuild(master, deck, out, route="xml")

    pic = _pictures(out)[0]
    assert pic.spPr.prstGeom is not None, "the photo came out a bare rectangle"
    assert pic.spPr.prstGeom.get("prst") == "ellipse"
    xfrm = pic.spPr.xfrm
    assert xfrm is not None, "the photo came out at the origin"
    assert (xfrm.x, xfrm.y, xfrm.cx, xfrm.cy) == FRAME


def test_the_photo_stops_being_a_placeholder(tmp_path: Path) -> None:
    """So nothing re-resolves it against a layout that has never heard of it,
    and `_claim` cannot hand it a text placeholder and then quietly copy
    nothing into it."""
    master, deck = _fixture(tmp_path)
    out = tmp_path / "out.pptx"

    rebuild(master, deck, out, route="xml")

    assert _pictures(out)[0].has_ph_elm is False


def test_the_ring_comes_down_with_it(tmp_path: Path) -> None:
    master, deck = _fixture(tmp_path)
    out = tmp_path / "out.pptx"

    rebuild(master, deck, out, route="xml")

    line = _pictures(out)[0].spPr.find(f"{{{A}}}ln")
    assert line is not None and line.get("w") == "38100"


def test_the_title_still_moves_into_the_new_layout(tmp_path: Path) -> None:
    """The freeze must not have cost the thing the rebuild is for: the title's
    words move across and its typeface is left to the master."""
    master, deck = _fixture(tmp_path)
    out = tmp_path / "out.pptx"

    result = rebuild(master, deck, out, route="xml")

    titles = [
        shape.text_frame.text
        for slide in Presentation(str(out)).slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text
    ]
    assert "Team & References" in titles
    assert not result.dropped
