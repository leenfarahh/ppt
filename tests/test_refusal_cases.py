"""Findings that used to be refused off one three-card slide, and a date that
was carried beside its slot instead of into it.

Off a real deck restyled onto a client master. Five of seven findings on the
slide came back "not taken", every one for a reason that did not hold:

- Card copy spilling a hair past its box was said to "run over" the card it
  sits on, so no box to grow into was offered.
- A bottom-anchored heading overflowing its TOP could not be grown, because
  growing only ever moved the bottom edge.
- A title was held back by the empty subtitle placeholder under it, which
  draws nothing.
- A paragraph in a tall bottom-anchored box had its last line drawn across the
  layout's footer rule, and nothing saw it, because the text never left its
  box.
- A stranded "Goals" was a paragraph of its own, pasted with a hard return at
  every line, and binding two words cannot bind one.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.extract import read_deck
from formatting_tool.extract.master_spec import derive_master_spec
from formatting_tool.linemetrics import ShapeKey, TextBounds
from formatting_tool.models import BrandGuidelines
from formatting_tool.rules import RuleContext
from formatting_tool.rules.space import TextCollisionRule, TextOverflowRule


class StubMetrics:
    def __init__(self, bounds=None, lines=None) -> None:
        self._bounds = bounds or {}
        self._lines = lines or {}

    @property
    def available(self) -> bool:
        return True

    def lines(self, key):
        return None

    def bounds(self, key):
        return self._bounds.get(key)

    def line_bounds(self, key):
        return self._lines.get(key)


def _spec():
    return derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())


def _found(rule, deck: Path) -> list:
    found = list(rule.check(RuleContext(deck=read_deck(deck), spec=_spec())))
    for issue in found:
        issue.id = issue.fingerprint()
    return found


def _apply(deck: Path, found: list, tmp_path: Path):
    from formatting_tool.apply import apply_fixes

    out = tmp_path / "fixed.pptx"
    result = apply_fixes(deck, found, out, selected=[i.id for i in found],
                         spec=_spec())
    return result, out


def _card_slide(tmp_path: Path, *, anchor=None, with_number=False):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.77),
                                  Inches(2.17), Inches(4.2), Inches(4.46))
    card.name = "Card"
    copy = slide.shapes.add_textbox(Inches(3.77), Inches(3.56), Inches(4.2),
                                    Inches(2.56))
    copy.name = "Card copy"
    copy.text_frame.word_wrap = True
    if anchor is not None:
        copy.text_frame.vertical_anchor = anchor
    copy.text_frame.text = "At ADFD, our people are the driving force " * 6
    copy.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    number = None
    if with_number:
        number = slide.shapes.add_textbox(Inches(6.4), Inches(5.94),
                                          Inches(1.57), Inches(0.71))
        number.name = "Number"
        number.text_frame.text = "01"
    path = tmp_path / "card.pptx"
    prs.save(str(path))
    return path, copy.shape_id, number.shape_id if number else None


def test_copy_spilling_onto_its_own_card_is_offered_room(tmp_path: Path) -> None:
    deck, copy_id, _ = _card_slide(tmp_path)
    metrics = StubMetrics({
        ShapeKey(1, copy_id): TextBounds(3.97, 3.56, 3.78, 2.64),  # to 6.20in
    })

    found = _found(TextOverflowRule(metrics=metrics), deck)

    assert len(found) == 1
    assert "running over" not in found[0].message
    assert found[0].expected.startswith("a box its copy fits")

    result, out = _apply(deck, found, tmp_path)
    assert len(result.applied) == 1, result.skipped
    assert "taller" in result.applied[0].detail


def test_a_bottom_anchored_box_grows_upward(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR

    deck, copy_id, _ = _card_slide(tmp_path, anchor=MSO_ANCHOR.BOTTOM)
    # Bottom-anchored copy overflows its top, not its bottom.
    metrics = StubMetrics({
        ShapeKey(1, copy_id): TextBounds(3.97, 3.40, 3.78, 2.72),
    })

    found = _found(TextOverflowRule(metrics=metrics), deck)
    result, out = _apply(deck, found, tmp_path)

    assert len(result.applied) == 1, result.skipped
    assert "upward" in result.applied[0].detail
    box = [s for s in Presentation(str(out)).slides[0].shapes
           if s.name == "Card copy"][0]
    # The bottom, where bottom-anchored copy sits, has not moved.
    assert (box.top + box.height) / 914400 == pytest.approx(6.12, abs=0.005)
    assert box.top / 914400 == pytest.approx(3.35, abs=0.01)


def test_a_line_that_misses_a_number_is_not_a_collision(tmp_path: Path) -> None:
    """The paragraph's rectangle reaches the number; its last line does not."""
    deck, copy_id, number_id = _card_slide(tmp_path, with_number=True)
    lines = [TextBounds(3.97, 3.56 + 0.24 * i, 3.78, 0.24) for i in range(10)]
    lines.append(TextBounds(3.97, 5.96, 3.26, 0.24))          # ends at 7.23in
    metrics = StubMetrics(
        {
            ShapeKey(1, copy_id): TextBounds(3.97, 3.56, 3.78, 2.64),
            ShapeKey(1, number_id): TextBounds(7.36, 5.99, 0.51, 0.60),
        },
        {ShapeKey(1, copy_id): lines},
    )

    assert _found(TextCollisionRule(metrics=metrics), deck) == []
    overflow = _found(TextOverflowRule(metrics=metrics), deck)
    assert overflow[0].expected.startswith("a box its copy fits")


def _paragraph_by_the_rule(tmp_path: Path):
    """A paragraph beside two cards, low and against the edge, over a rule."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    layout = prs.slide_layouts[5]                    # title only
    for placeholder in layout.placeholders:
        placeholder.left, placeholder.top = Inches(0.71), Inches(0.71)
        placeholder.width, placeholder.height = Inches(10.95), Inches(0.56)
    slide = prs.slides.add_slide(layout)
    # The footer rule belongs to the layout. python-pptx draws connectors on
    # slides only, so it is drawn here and moved across.
    rule = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.71),
                                      Inches(7.02), Inches(12.63), Inches(7.02))
    layout.shapes._spTree.append(rule._element)
    title = slide.shapes.title
    title.left, title.top = Inches(0.71), Inches(0.71)
    title.width, title.height = Inches(10.95), Inches(0.56)
    title.text = "Three point content"
    for i, left in enumerate((3.77, 8.43)):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left),
                                      Inches(2.17), Inches(4.2), Inches(4.46))
        card.name = f"Card {i + 1}"
    para = slide.shapes.add_textbox(Inches(0.0), Inches(2.65), Inches(2.75),
                                    Inches(4.46))
    para.name = "Context"
    para.text_frame.word_wrap = True
    para.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    para.text_frame.text = "This paragraph provides detailed context. " * 5
    path = tmp_path / "ruled.pptx"
    prs.save(str(path))
    return path, para.shape_id


def test_text_across_the_layouts_rule_moves_into_its_row(tmp_path: Path) -> None:
    from pptx import Presentation

    deck, para_id = _paragraph_by_the_rule(tmp_path)
    metrics = StubMetrics({
        ShapeKey(1, para_id): TextBounds(0.0, 5.01, 2.64, 2.10),  # to 7.11in
    })

    found = _found(TextCollisionRule(metrics=metrics), deck)

    assert len(found) == 1
    assert found[0].shape == "Context"
    assert "runs through its lines" in found[0].message
    # Onto the cards' top and the title's left edge, not a bare 0.11in nudge.
    assert "moved to 0.71, 2.17in" in found[0].expected

    result, out = _apply(deck, found, tmp_path)
    assert len(result.applied) == 1, result.skipped
    para = [s for s in Presentation(str(out)).slides[0].shapes
            if s.name == "Context"][0]
    # Off the slide edge and inside the frame. The second round may take it
    # further in, onto this test master's own safe margin.
    assert para.left / 914400 >= 0.70
    assert para.top / 914400 == pytest.approx(2.17, abs=0.01)


def test_a_rule_just_under_a_heading_is_not_through_it(tmp_path: Path) -> None:
    deck, para_id = _paragraph_by_the_rule(tmp_path)
    metrics = StubMetrics({
        # Ends 0.02in below the rule: touching, not crossed.
        ShapeKey(1, para_id): TextBounds(0.0, 5.01, 2.64, 2.03),
    })

    assert _found(TextCollisionRule(metrics=metrics), deck) == []


# --------------------------------------------------------------------------- #
# A stranded word that is a paragraph of its own
# --------------------------------------------------------------------------- #

def _hard_wrapped(tmp_path: Path, paragraphs: list[str]):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    box.name = "Copy"
    frame = box.text_frame
    frame.text = paragraphs[0]
    for text in paragraphs[1:]:
        frame.add_paragraph().text = text
    path = tmp_path / "wrapped.pptx"
    prs.save(str(path))
    return path


def _orphan(deck: Path, word: str):
    from formatting_tool.models import Category, Issue, Severity

    issue = Issue(
        rule_id="typography.orphan_widow", category=Category.TYPOGRAPHY,
        severity=Severity.WARNING,
        message=f"Last line is a single stranded word ('{word}').",
        slide=1, shape="Copy",
        shape_id=[s for s in read_deck(deck).slides[0].shapes
                  if s.name == "Copy"][0].shape_id,
        expected="more than 1 word(s) on the last line", found=word,
    )
    issue.id = issue.fingerprint()
    return issue


def test_a_hard_wrapped_sentence_is_rejoined(tmp_path: Path) -> None:
    from pptx import Presentation

    deck = _hard_wrapped(tmp_path, [
        "Our collaborations have been central to delivering",
        "transformative projects, in line with the United Nations Sustainable",
        "Goals",
    ])

    result, out = _apply(deck, [_orphan(deck, "Goals")], tmp_path)

    assert len(result.applied) == 1, result.skipped
    assert "rejoined 3 lines" in result.applied[0].detail
    box = Presentation(str(out)).slides[0].shapes[0]
    assert [p.text for p in box.text_frame.paragraphs] == [
        "Our collaborations have been central to delivering transformative "
        "projects, in line with the United Nations Sustainable Goals"
    ]


def test_a_list_of_short_items_is_not_merged(tmp_path: Path) -> None:
    deck = _hard_wrapped(tmp_path, ["Planning", "Delivery", "Review"])

    result, _ = _apply(deck, [_orphan(deck, "Review")], tmp_path)

    assert result.applied == []
    assert "only one word" in result.skipped[0].detail


def test_a_finished_sentence_is_not_merged(tmp_path: Path) -> None:
    deck = _hard_wrapped(tmp_path, [
        "Our collaborations have been central to delivering results.",
        "Thanks",
    ])

    result, _ = _apply(deck, [_orphan(deck, "Thanks")], tmp_path)

    assert result.applied == []


# --------------------------------------------------------------------------- #
# The date goes into the layout's date slot
# --------------------------------------------------------------------------- #

def test_a_date_goes_into_a_placeholder_prompted_for_one(tmp_path: Path) -> None:
    """The master's date slot is a text placeholder reading "DD/MM/YY"."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    from formatting_tool.rebuild.builder import claim_dates, _copy_text

    prs = Presentation()
    layout = prs.slide_layouts[0]                    # title slide
    subtitle = [p for p in layout.placeholders if p.placeholder_format.idx == 1][0]
    subtitle.text_frame.text = "DD/MM/YY"
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Advancing growth"
    loose = slide.shapes.add_textbox(Inches(0), Inches(5.8), Inches(3), Inches(0.4))
    loose.text_frame.text = "23-Sep-26"

    shapes = list(slide.shapes)
    pool = list(slide.placeholders)
    claims = claim_dates(shapes, pool)

    assert len(claims) == 1
    slot = claims[id(shapes[-1])]
    assert slot.placeholder_format.idx == 1
    assert slot not in pool
    assert _copy_text(shapes[-1], slot)
    assert slot.text_frame.text == "23-Sep-26"


def test_a_live_date_field_survives_the_copy(tmp_path: Path) -> None:
    pytest.importorskip("pptx")
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Inches

    from formatting_tool.rebuild.builder import _copy_text

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    source = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(3), Inches(1))
    target = slide.shapes.add_textbox(Inches(0), Inches(2), Inches(3), Inches(1))
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    p = source.text_frame.paragraphs[0]._p
    field = etree.SubElement(p, f"{{{a}}}fld", id="{F9F5D1D4-0000-0000-0000-000000000000}",
                             type="datetime1")
    rpr = etree.SubElement(field, f"{{{a}}}rPr", lang="en-US", sz="1600")
    etree.SubElement(rpr, f"{{{a}}}solidFill")
    etree.SubElement(field, f"{{{a}}}t").text = "23-Sep-26"
    target.text_frame.text = "x"

    assert _copy_text(source, target)

    fields = target.text_frame.paragraphs[0]._p.findall(f"{{{a}}}fld")
    assert len(fields) == 1 and fields[0].get("type") == "datetime1"
    # Emphasis would survive; size and colour are the layout's.
    assert fields[0].find(f"{{{a}}}rPr").get("sz") is None
    assert len(fields[0].find(f"{{{a}}}rPr")) == 0


def test_a_connector_joined_to_a_badge_is_not_a_collision(tmp_path: Path) -> None:
    """A process diagram: each number badge has a connector starting at its
    edge and running out to its label, straight through the number. That is
    the drawing. Moving the badge clear of it broke a diagram that was right."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5.71),
                               Inches(4.85), Inches(6.74), Inches(4.85))
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.71), Inches(4.51),
                                   Inches(0.68), Inches(0.68))
    badge.name = "Oval 20"
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    badge.text_frame.text = "03"
    path = tmp_path / "diagram.pptx"
    prs.save(str(path))
    metrics = StubMetrics({
        ShapeKey(1, badge.shape_id): TextBounds(5.90, 4.70, 0.30, 0.30),
    })

    assert _found(TextCollisionRule(metrics=metrics), path) == []


def test_a_cover_does_not_inherit_its_old_layouts_artwork() -> None:
    """The old cover's freeforms, rule and logo sit on the one layout the one
    cover uses, so repetition reads them as the slide's own. They are the old
    brand's cover design, and the new master brings a cover of its own."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    from formatting_tool.rebuild.pictures import inherit_artwork

    prs = Presentation()
    cover = prs.slide_layouts[0]                    # Title Slide
    content = prs.slide_layouts[5]                  # Title Only
    # Different artwork on each: the same shape on both would read as the
    # old brand's furniture and travel to neither.
    for layout, width in ((cover, 4), (content, 2)):
        scratch = prs.slides.add_slide(layout)
        art = scratch.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8), Inches(0),
                                       Inches(width), Inches(7))
        layout.shapes._spTree.append(art._element)
    for _ in range(len(prs.slides)):
        rid = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rid)
        prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
    prs.slides.add_slide(cover)
    prs.slides.add_slide(content)

    carried = inherit_artwork(prs)

    assert not any(line.startswith("slide 1:") for line in carried)
    assert any(line.startswith("slide 2:") for line in carried)
