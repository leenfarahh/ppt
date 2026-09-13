"""Arabic decks that read from the right, and the parts that still read left.

Two defects a deck carries after its direction has been set, both of which
survive `typography.rtl_not_set` because neither is the direction:

- a paragraph explicitly set flush left. `algn="l"` names an EDGE of the box,
  not a leading edge, so marking the paragraph right-to-left leaves the copy
  exactly where it was, against the margin the reader finishes at.
- a shape whose box starts at the English column. The copy inside it can be
  set flush right and the slide still reads from the wrong side of the page,
  which is what an Arabic deck built on an English master looks like. Pictures
  are the case nothing else here could see: they hold no text, and the
  alignment grid skips shapes that hold none.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.extract.master_spec import derive_master_spec
from formatting_tool.models import (
    BrandGuidelines,
    Category,
    DeckProfile,
    Geometry,
    Issue,
    LayoutProfile,
    ParagraphProfile,
    RunProfile,
    Severity,
    ShapeProfile,
    SlideProfile,
    Source,
    TableCell,
    TableProfile,
)
from formatting_tool.rules import RuleContext
from formatting_tool.rules.direction import RtlAlignmentRule, RtlLeadingEdgeRule

CANVAS_W, CANVAS_H = 13.333, 7.5

# The master's content column and its mirror. Every number below is one of
# these two: 0.92in is where an English reader starts, 12.41in is where an
# Arabic one does.
COLUMN = 0.92
MIRROR = round(CANVAS_W - COLUMN, 2)

AR = "التحول الصحي في نيوم"
EN = "Health Systems Operations"


# --------------------------------------------------------------------------- #
# Fixtures
# --------------------------------------------------------------------------- #

def _paragraph(text: str, alignment=None) -> ParagraphProfile:
    return ParagraphProfile(
        text=text,
        alignment=alignment,
        runs=[RunProfile(text=text)],
    )


def _shape(
    name: str,
    left: float,
    *,
    width: float = 4.0,
    top: float = 2.0,
    paragraphs=None,
    token=None,
    picture: bool = False,
    table: TableProfile | None = None,
) -> ShapeProfile:
    paragraphs = paragraphs or []
    return ShapeProfile(
        shape_id=abs(hash((name, left, top))) % 99999,
        name=name,
        shape_type="PICTURE (13)" if picture else "TEXT_BOX (17)",
        geometry=Geometry(left_in=left, top_in=top, width_in=width, height_in=1.0),
        placeholder_type=token,
        is_picture=picture,
        text="\n".join(p.text for p in paragraphs),
        paragraphs=paragraphs,
        table=table,
    )


def _layout(name: str, left: float, width: float, *, ps: bool = True) -> LayoutProfile:
    """One of the designer's layouts, with its content area stated.

    Stated twice when `ps` is on -- as a marked rectangle and as a placeholder
    -- which is what a master somebody has annotated looks like. With `ps` off
    it is the ordinary case: a finished master carrying placeholders and no
    annotation at all, which the rule has to be able to follow as drawn.
    """
    shapes = [
        _shape(f"{name} body", left, width=width, token="OBJECT (7)"),
        _shape(f"{name} page", 9.417, width=1.0, top=6.95, token="SLIDE_NUMBER (13)"),
    ]
    if ps:
        marked = _shape(f"{name} PS", left, width=width)
        marked.alt_text = "PS"
        shapes.insert(0, marked)
    return LayoutProfile(name=name, index=0, shapes=shapes)


def _master(*, symmetric: bool = True, ps: bool = True) -> DeckProfile:
    """A master with one content column, as the designer drew it.

    Symmetric is the normal case and the one the mirror depends on -- a master
    whose content area runs 0.92in to 12.41in declares both edges, so the
    mirror of its left column is a column it draws. `symmetric=False` is the
    master that stops short on the right, where there is no mirrored column to
    move anything onto and the rule has nothing to say.
    """
    right_width = (CANVAS_W - 2 * COLUMN) if symmetric else 6.0
    return DeckProfile(
        path="master.pptx",
        width_in=CANVAS_W,
        height_in=CANVAS_H,
        layouts=[_layout("content", COLUMN, right_width, ps=ps)],
    )


def _ctx(
    shapes,
    *,
    arabic: bool = True,
    master: DeckProfile | None = None,
    layout_name: str = "content",
):
    """A one-slide deck of the given shapes, read against the master.

    `arabic` decides what `DeckProfile.rtl` will answer, and it is measured off
    the copy rather than declared, so the slide carries a text box either way.
    It sits on no column of its own -- a fixture that decides the deck's
    direction should not also be one of the shapes under test.
    """
    voice = _shape("Copy", 5.0, top=5.0, paragraphs=[_paragraph(AR if arabic else EN)])
    deck = DeckProfile(
        path="messy.pptx",
        width_in=CANVAS_W,
        height_in=CANVAS_H,
        slides=[
            SlideProfile(number=1, layout_name=layout_name, shapes=[voice, *shapes])
        ],
    )
    spec = derive_master_spec(master or _master(), BrandGuidelines())
    return RuleContext(deck=deck, spec=spec)


def _found(rule, ctx) -> list[Issue]:
    return list(rule.check(ctx))


# --------------------------------------------------------------------------- #
# The master states both columns
# --------------------------------------------------------------------------- #

def test_the_mirror_of_the_content_column_is_a_column_the_master_draws() -> None:
    """The premise the whole geometry rule rests on. Without it there is
    nowhere to move a shape to that the master can be said to have drawn."""
    layout = _master().layouts[0]

    assert COLUMN in layout.declared_left_edges
    assert MIRROR in [round(edge, 2) for edge in layout.declared_right_edges]


def test_an_unannotated_master_is_followed_as_drawn() -> None:
    """Nothing is asked of the designer's file. A finished master states where
    content goes through its placeholders, and a rule that only worked on a
    master somebody had gone back and marked up would report nothing on most
    of them -- which is the same as not having the rule."""
    ctx = _ctx(
        [_shape("Picture 3", COLUMN, width=4.0, picture=True)],
        master=_master(ps=False),
    )

    found = _found(RtlLeadingEdgeRule(), ctx)

    assert len(found) == 1
    assert found[0].expected == f"right {MIRROR:.2f}in"


def test_a_slide_is_measured_against_the_layout_it_is_built_on() -> None:
    """Layouts declare different columns, and the slide says which one it is
    on. A shape sitting on a column some other layout offers is not on a
    column of its own layout, and moving it to that layout's mirror would put
    it somewhere nothing asked for."""
    master = DeckProfile(
        path="master.pptx",
        width_in=CANVAS_W,
        height_in=CANVAS_H,
        layouts=[
            _layout("content", COLUMN, CANVAS_W - 2 * COLUMN),
            _layout("inset", 2.0, CANVAS_W - 4.0),
        ],
    )
    ctx = _ctx(
        [_shape("Picture 3", COLUMN, width=4.0, picture=True)],
        master=master,
        layout_name="inset",
    )

    assert _found(RtlLeadingEdgeRule(), ctx) == []


# --------------------------------------------------------------------------- #
# Paragraphs set flush left
# --------------------------------------------------------------------------- #

def test_arabic_set_flush_left_is_reported() -> None:
    ctx = _ctx([_shape("Body", COLUMN, paragraphs=[_paragraph(AR, "LEFT (1)")])])

    found = _found(RtlAlignmentRule(), ctx)

    assert len(found) == 1
    assert found[0].expected == "right aligned on every Arabic paragraph"


def test_an_unstated_alignment_is_not() -> None:
    """It inherits, and what a right-to-left paragraph inherits is flush
    right. Reporting it would fault a paragraph for saying nothing."""
    ctx = _ctx([_shape("Body", COLUMN, paragraphs=[_paragraph(AR)])])

    assert _found(RtlAlignmentRule(), ctx) == []


def test_centred_and_justified_arabic_are_left_alone() -> None:
    """Neither has a side: centred copy reads the same either way, and
    justified copy justifies to whichever direction the paragraph runs."""
    ctx = _ctx([
        _shape("Centred", COLUMN, paragraphs=[_paragraph(AR, "CENTER (2)")]),
        _shape("Justified", COLUMN, top=3.5, paragraphs=[_paragraph(AR, "JUSTIFY (4)")]),
    ])

    assert _found(RtlAlignmentRule(), ctx) == []


def test_english_set_flush_left_is_not_a_defect() -> None:
    """The rule is per paragraph, not per deck: an English line in an Arabic
    deck is against the margin it reads from."""
    ctx = _ctx([_shape("Body", COLUMN, paragraphs=[_paragraph(EN, "LEFT (1)")])])

    assert _found(RtlAlignmentRule(), ctx) == []


def test_an_arabic_paragraph_in_an_english_deck_is_still_reported() -> None:
    """A bilingual deck is the normal case here. An Arabic pull-quote on an
    otherwise English slide hangs off its own margin just the same."""
    ctx = _ctx(
        [_shape("Quote", COLUMN, paragraphs=[_paragraph(AR, "LEFT (1)")])],
        arabic=False,
    )

    assert len(_found(RtlAlignmentRule(), ctx)) == 1


def test_table_cells_count_and_the_header_row_does_not() -> None:
    """Most of the Arabic in a consulting deck is in table cells. Row 0 is
    left out because `table.header_alignment` already reports it, and a
    designer told to fix one row twice fixes it once and doubts the report."""
    cells = [
        TableCell(row=0, column=0, text=AR, paragraphs=[_paragraph(AR, "LEFT (1)")]),
        TableCell(row=1, column=0, text=AR, paragraphs=[_paragraph(AR, "LEFT (1)")]),
        TableCell(row=2, column=0, text=AR, paragraphs=[_paragraph(AR, "LEFT (1)")]),
    ]
    table = TableProfile(rows=3, columns=1, cells=cells)
    ctx = _ctx([_shape("Table", COLUMN, table=table)])

    found = _found(RtlAlignmentRule(), ctx)

    assert len(found) == 1
    assert "2 of 2" in (found[0].found or "")      # the two body rows, not three


# --------------------------------------------------------------------------- #
# Shapes that lead from the left
# --------------------------------------------------------------------------- #

def test_an_image_on_the_english_column_is_reported_with_its_mirror() -> None:
    """The case no other rule here can see: a picture holds no text, and the
    alignment grid reads copy."""
    ctx = _ctx([_shape("Picture 3", COLUMN, width=4.0, picture=True)])

    found = _found(RtlLeadingEdgeRule(), ctx)

    assert len(found) == 1
    assert found[0].expected == f"right {MIRROR:.2f}in"
    assert "image" in found[0].message


def test_an_english_deck_is_not_turned_round() -> None:
    ctx = _ctx([_shape("Picture 3", COLUMN, width=4.0, picture=True)], arabic=False)

    assert _found(RtlLeadingEdgeRule(), ctx) == []


def test_a_shape_already_ending_on_a_column_is_left_alone() -> None:
    """Both a full-width shape and one already moved across: each ends on a
    column the master draws, which is the whole test for leading from the
    right."""
    ctx = _ctx([
        _shape("Full width", COLUMN, width=CANVAS_W - 2 * COLUMN),
        _shape("Moved", MIRROR - 4.0, top=3.5, width=4.0, picture=True),
    ])

    assert _found(RtlLeadingEdgeRule(), ctx) == []


def test_a_shape_on_no_column_at_all_is_left_alone() -> None:
    """Placed somewhere of its own, so there is no evidence it was placed
    against the English frame and nothing to mirror it onto."""
    ctx = _ctx([_shape("Loose", 3.17, width=4.0, picture=True)])

    assert _found(RtlLeadingEdgeRule(), ctx) == []


def test_nothing_is_said_when_the_master_draws_no_mirrored_column() -> None:
    """An asymmetric master states a right-hand frame that is genuinely
    different, and inventing a position there would be this layer guessing."""
    ctx = _ctx(
        [_shape("Picture 3", COLUMN, width=4.0, picture=True)],
        master=_master(symmetric=False),
    )

    assert _found(RtlLeadingEdgeRule(), ctx) == []


def test_chrome_stays_where_the_master_puts_it() -> None:
    """A footer, a page number and a date live where the master puts them.
    Which side of an Arabic page pagination belongs on is a decision this
    layer has no evidence to overturn."""
    ctx = _ctx([
        _shape("Page 1", COLUMN, top=6.95, width=1.0, token="SLIDE_NUMBER (13)"),
        _shape("Footer 1", COLUMN, top=6.95, width=3.0, token="FOOTER (15)"),
    ])

    assert _found(RtlLeadingEdgeRule(), ctx) == []


def test_a_deck_of_a_different_page_size_is_not_measured() -> None:
    """Two page sizes are two coordinate spaces, and mirroring a 4:3 slide
    about a 16:9 master's centre puts the shape somewhere neither asked for."""
    ctx = _ctx([_shape("Picture 3", COLUMN, width=4.0, picture=True)])
    ctx.deck.width_in = 10.0

    assert _found(RtlLeadingEdgeRule(), ctx) == []


# --------------------------------------------------------------------------- #
# The fixes
# --------------------------------------------------------------------------- #

def _issue(rule_id: str, **kwargs) -> Issue:
    issue = Issue(
        category=kwargs.pop("category", Category.SPACE),
        severity=kwargs.pop("severity", Severity.WARNING),
        message=kwargs.pop("message", f"{rule_id} finding"),
        source=Source.RULE,
        rule_id=rule_id,
        deck="messy.pptx",
        **kwargs,
    )
    issue.id = issue.fingerprint()
    return issue


def test_the_fix_sets_the_arabic_paragraphs_right_aligned(tmp_path: Path) -> None:
    """Nothing moves and no words change: the copy swaps which edge of its own
    box it sits against."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    from formatting_tool.apply import apply_fixes

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(COLUMN), Inches(2), Inches(4), Inches(1))
    box.name = "Body"
    frame = box.text_frame
    frame.text = AR
    frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    english = frame.add_paragraph()
    english.text = EN
    english.alignment = PP_ALIGN.LEFT
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    out = tmp_path / "fixed.pptx"
    result = apply_fixes(
        deck,
        [_issue(
            "typography.rtl_alignment",
            category=Category.TYPOGRAPHY,
            slide=1,
            shape="Body",
            shape_id=box.shape_id,
            expected="right aligned on every Arabic paragraph",
        )],
        out,
        selected=None,
    )

    assert len(result.applied) == 1
    fixed = Presentation(str(out)).slides[0].shapes[0].text_frame.paragraphs
    assert fixed[0].alignment == PP_ALIGN.RIGHT     # the Arabic line
    assert fixed[1].alignment == PP_ALIGN.LEFT      # the English one, untouched


def test_the_fix_moves_an_image_onto_the_mirrored_column(tmp_path: Path) -> None:
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    from formatting_tool.apply import apply_fixes

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(CANVAS_W), Inches(CANVAS_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(COLUMN), Inches(2), Inches(4), Inches(1))
    box.name = "Picture 3"
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    out = tmp_path / "fixed.pptx"
    result = apply_fixes(
        deck,
        [_issue(
            "space.rtl_leading_edge",
            slide=1,
            shape="Picture 3",
            shape_id=box.shape_id,
            expected=f"right {MIRROR:.2f}in",
            found=f"right {COLUMN + 4.0:.2f}in",
        )],
        out,
        selected=None,
    )

    assert len(result.applied) == 1
    moved = Presentation(str(out)).slides[0].shapes[0]
    assert round((moved.left + moved.width) / 914400, 2) == MIRROR
    assert round(moved.width / 914400, 2) == 4.0    # the box keeps its size
