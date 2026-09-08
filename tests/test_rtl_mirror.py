"""Turning a rebuilt deck round for copy that reads right to left.

Applying an English master to an Arabic deck and stopping there gets the type
and the colour right and the layout backwards: the master puts its title
0.92in from the left because that is where an English reader starts.

What turns is the FRAME, and getting that wrong was the first thing this did.
Mirroring every shape on the slide moved an Arabic deck's own copy -- already
set flush right by whoever wrote it -- over to the left, while correctly
bringing the master's title across. The deck was right and the master was
wrong, and turning both fixed one and broke the other.

The second mistake was turning the frame twice: the layout moved, and then the
placeholder that INHERITS from the layout read its new position and moved
again, landing back where it started.
"""

from __future__ import annotations

from pathlib import Path

import pytest

AR = "التحول الصحي في نيوم"
EN = "Health Systems Operations"
WIDTH_IN = 13.333


def _deck(tmp_path: Path):
    """A master-shaped deck: a layout with a left-hand title, content on it."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH_IN), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])   # Title Only
    slide.shapes.title.text = AR
    # The author's own copy, already set flush right the way an Arabic deck is.
    box = slide.shapes.add_textbox(
        Inches(9.40), Inches(2.0), Inches(3.0), Inches(0.6)
    )
    box.name = "Authored"
    box.text_frame.text = AR
    return prs, slide


def _right(shape) -> float:
    return round((shape.left + shape.width) / 914400, 2)


def _left(shape) -> float:
    return round(shape.left / 914400, 2)


def _find(container, name):
    return next(s for s in container.shapes if s.name == name)


# --------------------------------------------------------------------------- #
# Frame, not content
# --------------------------------------------------------------------------- #

def test_the_authors_own_shapes_do_not_move(tmp_path: Path) -> None:
    """An Arabic deck's author has already decided which side their copy sits
    on. The rebuild has no business arguing."""
    from formatting_tool.rebuild.rtl import mirror_presentation

    prs, slide = _deck(tmp_path)
    before = _left(_find(slide, "Authored"))

    mirror_presentation(prs)

    assert _left(_find(slide, "Authored")) == before


def test_the_layout_turns_round(tmp_path: Path) -> None:
    from formatting_tool.rebuild.rtl import mirror_presentation

    prs, slide = _deck(tmp_path)
    layout = slide.slide_layout
    title = _find(layout, "Title 1")
    was_left, was_width = _left(title), title.width

    mirror_presentation(prs)

    turned = _find(layout, "Title 1")
    assert _right(turned) == round(WIDTH_IN - was_left, 2)
    assert turned.width == was_width        # moved, not resized


def test_an_inheriting_placeholder_turns_once(tmp_path: Path) -> None:
    """The layout moves and the placeholder reads its new position. Turning
    that too puts it back where it started, which is what happened."""
    from formatting_tool.rebuild.rtl import mirror_presentation

    prs, slide = _deck(tmp_path)
    layout_title = _find(slide.slide_layout, "Title 1")
    was_left = _left(layout_title)

    mirror_presentation(prs)

    slide_title = slide.shapes.title
    assert _right(slide_title) == round(WIDTH_IN - was_left, 2)
    assert _left(slide_title) != was_left


def test_a_placeholder_with_its_own_position_turns_too(tmp_path: Path) -> None:
    """One that states where it sits is not inheriting, so the layout's move
    does not carry it and it has to be moved itself."""
    from pptx.util import Inches

    from formatting_tool.rebuild.rtl import mirror_presentation

    prs, slide = _deck(tmp_path)
    title = slide.shapes.title
    title.left, title.width = Inches(1.0), Inches(4.0)

    mirror_presentation(prs)

    assert _right(slide.shapes.title) == round(WIDTH_IN - 1.0, 2)


def test_nothing_is_resized_or_moved_vertically(tmp_path: Path) -> None:
    """A mirror about the vertical centre changes nothing about how high a
    shape sits or how tall it is."""
    from formatting_tool.rebuild.rtl import mirror_presentation

    prs, slide = _deck(tmp_path)
    layout = slide.slide_layout
    before = {s.name: (s.top, s.width, s.height) for s in layout.shapes}

    mirror_presentation(prs)

    after = {s.name: (s.top, s.width, s.height) for s in layout.shapes}
    assert after == before


# --------------------------------------------------------------------------- #
# Text turns with the page
# --------------------------------------------------------------------------- #

def test_flush_left_becomes_flush_right(tmp_path: Path) -> None:
    from pptx.enum.text import PP_ALIGN

    from formatting_tool.rebuild.rtl import mirror_presentation

    prs, slide = _deck(tmp_path)
    paragraph = _find(slide, "Authored").text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT

    mirror_presentation(prs)

    assert paragraph.alignment == PP_ALIGN.RIGHT


def test_centred_copy_has_no_side(tmp_path: Path) -> None:
    from pptx.enum.text import PP_ALIGN

    from formatting_tool.rebuild.rtl import mirror_presentation

    prs, slide = _deck(tmp_path)
    paragraph = _find(slide, "Authored").text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER

    mirror_presentation(prs)

    assert paragraph.alignment == PP_ALIGN.CENTER


def test_alignment_that_says_nothing_is_left_to_inherit(tmp_path: Path) -> None:
    """A right-to-left paragraph inherits flush right, so what it inherits is
    already correct. Setting it here would bake in a value the layout is
    entitled to change."""
    from formatting_tool.rebuild.rtl import mirror_presentation

    prs, slide = _deck(tmp_path)
    paragraph = _find(slide, "Authored").text_frame.paragraphs[0]
    assert paragraph.alignment is None

    mirror_presentation(prs)

    assert paragraph.alignment is None


def test_arabic_paragraphs_come_out_marked_right_to_left(tmp_path: Path) -> None:
    from formatting_tool.rebuild.rtl import mirror_presentation

    prs, slide = _deck(tmp_path)
    mirror_presentation(prs)

    paragraph = _find(slide, "Authored").text_frame.paragraphs[0]
    properties = paragraph._p.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}pPr"
    )
    assert properties is not None and properties.get("rtl") == "1"


# --------------------------------------------------------------------------- #
# When it happens at all
# --------------------------------------------------------------------------- #

def test_an_english_deck_is_not_turned_round(tmp_path: Path) -> None:
    """Decided from the deck rather than asked for, and turning an English
    deck round would be the defect."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    from formatting_tool.extract import read_deck

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH_IN), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = EN
    path = tmp_path / "english.pptx"
    prs.save(str(path))

    assert not read_deck(path).rtl


def test_an_arabic_deck_is(tmp_path: Path) -> None:
    from formatting_tool.extract import read_deck

    prs, _slide = _deck(tmp_path)
    path = tmp_path / "arabic.pptx"
    prs.save(str(path))

    assert read_deck(path).rtl
