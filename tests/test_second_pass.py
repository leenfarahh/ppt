"""Measuring the deck that was written, not the one that was measured.

The report a designer ticks describes the deck as it arrived. By the time the
fixes and the master have been applied it is a different file: type at the
master's size instead of the deck's, a box that no longer shrinks its text to
fit, a shape that no longer clears its neighbour. None of that is in the
report, because none of it was true when the report was written.

On a real deck the second pass came back with eleven findings the first pass
could not have had, among them a title overlapping the subtitle by 4 square
inches -- which is what the designer was looking at when they asked for this.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.apply import apply_fixes
from formatting_tool.models import BrandGuidelines, Category, Issue, Severity, Source


def _issue(rule_id: str, slide=None, shape=None) -> Issue:
    issue = Issue(
        category=Category.SPACE,
        severity=Severity.ERROR,
        message=f"{rule_id} finding",
        source=Source.RULE,
        rule_id=rule_id,
        slide=slide,
        shape=shape,
        deck="messy.pptx",
    )
    issue.id = issue.fingerprint()
    return issue


def _spec():
    from formatting_tool.extract import read_deck
    from formatting_tool.extract.master_spec import derive_master_spec

    return derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())


def _deck(tmp_path: Path) -> Path:
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.name = "TextBox 1"
    box.text_frame.text = "NEOM's Vision "
    path = tmp_path / "messy.pptx"
    prs.save(str(path))
    return path


# --------------------------------------------------------------------------- #
# It runs, and only when it can
# --------------------------------------------------------------------------- #

def test_the_output_is_measured_after_the_fixes_land(tmp_path: Path) -> None:
    deck = _deck(tmp_path)
    issue = _issue("typography.whitespace", slide=1, shape="TextBox 1")

    result = apply_fixes(
        deck, [issue], tmp_path / "out.pptx", selected=[issue.id], spec=_spec()
    )

    assert result.rechecked
    assert result.recheck is not None


def test_without_a_master_there_is_nothing_to_measure_against(
    tmp_path: Path,
) -> None:
    """Honest rather than hopeful: the rules compare a deck to a master, and
    there is no master here."""
    deck = _deck(tmp_path)
    issue = _issue("typography.whitespace", slide=1, shape="TextBox 1")

    result = apply_fixes(deck, [issue], tmp_path / "out.pptx", selected=[issue.id])

    assert not result.rechecked
    assert result.recheck == []


# --------------------------------------------------------------------------- #
# What is new versus what was inherited
# --------------------------------------------------------------------------- #

def test_a_finding_the_deck_already_had_is_not_introduced() -> None:
    from formatting_tool.apply.applier import ApplyResult

    old = _issue("space.overlap", slide=2, shape="Title 1")
    result = ApplyResult(deck="messy.pptx", output=Path("out.pptx"))
    result.before = [old]
    result.recheck = [_issue("space.overlap", slide=2, shape="Title 1")]

    assert result.introduced == []


def test_a_finding_only_the_output_has_is_introduced() -> None:
    from formatting_tool.apply.applier import ApplyResult

    result = ApplyResult(deck="messy.pptx", output=Path("out.pptx"))
    result.before = [_issue("space.overlap", slide=2, shape="Title 1")]
    new = _issue("space.overlap", slide=4, shape="Title 1")
    result.recheck = [result.before[0], new]

    assert result.introduced == [new]


def test_the_same_defect_measured_differently_is_still_the_same_defect() -> None:
    """Matched on rule and place, not on id. An id carries the message, and a
    message carrying a measurement changes when the measurement does -- so an
    overlap that grew from 2.79 to 3.10 sq in would read as a new finding
    rather than the same one made worse."""
    from formatting_tool.apply.applier import ApplyResult

    before = _issue("space.overlap", slide=2, shape="Title 1")
    before.message = "overlap by 2.79 sq in"
    after = _issue("space.overlap", slide=2, shape="Title 1")
    after.message = "overlap by 3.10 sq in"
    assert before.fingerprint() != after.fingerprint()

    result = ApplyResult(deck="messy.pptx", output=Path("out.pptx"))
    result.before = [before]
    result.recheck = [after]

    assert result.introduced == []


def test_a_run_that_fixed_nothing_still_reports_the_state_of_the_file(
    tmp_path: Path,
) -> None:
    """The point is the deck being sent, not the diff. A run where every fix
    was refused still has to say what the file measures."""
    deck = _deck(tmp_path)
    issue = _issue("space.overlap", slide=1, shape="TextBox 1")   # no fixer

    result = apply_fixes(
        deck, [issue], tmp_path / "out.pptx", selected=[issue.id], spec=_spec()
    )

    assert not result.applied
    assert result.rechecked
