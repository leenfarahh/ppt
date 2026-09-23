"""What each shape IS, and what the restyle is then allowed to do with it.

THE THREE DEFECTS THIS ANSWERS, all from one real consulting deck restyled
onto a client master:

  - "Source: Oxford Economics, Team analysis" came out of the foot of the page
    and printed under the title, on top of the subtitle. It is a full-width
    box, so it measures exactly like a standfirst, and nothing in a .pptx says
    "this is small print".
  - The subtitle region was filled from whichever full-width box happened to
    be drawn over it, which on a chart slide is the chart's own caption.
  - A chevron banner carrying "1. What is the new ambition?" lost its chevron:
    the words were claimed into a body region and the shape they came out of
    was deleted with them.

Each is invisible in the file and obvious in a picture, which is why the
answer is a model reading the render (`ai.roles`) with a deterministic reading
of small print behind it for the runs where there is no model.
"""

from __future__ import annotations

import pytest

from formatting_tool.ai.roles import (
    PROTECTED,
    ROLES,
    ShapeRole,
    SlideRoles,
    normalize_copy,
    roles_from_response,
)


# --------------------------------------------------------------------------- #
# Reading the answer
# --------------------------------------------------------------------------- #

class _Geometry:
    def __init__(self, left, top, width, height):
        self.left_in, self.top_in = left, top
        self.width_in, self.height_in = width, height


class _Shape:
    def __init__(self, shape_id, name, text="", box=(0.0, 0.0, 1.0, 1.0)):
        self.shape_id, self.name, self.text = shape_id, name, text
        self.geometry = _Geometry(*box)


class _Listed:
    def __init__(self, shape, path):
        self.shape, self.path = shape, path

    @property
    def name(self):
        return self.shape.name

    @property
    def shape_id(self):
        return self.shape.shape_id


def _refs(*shapes):
    return {
        f"s{i}": _Listed(shape, (i,))
        for i, shape in enumerate(shapes, start=1)
    }


def test_a_word_outside_the_vocabulary_reads_as_other() -> None:
    """Whitelisted on the way back rather than trusted. A word nobody defined
    has to read as unlabelled, not as a new behaviour in the restyle."""
    refs = _refs(_Shape(4, "Rectangle 4"))
    got = roles_from_response({"shapes": [{"shape": "s1", "role": "banner"}]},
                              1, refs)
    assert [entry.role for entry in got.shapes] == ["other"]
    assert "banner" not in ROLES


def test_a_ref_that_was_never_sent_is_dropped() -> None:
    """The model having invented a shape. The nearest real shape to an
    invented one is a different shape on somebody's client deck."""
    refs = _refs(_Shape(4, "Rectangle 4"))
    got = roles_from_response(
        {"shapes": [{"shape": "s9", "role": "source"}]}, 1, refs
    )
    assert got.shapes == () and got.reviewed


def test_a_second_answer_for_one_shape_is_dropped() -> None:
    refs = _refs(_Shape(4, "Rectangle 4"))
    got = roles_from_response(
        {"shapes": [{"shape": "s1", "role": "source"},
                    {"shape": "s1", "role": "body"}]}, 1, refs
    )
    assert [entry.role for entry in got.shapes] == ["source"]


def test_the_copy_is_carried_as_an_address() -> None:
    """THE ONLY ADDRESS THAT SURVIVES THE RESTYLE. These roles are read off the
    deck as it arrived, and the restyle then hands the file to PowerPoint,
    which duplicates every slide and renumbers any id it finds twice."""
    refs = _refs(_Shape(4, "Rectangle 4", text="Source:  Oxford\nEconomics"))
    got = roles_from_response(
        {"shapes": [{"shape": "s1", "role": "source"}]}, 1, refs
    )
    assert got.shapes[0].text == "source: oxford economics"
    assert got.protected_texts == {"source: oxford economics"}


def test_normalising_is_one_function_for_both_sides() -> None:
    """A round trip through PowerPoint turns a soft return into a line feed
    and back, and two stages disagreeing about that is a protection that
    silently stops working."""
    assert normalize_copy("  Source:\r\n Oxford  ") == normalize_copy(
        "source: oxford"
    )


# --------------------------------------------------------------------------- #
# What the answer protects
# --------------------------------------------------------------------------- #

def test_the_protected_roles_are_the_ones_that_must_not_move() -> None:
    """`source` is protected and `body` is not, and that is the distinction
    the whole module draws. A body box SHOULD move into the region the master
    gives it -- that is what applying a master is."""
    assert PROTECTED == {"source", "chart", "decoration"}
    assert "body" not in PROTECTED and "title" not in PROTECTED


def test_a_shape_inside_a_chart_is_protected_though_nobody_named_it() -> None:
    """A chart drawn as forty shapes does not come back as forty answers every
    time, and a chart with thirty-eight parts protected and two free to be
    levelled is a chart this tool is allowed to rewrite."""
    roles = SlideRoles(
        slide=1,
        shapes=(ShapeRole(ref="s1", role="chart", shape_id=5,
                          box=(1.0, 1.0, 6.0, 5.0)),),
        reviewed=True,
    )
    assert roles.protects(5)                              # named
    assert roles.protects(77, (1.5, 1.5, 2.0, 2.0))       # a label inside it
    assert not roles.protects(78, (8.0, 1.0, 9.0, 2.0))   # a caption beside it


def test_a_shape_half_in_a_chart_is_not_inside_it() -> None:
    roles = SlideRoles(
        slide=1,
        shapes=(ShapeRole(ref="s1", role="chart", shape_id=5,
                          box=(1.0, 1.0, 5.0, 5.0)),),
        reviewed=True,
    )
    assert not roles.inside_a_chart((4.0, 2.0, 7.0, 3.0))


def test_two_subtitles_are_no_subtitle() -> None:
    """The word is defined as at most one, so two of them is the model
    disagreeing with itself -- and filling the region from either is a coin
    toss that puts a heading at the top of somebody's page."""
    roles = SlideRoles(
        slide=1,
        shapes=(
            ShapeRole(ref="s1", role="subtitle", shape_id=1, text="one"),
            ShapeRole(ref="s2", role="subtitle", shape_id=2, text="two"),
        ),
        reviewed=True,
    )
    assert roles.subtitle is None


def test_one_subtitle_is_the_subtitle() -> None:
    roles = SlideRoles(
        slide=1,
        shapes=(ShapeRole(ref="s1", role="subtitle", shape_id=1, text="a kicker"),),
        reviewed=True,
    )
    assert roles.subtitle.text == "a kicker"


# --------------------------------------------------------------------------- #
# The deterministic reading, for the runs with no model
# --------------------------------------------------------------------------- #

pytest.importorskip("pptx")

from formatting_tool.rebuild.pictures import reads_as_a_note      # noqa: E402


class _Run:
    def __init__(self, size=None):
        class _Font:
            pass
        self.font = _Font()
        self.font.size = size


class _Paragraph:
    def __init__(self, runs):
        self.runs = runs


class _TextFrame:
    def __init__(self, text, runs=()):
        self.text = text
        self.paragraphs = [_Paragraph(list(runs))]


class _Box:
    def __init__(self, text, top_in=6.9, runs=()):
        self.has_text_frame = True
        self.text_frame = _TextFrame(text, runs)
        self.top = int(top_in * 914400)


CANVAS = (13.333, 7.5)


def test_a_source_line_is_read_as_small_print() -> None:
    assert reads_as_a_note(_Box("Source: Oxford Economics, Team analysis"), CANVAS)


def test_the_opener_is_enough_wherever_it_sits() -> None:
    """"Source:" is not a sentence anybody starts a body paragraph with, so
    its position does not have to agree before it counts."""
    assert reads_as_a_note(_Box("Note: figures are nominal", top_in=1.0), CANVAS)


def test_a_footnote_marker_before_the_opener_is_stepped_over() -> None:
    assert reads_as_a_note(_Box("* Source: team analysis"), CANVAS)


def test_small_type_at_the_foot_of_the_page_counts() -> None:
    """The methodology line that opens with no marker at all."""
    from pptx.util import Pt

    box = _Box("Figures exclude intra-group transfers", runs=[_Run(Pt(8))])
    assert reads_as_a_note(box, CANVAS)


def test_large_type_at_the_foot_of_the_page_does_not() -> None:
    """A closing statement somebody set large on purpose is not small print,
    and swallowing it would leave a region empty and a line stranded."""
    from pptx.util import Pt

    box = _Box("Thank you", runs=[_Run(Pt(28))])
    assert not reads_as_a_note(box, CANVAS)


def test_small_type_high_on_the_page_does_not() -> None:
    """A caption is small and is not a footnote; either signal on its own is
    ordinary, which is why both are required."""
    from pptx.util import Pt

    box = _Box("Manufacturing value add, 2020-2025",
               top_in=1.4, runs=[_Run(Pt(9))])
    assert not reads_as_a_note(box, CANVAS)


def test_an_empty_box_is_not_a_note() -> None:
    assert not reads_as_a_note(_Box("   "), CANVAS)
