"""Header rows, in real tables and in components drawn to look like them.

Until the extractor read tables, a table was one opaque rectangle to the whole
tool: a seven-column, forty-two-cell table came back as a single 12.27 x 5.23in
box, so nothing here could see a heading, a column or a row inside one.

The defect off a real deck, on that table: two headings set on one line and
left aligned, five set on two lines and centred. The header row reads as ragged
and no two column labels start at the same place.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.extract.master_spec import derive_master_spec
from formatting_tool.models import (
    BrandGuidelines,
    DeckProfile,
    Geometry,
    ParagraphProfile,
    RunProfile,
    ShapeProfile,
    SlideProfile,
    TableCell,
    TableProfile,
)
from formatting_tool.rules import RuleContext
from formatting_tool.rules.tables import (
    TableHeaderAlignmentRule,
    TableHeaderRowsRule,
)

CANVAS_W, CANVAS_H = 13.333, 7.5
HEADER_FILL = "433626"
BODY_FILL = None


def _cell(row, column, text, *, align="LEFT (1)", fill=HEADER_FILL) -> TableCell:
    parts = text.split("\n")
    return TableCell(
        row=row,
        column=column,
        text=text,
        paragraphs=[
            ParagraphProfile(text=p, alignment=align, runs=[RunProfile(text=p)])
            for p in parts
        ],
        geometry=Geometry(
            left_in=0.5 + column * 1.7, top_in=1.5 + row * 0.8,
            width_in=1.7, height_in=0.8,
        ),
        fill_hex=fill,
    )


def _table(header: list[tuple], *, body_fill=BODY_FILL, flagged=False) -> ShapeProfile:
    """A two-row table: the given header cells over a plain body row.

    `header` is (text, alignment, fill) per column.
    """
    cells = [
        _cell(0, i, text, align=align, fill=fill)
        for i, (text, align, fill) in enumerate(header)
    ]
    cells += [
        _cell(1, i, "body", align=None, fill=body_fill)
        for i in range(len(header))
    ]
    return ShapeProfile(
        shape_id=17,
        name="Table 17",
        shape_type="TABLE (19)",
        geometry=Geometry(left_in=0.5, top_in=1.5,
                          width_in=1.7 * len(header), height_in=1.6),
        table=TableProfile(
            rows=2,
            columns=len(header),
            row_heights_in=[0.8, 0.8],
            column_widths_in=[1.7] * len(header),
            cells=cells,
            first_row_header=flagged,
        ),
    )


def _copy(text: str, n: int = 3) -> list[ShapeProfile]:
    """Text-bearing shapes, which is what `DeckProfile.rtl` counts."""
    return [
        ShapeProfile(
            shape_id=100 + i,
            name=f"Body {i}",
            shape_type="TEXT_BOX (17)",
            geometry=Geometry(left_in=0.5, top_in=4.0 + i * 0.5,
                              width_in=4.0, height_in=0.4),
            text=text,
            paragraphs=[ParagraphProfile(text=text, runs=[RunProfile(text=text)])],
        )
        for i in range(n)
    ]


def _ctx(shapes, *, hidden=False) -> RuleContext:
    master = DeckProfile(path="master.pptx", width_in=CANVAS_W, height_in=CANVAS_H)
    deck = DeckProfile(
        path="messy.pptx", width_in=CANVAS_W, height_in=CANVAS_H,
        slides=[SlideProfile(number=1, shapes=shapes, hidden=hidden)],
    )
    return RuleContext(deck=deck, spec=derive_master_spec(master, BrandGuidelines()))


def _rows(ctx):
    return list(TableHeaderRowsRule().check(ctx))


def _align(ctx):
    return list(TableHeaderAlignmentRule().check(ctx))


# --------------------------------------------------------------------------- #
# The defect off the real deck
# --------------------------------------------------------------------------- #

REAL = [
    ("Cell chemistry", "LEFT (1)", HEADER_FILL),
    ("Demand drivers", "LEFT (1)", HEADER_FILL),
    ("Estimated supply gap by 2030\n(over/undersupply)", "CENTER (2)", HEADER_FILL),
    ("Captured by China\n(% of capacity)", "CENTER (2)", HEADER_FILL),
    ("Demand CAGR\n(2025-30)", "CENTER (2)", HEADER_FILL),
    ("Capacity CAGR\n(2025-30)", "CENTER (2)", HEADER_FILL),
    ("EBITDA margins\n(top players)", "CENTER (2)", HEADER_FILL),
]


def test_headings_on_more_than_one_line_are_reported() -> None:
    issues = _rows(_ctx([_table(REAL)]))

    assert len(issues) == 1
    assert "5 of 7 headings" in issues[0].message
    assert issues[0].shape == "Table 17"
    assert issues[0].expected == "one line per heading"
    assert issues[0].found == "up to 2 lines"


def test_headings_off_the_leading_edge_are_reported() -> None:
    issues = _align(_ctx([_table(REAL)]))

    assert len(issues) == 1
    assert "5 of 7 headings" in issues[0].message
    assert "not left aligned" in issues[0].message
    assert issues[0].expected == "left aligned, as the rest of the header"


def test_one_finding_per_header_not_one_per_cell() -> None:
    """Five cells, one act to fix: select the header row and set it. Five
    findings about one row is the noise this codebase keeps out of reports."""
    ctx = _ctx([_table(REAL)])

    assert len(_rows(ctx)) == 1
    assert len(_align(ctx)) == 1


def test_a_tidy_header_says_nothing() -> None:
    header = [(f"Column {n}", "LEFT (1)", HEADER_FILL) for n in range(4)]
    ctx = _ctx([_table(header)])

    assert _rows(ctx) == []
    assert _align(ctx) == []


def test_an_unstated_alignment_is_not_a_defect() -> None:
    """A cell that states no alignment inherits one, and in a left-to-right
    deck what it inherits renders flush left. Faulting "unstated" would report
    most of the body of most tables for being ordinary."""
    header = [(f"Column {n}", None, HEADER_FILL) for n in range(4)]

    assert _align(_ctx([_table(header)])) == []


# --------------------------------------------------------------------------- #
# What makes a row a header
# --------------------------------------------------------------------------- #

def test_a_top_row_filled_like_the_body_is_not_a_header() -> None:
    """Nothing sets it apart, so nothing says these are headings rather than
    the first row of data."""
    header = [("Centred", "CENTER (2)", BODY_FILL) for _ in range(4)]

    assert _align(_ctx([_table(header, body_fill=BODY_FILL)])) == []


def test_the_banding_flag_alone_can_carry_a_plain_top_row() -> None:
    """Weak on its own, which is why it does not define a header. It is still
    the deck author saying row 0 is one."""
    header = [("Centred", "CENTER (2)", BODY_FILL) for _ in range(4)]

    assert len(_align(_ctx([_table(header, body_fill=BODY_FILL, flagged=True)]))) == 1


def test_a_row_of_differently_coloured_boxes_is_not_a_header_band() -> None:
    """Three fills across the top row is a diagram, not a header, and holding
    its contents to a heading's rules would fault it for not being a table."""
    header = [
        ("One", "CENTER (2)", "FF0000"),
        ("Two", "CENTER (2)", "00FF00"),
        ("Three", "CENTER (2)", "0000FF"),
    ]

    assert _align(_ctx([_table(header)])) == []


def test_a_colour_coded_body_does_not_hide_its_header() -> None:
    """A status table colours its cells by value, which is the point of it.
    Demanding one fill across the body threw out exactly the components most
    likely to carry this defect."""
    shape = _table(REAL)
    for cell in shape.table.cells:
        if cell.row == 1:
            cell.fill_hex = {0: "F5CACA", 1: None, 2: "C8E6C9"}.get(
                cell.column % 3, None
            )

    assert len(_rows(_ctx([shape]))) == 1


def test_a_hidden_slide_is_left_alone() -> None:
    assert _rows(_ctx([_table(REAL)], hidden=True)) == []


def test_a_one_row_table_has_no_body_to_be_set_apart_from() -> None:
    shape = _table(REAL)
    shape.table.cells = [c for c in shape.table.cells if c.row == 0]
    shape.table.rows = 1

    assert _rows(_ctx([shape])) == []


# --------------------------------------------------------------------------- #
# Which edge is the leading one
# --------------------------------------------------------------------------- #

def test_a_right_to_left_deck_leads_on_the_right() -> None:
    """Arabic copy is set flush right, so a right-aligned heading is the
    correct one and a left-aligned heading is the defect. Reading this rule on
    the left edge would give an Arabic deck a clean bill of health, which is
    worse than saying the wrong thing."""
    arabic = _copy("مرحبا بالعالم في هذا العرض")
    header = [("عنوان", "RIGHT (3)", HEADER_FILL) for _ in range(4)]

    ctx = _ctx([_table(header)] + arabic)
    assert ctx.deck.rtl, "fixture is not reading as an Arabic deck"
    assert _align(ctx) == []

    left = [("عنوان", "LEFT (1)", HEADER_FILL) for _ in range(4)]
    issues = _align(_ctx([_table(left)] + arabic))

    assert len(issues) == 1
    assert "not right aligned" in issues[0].message


# --------------------------------------------------------------------------- #
# Reading a table out of a real file
# --------------------------------------------------------------------------- #

def test_a_real_table_comes_back_as_rows_of_cells(tmp_path: Path) -> None:
    """The extractor change this all rests on. Before it, every one of these
    cells was a single opaque rectangle."""
    pytest.importorskip("pptx")
    from formatting_tool.extract import read_deck
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_table(2, 3, Inches(0.5), Inches(1.5),
                                   Inches(6.0), Inches(1.6))
    frame.name = "Table 17"
    table = frame.table
    table.first_row = False
    for column, text in enumerate(["Chemistry", "Drivers", "Supply gap\nby 2030"]):
        cell = table.cell(0, column)
        cell.text_frame.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x43, 0x36, 0x26)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER if column == 2 else PP_ALIGN.LEFT
    for column in range(3):
        body = table.cell(1, column)
        body.text_frame.text = "body"
        body.fill.background()
    path = tmp_path / "table.pptx"
    prs.save(str(path))

    shape = next(
        s for s in read_deck(path).slides[0].shapes if s.table is not None
    )

    assert (shape.table.rows, shape.table.columns) == (2, 3)
    header = shape.table.row_at(0)
    assert [c.lines for c in header] == [1, 1, 2]
    assert [c.geometry.left_in for c in header] == pytest.approx([0.5, 2.5, 4.5])
    assert all(c.fill_hex == "433626" for c in header)
    assert header[2].paragraphs[0].alignment == "CENTER (2)"


def test_the_real_table_is_reported_end_to_end(tmp_path: Path) -> None:
    pytest.importorskip("pptx")
    from formatting_tool.extract import read_deck
    from formatting_tool.linemetrics import NullLineMetrics
    from formatting_tool.rules import build_default_rules, run_rules
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_table(2, 3, Inches(0.5), Inches(1.5),
                                   Inches(6.0), Inches(1.6))
    table = frame.table
    table.first_row = False
    for column, text in enumerate(["Chemistry", "Drivers", "Supply gap\nby 2030"]):
        cell = table.cell(0, column)
        cell.text_frame.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x43, 0x36, 0x26)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER if column == 2 else PP_ALIGN.LEFT
    for column in range(3):
        table.cell(1, column).text_frame.text = "body"
        table.cell(1, column).fill.background()
    path = tmp_path / "table.pptx"
    prs.save(str(path))

    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    ctx = RuleContext(deck=read_deck(path), spec=spec)
    found = {
        i.rule_id for i in run_rules(ctx, build_default_rules(NullLineMetrics()))
    }

    assert "table.header_rows" in found
    assert "table.header_alignment" in found


def test_a_deck_whose_arabic_lives_in_its_tables_still_reads_right_to_left() -> None:
    """A table's copy is not on the shape that holds it. A comparison-heavy
    Arabic deck is mostly tables, so counting shapes alone left almost nothing
    to count and the deck came back left-to-right -- which reports a correctly
    set Arabic deck as clean."""
    header = [("عنوان", "RIGHT (3)", HEADER_FILL) for _ in range(4)]
    shape = _table(header)
    for cell in shape.table.cells:
        if cell.row == 1:
            cell.text = "قيمة"

    ctx = _ctx([shape])

    assert ctx.deck.rtl
    assert _align(ctx) == []


# --------------------------------------------------------------------------- #
# And the fix
# --------------------------------------------------------------------------- #

def _deck_with_a_mixed_header(tmp_path: Path, rtl: bool = False) -> Path:
    """Two headings on the leading edge, two centred: the real shape of it."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_table(2, 4, Inches(0.5), Inches(1.5),
                                   Inches(8.0), Inches(1.6))
    frame.name = "Table 17"
    table = frame.table
    table.first_row = False
    lead = PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT
    headings = (
        ["عنوان الأول", "عنوان الثاني", "عنوان الثالث", "عنوان الرابع"] if rtl
        else ["Cell chemistry", "Demand drivers", "Captured by China",
              "Demand CAGR"]
    )
    for column, text in enumerate(headings):
        cell = table.cell(0, column)
        cell.text_frame.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x43, 0x36, 0x26)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = lead if column < 2 else PP_ALIGN.CENTER
    body = ["قيمة", "قيمة", "قيمة", "قيمة"] if rtl else ["body"] * 4
    for column in range(4):
        table.cell(1, column).text_frame.text = body[column]
        table.cell(1, column).fill.background()
    path = tmp_path / ("rtl.pptx" if rtl else "table.pptx")
    prs.save(str(path))
    return path


def _apply(deck: Path, out: Path):
    from formatting_tool.apply import apply_fixes
    from formatting_tool.extract import read_deck
    from formatting_tool.linemetrics import NullLineMetrics
    from formatting_tool.rules import build_default_rules, run_rules

    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    ctx = RuleContext(deck=read_deck(deck), spec=spec)
    found = [
        i for i in run_rules(ctx, build_default_rules(NullLineMetrics()))
        if i.rule_id == "table.header_alignment"
    ]
    for issue in found:
        issue.id = issue.fingerprint()
    return found, apply_fixes(deck, found, out, selected=[i.id for i in found])


def _alignments(path: Path) -> list:
    from pptx import Presentation

    table = next(
        s.table for s in Presentation(str(path)).slides[0].shapes if s.has_table
    )
    return [
        str(p.alignment)
        for cell in table.rows[0].cells
        for p in cell.text_frame.paragraphs
    ]


def test_the_header_row_is_set_to_the_leading_edge(tmp_path: Path) -> None:
    deck = _deck_with_a_mixed_header(tmp_path)
    out = tmp_path / "fixed.pptx"

    found, result = _apply(deck, out)

    assert len(found) == 1
    assert len(result.applied) == 1
    assert "set the header row left aligned" in result.applied[0].detail
    assert _alignments(out) == ["LEFT (1)"] * 4


def test_the_whole_row_is_set_not_only_the_cells_that_were_wrong(
    tmp_path: Path,
) -> None:
    """Two of four were already correct. Touching only the other two would
    leave the row in a third state nobody chose."""
    deck = _deck_with_a_mixed_header(tmp_path)
    out = tmp_path / "fixed.pptx"

    _, result = _apply(deck, out)

    # Two changed, two already correct and left as they were.
    assert "changing 2 heading line(s)" in result.applied[0].detail
    assert len(set(_alignments(out))) == 1


def test_an_arabic_header_is_set_right_not_left(tmp_path: Path) -> None:
    """A fixer that assumed left would take a correct Arabic table and break
    it. Which edge leads is read back off the finding."""
    deck = _deck_with_a_mixed_header(tmp_path, rtl=True)
    out = tmp_path / "fixed.pptx"

    found, result = _apply(deck, out)

    assert len(found) == 1
    assert "set the header row right aligned" in result.applied[0].detail
    assert _alignments(out) == ["RIGHT (3)"] * 4


def test_nothing_is_moved_or_resized(tmp_path: Path) -> None:
    from pptx import Presentation

    deck = _deck_with_a_mixed_header(tmp_path)
    out = tmp_path / "fixed.pptx"

    def box(path):
        s = next(x for x in Presentation(str(path)).slides[0].shapes if x.has_table)
        return (s.left, s.top, s.width, s.height)

    was = box(deck)
    _apply(deck, out)

    assert box(out) == was


def test_applying_it_twice_changes_nothing_the_second_time(
    tmp_path: Path,
) -> None:
    deck = _deck_with_a_mixed_header(tmp_path)
    out = tmp_path / "fixed.pptx"
    _apply(deck, out)

    found, _ = _apply(out, tmp_path / "again.pptx")

    assert found == []
