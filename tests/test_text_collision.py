"""Text that leaves its own box and lands on the shape below it.

Off a real slide: four status columns, each a line of copy over a thin red
progress bar. In two of them the copy ran to a third line, left its box, and
was drawn across the bar. Nothing reported it, and both rules that might have
had a reason not to:

- `space.overlap` measures STORED boxes and only fires between two
  text-bearing shapes. The bar is a coloured panel with no text, and on that
  slide the stored boxes did not touch anyway -- the copy box ended at
  5.02in and the bar started at 5.06in. The collision only exists once the
  text is drawn.
- `space.text_overflow` estimated the wrapped height from an average glyph
  advance and compared it against the box, so it could neither be sure the
  text overflowed nor say what the overflow landed on.

What settles it is asking the renderer where it actually put the text.
PowerPoint reports that per shape, and the difference is stark: measured, the
same slide reads "Text is drawn 0.40in tall in a 0.22in box, overhanging by
0.23in and is drawn across 'Status bar 1'".

These use a stub provider rather than driving PowerPoint, so what is under
test is the arithmetic and the judgement rather than the COM plumbing.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.extract import read_deck
from formatting_tool.extract.master_spec import derive_master_spec
from formatting_tool.linemetrics import ShapeKey, TextBounds
from formatting_tool.models import BrandGuidelines
from formatting_tool.rules import RuleContext, run_rules
from formatting_tool.rules.space import TextOverflowRule

STATUS = ("Faress and Mohamad are currently finalizing the data and deriving "
          "the trend accordingly")


class StubMetrics:
    """A renderer's answers, without a renderer."""

    def __init__(self, bounds: dict[ShapeKey, TextBounds] | None = None) -> None:
        self._bounds = bounds or {}

    @property
    def available(self) -> bool:
        return True

    def lines(self, key: ShapeKey):
        return None

    def bounds(self, key: ShapeKey):
        return self._bounds.get(key)


def _deck(tmp_path: Path, *, bar_top_in: float = 5.06,
          copy_height_in: float = 0.22):
    """A status line over a progress bar, as the slide had it."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    copy = slide.shapes.add_textbox(
        Inches(2.0), Inches(4.80), Inches(2.2), Inches(copy_height_in)
    )
    copy.name = "Status copy 1"
    copy.text_frame.word_wrap = True
    copy.text_frame.text = STATUS
    copy.text_frame.paragraphs[0].runs[0].font.size = Pt(8)

    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0),
                                 Inches(bar_top_in), Inches(2.2), Inches(0.06))
    bar.name = "Status bar 1"
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string("C00000")

    path = tmp_path / "status.pptx"
    prs.save(str(path))
    return path, copy.shape_id


def _run(deck: Path, metrics) -> list:
    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    ctx = RuleContext(deck=read_deck(deck), spec=spec)
    return list(TextOverflowRule(metrics=metrics).check(ctx))


# The text as PowerPoint actually drew it on that slide: 0.40in of it,
# starting a hair below the box top, so it ends at 5.25in.
def _overflowing(shape_id: int) -> StubMetrics:
    return StubMetrics({
        ShapeKey(1, shape_id): TextBounds(
            left_in=2.10, top_in=4.85, width_in=1.87, height_in=0.40
        )
    })


def _fitting(shape_id: int) -> StubMetrics:
    return StubMetrics({
        ShapeKey(1, shape_id): TextBounds(
            left_in=2.10, top_in=4.85, width_in=1.87, height_in=0.14
        )
    })


# --------------------------------------------------------------------------- #
# The case off the slide
# --------------------------------------------------------------------------- #

def test_the_overflow_onto_the_bar_is_reported(tmp_path: Path) -> None:
    deck, shape_id = _deck(tmp_path)

    found = _run(deck, _overflowing(shape_id))

    assert len(found) == 1
    assert found[0].rule_id == "space.text_overflow"
    assert "Status bar 1" in found[0].message
    assert "0.22in box" in found[0].message


def test_it_says_how_far_past_the_box_the_text_goes(tmp_path: Path) -> None:
    """A designer has to be able to tell a hair from a whole line."""
    deck, shape_id = _deck(tmp_path)

    found = _run(deck, _overflowing(shape_id))

    # Box bottom is 5.02in, the text is drawn to 5.25in.
    assert "0.23in" in found[0].message
    assert "0.23in past the bottom" in found[0].found
    assert "overhanging the bottom" in found[0].message


def test_the_overflow_finding_is_a_warning_either_way(tmp_path: Path) -> None:
    """It is one thing, and one severity: the box is too small for its copy.

    Whether that spills onto something is a different finding with a
    different remedy -- `space.text_collision`, which is reported on the
    shape that has to move and can be applied. This one names what it runs
    over for context and stops there, because shortening copy is a design
    call whatever is underneath.
    """
    from formatting_tool.models import Severity

    on_something, shape_id = _deck(tmp_path)
    hit = _run(on_something, _overflowing(shape_id))

    empty_dir = tmp_path / "empty"
    empty_dir.mkdir()
    nothing, other_id = _deck(empty_dir, bar_top_in=6.80)
    miss = _run(nothing, _overflowing(other_id))

    assert hit[0].severity is Severity.WARNING
    assert "Status bar 1" in hit[0].message      # context
    assert miss[0].severity is Severity.WARNING
    assert "Status bar 1" not in miss[0].message


# --------------------------------------------------------------------------- #
# What it must not report
# --------------------------------------------------------------------------- #

def test_text_that_fits_is_not_reported(tmp_path: Path) -> None:
    """And the estimate does not get to second-guess the measurement: a
    measured box that fits is finished, whatever the average-glyph guess
    would have said about it."""
    deck, shape_id = _deck(tmp_path)

    assert _run(deck, _fitting(shape_id)) == []


def test_a_caption_deliberately_over_a_panel_is_left_alone(
    tmp_path: Path,
) -> None:
    """Text over a coloured panel is ordinary layout -- a label on a band, a
    caption on a photo -- and `space.overlap` declines to report it for that
    reason. Only the part of the text below its OWN box is tested here, so a
    box that overlaps the panel on purpose says nothing.
    """
    deck, shape_id = _deck(tmp_path, bar_top_in=4.90, copy_height_in=0.60)

    # The text fits its tall box, and the box happens to sit over the bar.
    assert _run(deck, _fitting(shape_id)) == []


def test_with_no_renderer_the_old_estimate_still_runs(tmp_path: Path) -> None:
    """The measurement is an upgrade, not a dependency. A host with no
    PowerPoint keeps the proxy it had."""
    from formatting_tool.linemetrics import NullLineMetrics

    deck, shape_id = _deck(tmp_path)

    found = _run(deck, NullLineMetrics())

    assert len(found) == 1
    # The estimate cannot name what it hits, and says so in vaguer terms.
    assert "Status bar 1" not in found[0].message


def test_a_shape_the_renderer_said_nothing_about_falls_back(
    tmp_path: Path,
) -> None:
    """A provider that is available but has no answer for this shape -- a
    placeholder PowerPoint will not talk about -- must not silence the rule.
    """
    deck, shape_id = _deck(tmp_path)

    found = _run(deck, StubMetrics({}))      # available, but knows nothing

    assert len(found) == 1                   # the estimate answered instead


# --------------------------------------------------------------------------- #
# End to end, against the real renderer
# --------------------------------------------------------------------------- #

def test_the_real_renderer_finds_it_too(tmp_path: Path) -> None:
    """The stub above encodes numbers taken from PowerPoint. This checks they
    were not wishful thinking."""
    from formatting_tool import powerpoint
    from formatting_tool.linemetrics import PowerPointComMetrics

    if not powerpoint.available():
        pytest.skip("desktop PowerPoint is needed to measure text")

    deck, shape_id = _deck(tmp_path)
    metrics = PowerPointComMetrics(deck)

    bounds = metrics.bounds(ShapeKey(1, shape_id))
    assert bounds is not None
    assert bounds.height_in > 0.30           # three lines, not the two it fits
    assert bounds.bottom_in > 5.06           # past where the bar starts

    found = _run(deck, metrics)
    assert len(found) == 1
    assert "Status bar 1" in found[0].message


# --------------------------------------------------------------------------- #
# The collision, and moving something clear of it
# --------------------------------------------------------------------------- #
#
# This is the half that was missing. Detection alone left the defect in the
# deck: "nothing changed" was the whole complaint, and it was fair.
#
# The remedy is the one a designer reaches for first -- look at the space
# around the collision, and if something can move into it, move the shape that
# landed on the other. What makes that safe rather than reckless is that the
# applier already checks every geometric move against the shapes around it and
# puts it back if it lands on a neighbour, so "if there is room" is enforced
# rather than assumed.


def _collisions(deck: Path, metrics) -> list:
    from formatting_tool.rules.space import TextCollisionRule

    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    ctx = RuleContext(deck=read_deck(deck), spec=spec)
    found = list(TextCollisionRule(metrics=metrics).check(ctx))
    for issue in found:
        issue.id = issue.fingerprint()
    return found


def test_the_collision_is_reported_on_the_shape_that_moves(
    tmp_path: Path,
) -> None:
    """Geometry cannot say which of two shapes is in the wrong place. z-order
    says which one landed on the other, and the bar is drawn last."""
    deck, shape_id = _deck(tmp_path)

    found = _collisions(deck, _overflowing(shape_id))

    assert len(found) == 1
    assert found[0].shape == "Status bar 1"
    assert "is in front" in found[0].message


def test_the_finding_carries_the_position_to_move_to(tmp_path: Path) -> None:
    """The arithmetic belongs where the measurement is, so the fixer applies a
    position rather than deriving one."""
    deck, shape_id = _deck(tmp_path)

    found = _collisions(deck, _overflowing(shape_id))

    # Text is drawn to 5.25in, so the bar clears it at 5.27in.
    assert "moved to 2.00, 5.27in" in found[0].expected


def test_the_shape_is_actually_moved_clear(tmp_path: Path) -> None:
    from pptx import Presentation

    from formatting_tool.apply import apply_fixes

    deck, shape_id = _deck(tmp_path)
    found = _collisions(deck, _overflowing(shape_id))
    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, found, out, selected=[i.id for i in found],
                         spec=spec)

    assert len(result.applied) == 1
    assert "clear of the text drawn over it" in result.applied[0].detail
    bar = [s for s in Presentation(str(out)).slides[0].shapes
           if s.name == "Status bar 1"][0]
    # Down from 5.06in to past where the text is drawn, 5.25in.
    assert bar.top / 914400 > 5.25


def test_a_shape_with_nowhere_to_go_is_left_alone(tmp_path: Path) -> None:
    """The guard that makes this safe to apply by default.

    A bar with a footnote pinned right under it cannot move down, and shoving
    the footnote instead would trade one collision for another. The applier
    reverts the move and names what stopped it.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    from formatting_tool.apply import apply_fixes

    deck, shape_id = _deck(tmp_path)
    # Add a footnote immediately below the bar, then re-read.
    prs = Presentation(str(deck))
    slide = prs.slides[0]
    note = slide.shapes.add_textbox(Inches(2.0), Inches(5.14), Inches(2.2),
                                    Inches(0.20))
    note.name = "Footnote 1"
    note.text_frame.text = "1. Electric Vehicles"
    note.text_frame.paragraphs[0].runs[0].font.size = Pt(7)
    prs.save(str(deck))

    found = _collisions(deck, _overflowing(shape_id))
    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, found, out, selected=[i.id for i in found],
                         spec=spec)

    assert result.applied == []
    assert "Footnote 1" in result.skipped[0].detail
    bar = [s for s in Presentation(str(out)).slides[0].shapes
           if s.name == "Status bar 1"][0]
    assert bar.top / 914400 == pytest.approx(5.06, abs=0.001)


def test_text_that_fits_produces_no_collision(tmp_path: Path) -> None:
    """A bar under a box whose text stays inside it is a layout, not a defect,
    however close the two sit."""
    deck, shape_id = _deck(tmp_path)

    assert _collisions(deck, _fitting(shape_id)) == []


def test_no_renderer_means_no_collision_finding(tmp_path: Path) -> None:
    """It cannot be inferred from the file. Without a measurement the rule
    says nothing rather than guessing at where the text was drawn."""
    from formatting_tool.linemetrics import NullLineMetrics

    deck, shape_id = _deck(tmp_path)

    assert _collisions(deck, NullLineMetrics()) == []


def test_the_real_renderer_drives_the_whole_fix(tmp_path: Path) -> None:
    """End to end on the real thing: measure, decide, move, and measure again."""
    from pptx import Presentation

    from formatting_tool import powerpoint
    from formatting_tool.apply import apply_fixes
    from formatting_tool.linemetrics import PowerPointComMetrics

    if not powerpoint.available():
        pytest.skip("desktop PowerPoint is needed to measure text")

    deck, shape_id = _deck(tmp_path)
    found = _collisions(deck, PowerPointComMetrics(deck))
    assert len(found) == 1

    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    out = tmp_path / "fixed.pptx"
    result = apply_fixes(deck, found, out, selected=[i.id for i in found],
                         spec=spec)
    assert len(result.applied) == 1

    # Measured on the written deck: the text is where it was, the bar is not,
    # and they no longer touch.
    after = PowerPointComMetrics(out)
    ink = after.bounds(ShapeKey(1, shape_id))
    bar = [s for s in Presentation(str(out)).slides[0].shapes
           if s.name == "Status bar 1"][0]
    assert bar.top / 914400 > ink.bottom_in


# --------------------------------------------------------------------------- #
# Any edge, not just the bottom
# --------------------------------------------------------------------------- #

def _wide(shape_id: int) -> StubMetrics:
    """A line drawn 2.90in wide out of a 2.20in box: wrapping is off."""
    return StubMetrics({
        ShapeKey(1, shape_id): TextBounds(
            left_in=2.10, top_in=4.85, width_in=2.90, height_in=0.17
        )
    })


def test_text_running_off_the_right_is_reported(tmp_path: Path) -> None:
    """A wrapped paragraph can only grow downwards, so this began as a test of
    the bottom edge. Wrapping off is ordinary in a hand-built deck, and then a
    long line runs off the side instead."""
    deck, shape_id = _deck(tmp_path)

    found = _run(deck, _wide(shape_id))

    assert len(found) == 1
    assert "overhanging the right" in found[0].message


def test_the_way_out_of_a_sideways_collision_is_sideways(tmp_path: Path) -> None:
    from formatting_tool.rules.space import TextCollisionRule

    deck, shape_id = _deck(tmp_path, bar_top_in=4.85)   # beside the text now
    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    ctx = RuleContext(deck=read_deck(deck), spec=spec)

    found = list(TextCollisionRule(metrics=_wide(shape_id)).check(ctx))

    # Down is an alignment here -- the bar sits inside the ink vertically --
    # so the only way out is across, and across is 2.12in. That is a
    # relocation, not a nudge, so no target is offered and the overflow
    # finding is left to speak for it.
    assert found == []


# --------------------------------------------------------------------------- #
# The row moves as one
# --------------------------------------------------------------------------- #

def test_a_bar_in_a_row_takes_the_row_with_it(tmp_path: Path) -> None:
    """Nudging one bar of four clear of the copy above it breaks the row,
    which is a defect traded for a defect. `space.text_collision` is in COHORT
    so the set moves together or not at all.
    """
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    from formatting_tool.apply import apply_fixes
    from formatting_tool.rules.space import TextCollisionRule

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    first = None
    for col in range(4):
        x = 0.8 + col * 3.1
        copy = slide.shapes.add_textbox(Inches(x), Inches(4.80), Inches(2.2),
                                        Inches(0.22))
        copy.name = f"Status copy {col + 1}"
        copy.text_frame.word_wrap = True
        copy.text_frame.text = STATUS if col == 0 else "Casebook is finalized"
        copy.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        if col == 0:
            first = copy
    for col in range(4):
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col * 3.1),
            Inches(5.06), Inches(2.2), Inches(0.06),
        )
        bar.name = f"Status bar {col + 1}"
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string("C00000")
    deck = tmp_path / "row.pptx"
    prs.save(str(deck))

    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    ctx = RuleContext(deck=read_deck(deck), spec=spec)
    # The stub's ink has to sit inside THIS deck's first column, at x=0.8,
    # or the overhang comes out sideways and the test measures the wrong thing.
    ink = StubMetrics({
        ShapeKey(1, first.shape_id): TextBounds(
            left_in=0.90, top_in=4.85, width_in=1.87, height_in=0.40
        )
    })
    found = list(TextCollisionRule(metrics=ink).check(ctx))
    for issue in found:
        issue.id = issue.fingerprint()
    assert len(found) == 1

    out = tmp_path / "fixed.pptx"
    result = apply_fixes(deck, found, out, selected=[i.id for i in found],
                         spec=spec)

    assert len(result.applied) == 1
    assert "moves as one" in result.applied[0].detail
    tops = {
        s.name: round(s.top / 914400, 3)
        for s in Presentation(str(out)).slides[0].shapes
        if s.name.startswith("Status bar")
    }
    assert len(set(tops.values())) == 1, f"the row came apart: {tops}"
    assert set(tops.values()) != {5.06}, "nothing moved at all"


# --------------------------------------------------------------------------- #
# Growing the box, where there is room
# --------------------------------------------------------------------------- #

def test_the_box_grows_into_empty_space(tmp_path: Path) -> None:
    """The one remedy that is arithmetic rather than judgement: the copy and
    the type are untouched, and the stored box stops lying about what is
    drawn -- which is what breaks a rebuild onto a master."""
    from pptx import Presentation

    from formatting_tool.apply import apply_fixes

    deck, shape_id = _deck(tmp_path, bar_top_in=6.80)   # nothing in the way
    found = _run(deck, _overflowing(shape_id))
    for issue in found:
        issue.id = issue.fingerprint()
    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, found, out, selected=[i.id for i in found],
                         spec=spec)

    assert len(result.applied) == 1
    assert "taller" in result.applied[0].detail
    box = [s for s in Presentation(str(out)).slides[0].shapes
           if s.name == "Status copy 1"][0]
    # Text is drawn to 5.25in from a box top of 4.80in, plus the inset.
    assert box.height / 914400 == pytest.approx(0.50, abs=0.02)


def test_the_box_does_not_grow_into_what_the_text_runs_over(
    tmp_path: Path,
) -> None:
    """Growing reaches in the same direction the text already spilled.

    So where the text runs over something, the rule offers no box to grow
    into at all and the remedy is `space.text_collision` moving that shape.
    Refused in the rule rather than the fixer because the fixer sees the
    top-level shapes and the rule has looked at all of them.
    """
    from pptx import Presentation

    from formatting_tool.apply import apply_fixes

    deck, shape_id = _deck(tmp_path)          # the bar is right below
    found = _run(deck, _overflowing(shape_id))
    for issue in found:
        issue.id = issue.fingerprint()
    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, found, out, selected=[i.id for i in found],
                         spec=spec)

    assert result.applied == []
    assert "no box to grow into" in result.skipped[0].detail
    box = [s for s in Presentation(str(out)).slides[0].shapes
           if s.name == "Status copy 1"][0]
    assert box.height / 914400 == pytest.approx(0.22, abs=0.001)


def test_a_box_that_already_fits_is_not_grown(tmp_path: Path) -> None:
    deck, shape_id = _deck(tmp_path, bar_top_in=6.80)

    assert _run(deck, _fitting(shape_id)) == []


# --------------------------------------------------------------------------- #
# Growing a box moves the text unless it is top-anchored
# --------------------------------------------------------------------------- #
#
# The regression that made the deck worse rather than better. Growing a box
# looks like it only adds room, and it does -- for top-anchored text. For
# middle-anchored text it moves the copy down by half the growth, and for
# bottom-anchored text by all of it. Measured on the box off that slide,
# growing 0.22in to 0.50in moved bottom-anchored copy from ending at 4.97in to
# ending at 5.25in: onto a bar at 5.06in it had not been touching at all.
#
# So the fix CREATED the collision it was there to help with, which is worse
# than doing nothing, and "it did the opposite" was the right description.


def _anchored_deck(tmp_path: Path, anchor, *, with_bar: bool):
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    copy = slide.shapes.add_textbox(Inches(2.0), Inches(4.80), Inches(2.2),
                                    Inches(0.22))
    copy.name = "Status copy 1"
    copy.text_frame.word_wrap = True
    copy.text_frame.vertical_anchor = anchor
    copy.text_frame.text = STATUS
    copy.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    if with_bar:
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0),
                                     Inches(5.06), Inches(2.2), Inches(0.06))
        bar.name = "Status bar 1"
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string("C00000")
    path = tmp_path / "anchored.pptx"
    prs.save(str(path))
    return path, copy.shape_id


@pytest.mark.parametrize("anchor_name", ["MIDDLE", "BOTTOM"])
def test_a_box_whose_text_would_move_is_not_grown(
    tmp_path: Path, anchor_name: str
) -> None:
    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR

    from formatting_tool.apply import apply_fixes

    deck, shape_id = _anchored_deck(
        tmp_path, getattr(MSO_ANCHOR, anchor_name), with_bar=False
    )
    found = _run(deck, _overflowing(shape_id))
    for issue in found:
        issue.id = issue.fingerprint()
    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, found, out, selected=[i.id for i in found],
                         spec=spec)

    assert result.applied == []
    assert "anchored" in result.skipped[0].detail
    box = [s for s in Presentation(str(out)).slides[0].shapes
           if s.name == "Status copy 1"][0]
    assert box.height / 914400 == pytest.approx(0.22, abs=0.001)


def test_growing_a_top_anchored_box_leaves_the_text_where_it_was(
    tmp_path: Path,
) -> None:
    """The measurement that makes the top-anchored case safe, asserted against
    the real renderer rather than assumed."""
    from pptx.enum.text import MSO_ANCHOR

    from formatting_tool import powerpoint
    from formatting_tool.apply import apply_fixes
    from formatting_tool.linemetrics import PowerPointComMetrics

    if not powerpoint.available():
        pytest.skip("desktop PowerPoint is needed to measure text")

    deck, shape_id = _anchored_deck(tmp_path, MSO_ANCHOR.TOP, with_bar=False)
    was = PowerPointComMetrics(deck).bounds(ShapeKey(1, shape_id))

    found = _run(deck, PowerPointComMetrics(deck))
    for issue in found:
        issue.id = issue.fingerprint()
    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    out = tmp_path / "fixed.pptx"
    result = apply_fixes(deck, found, out, selected=[i.id for i in found],
                         spec=spec)

    assert len(result.applied) == 1
    assert "taller" in result.applied[0].detail
    now = PowerPointComMetrics(out).bounds(ShapeKey(1, shape_id))
    # The box changed; the drawn text did not move a thousandth of an inch.
    assert now.bottom_in == pytest.approx(was.bottom_in, abs=0.005)


def test_a_bottom_anchored_box_over_a_bar_is_not_made_worse(
    tmp_path: Path,
) -> None:
    """End to end on the exact regression, against the real renderer: the
    text must not end up lower than it started."""
    from pptx.enum.text import MSO_ANCHOR

    from formatting_tool import powerpoint
    from formatting_tool.apply import apply_fixes
    from formatting_tool.linemetrics import PowerPointComMetrics
    from formatting_tool.rules import build_default_rules, run_rules

    if not powerpoint.available():
        pytest.skip("desktop PowerPoint is needed to measure text")

    deck, shape_id = _anchored_deck(tmp_path, MSO_ANCHOR.BOTTOM, with_bar=True)
    was = PowerPointComMetrics(deck).bounds(ShapeKey(1, shape_id))

    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    ctx = RuleContext(deck=read_deck(deck), spec=spec)
    found = [
        i for i in run_rules(ctx, build_default_rules(PowerPointComMetrics(deck)))
        if i.rule_id in ("space.text_overflow", "space.text_collision")
    ]
    for issue in found:
        issue.id = issue.fingerprint()
    out = tmp_path / "fixed.pptx"

    apply_fixes(deck, found, out, selected=[i.id for i in found], spec=spec)

    now = PowerPointComMetrics(out).bounds(ShapeKey(1, shape_id))
    assert now.bottom_in <= was.bottom_in + 0.005, (
        f"the copy was pushed down, from {was.bottom_in:.3f}in to "
        f"{now.bottom_in:.3f}in"
    )
