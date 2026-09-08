"""Text that does not fit the box it is in.

The defect a rebuild creates more than any other, and the rule that was a stub
returning nothing while it did. A messy deck hides overlong copy behind
shrink-to-fit; put that slide on the master and the copy arrives at the
master's size with nothing shrinking it, and it runs out of its box and over
whatever is under it.

The three autofit modes are three different claims and are not treated alike,
which is most of what these pin down.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.models import BrandGuidelines

LONG = (
    "The new future that NEOM is creating embodies sustainability in all that "
    "we do - regenerating the environment, transforming our economy and "
    "promoting human wellbeing as we seek to overcome pressing global "
    "challenges. "
) * 2


def _findings(tmp_path: Path, boxes):
    """Build a one-slide deck of text boxes and return its overflow findings."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt

    from formatting_tool.extract import read_deck
    from formatting_tool.extract.master_spec import derive_master_spec
    from formatting_tool.rules import RuleContext
    from formatting_tool.rules.space import TextOverflowRule

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for name, (x, y, w, h), text, wrap, autofit in boxes:
        shape = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.name = name
        frame = shape.text_frame
        frame.text = text
        frame.word_wrap = wrap
        frame.auto_size = autofit
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(18)

    path = tmp_path / "overflow.pptx"
    prs.save(str(path))

    spec = derive_master_spec(read_deck("test_master1.pptx"), BrandGuidelines())
    context = RuleContext(deck=read_deck(path), spec=spec)
    return {i.shape: i for i in TextOverflowRule().check(context)}


def _autofit():
    from pptx.enum.text import MSO_AUTO_SIZE

    return MSO_AUTO_SIZE


def test_copy_taller_than_its_box_is_reported(tmp_path: Path) -> None:
    found = _findings(tmp_path, [
        ("Body", (0.5, 1.0, 4.0, 0.8), LONG, True, _autofit().NONE),
    ])

    assert "Body" in found
    assert "runs outside it" in found["Body"].message


def test_copy_that_fits_is_not(tmp_path: Path) -> None:
    found = _findings(tmp_path, [
        ("Body", (0.5, 1.0, 4.0, 2.5), "NEOM's Vision", True, _autofit().NONE),
    ])

    assert found == {}


def test_wrapping_off_is_measured_across_not_down(tmp_path: Path) -> None:
    """One line however long it is, so the box's width is the constraint. A
    height check would call this box roomy."""
    found = _findings(tmp_path, [
        (
            "Heading", (6.0, 1.0, 2.0, 2.0),
            "Partnership Challenges and Mitigation Strategies",
            False, _autofit().NONE,
        ),
    ])

    assert "Heading" in found
    assert "Wrapping is off" in found["Heading"].message


def test_shrink_to_fit_is_not_an_overflow(tmp_path: Path) -> None:
    """The text is made to fit. That it had to be is a finding, and
    `size.autofit_shrink` is the one that makes it."""
    found = _findings(tmp_path, [
        ("Body", (6.0, 5.0, 3.0, 0.6), LONG, True, _autofit().TEXT_TO_FIT_SHAPE),
    ])

    assert found == {}


def test_grow_to_fit_is_reported_as_a_box_that_lies_about_its_size(
    tmp_path: Path,
) -> None:
    """The text does not overflow -- the box grows. But it is stored smaller
    than it renders, so it grows into its neighbour and the overlap rule,
    measuring the stored box, sees nothing."""
    found = _findings(tmp_path, [
        ("Body", (6.0, 3.0, 4.0, 0.6), LONG, True, _autofit().SHAPE_TO_FIT_TEXT),
    ])

    assert "Body" in found
    assert "grows to fit" in found["Body"].message
    assert found["Body"].severity.value == "warning"


def test_an_empty_box_says_nothing(tmp_path: Path) -> None:
    found = _findings(tmp_path, [
        ("Empty", (0.5, 1.0, 4.0, 0.4), "   ", True, _autofit().NONE),
    ])

    assert found == {}


def test_the_estimate_is_given_a_line_of_slack(tmp_path: Path) -> None:
    """The wrap is estimated from an average glyph advance, not rendered. An
    estimate half a line out must not produce a finding, or a report fills
    with maybes."""
    from formatting_tool.rules.space import _SLACK_IN

    assert _SLACK_IN >= 0.2
