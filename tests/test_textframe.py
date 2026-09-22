"""The text inside a shape, as opposed to where the shape sits.

Three settings nothing in this tool read, each of which defeats something the
tool otherwise does well: the anchor decides whether a failing box can be
corrected at all, the autofit scale decides whether the size in the file is the
size on the screen, and the insets decide whether a row of cards reads as a row
while every box on it is exactly where it should be.
"""

from __future__ import annotations

import pytest

from formatting_tool.apply import fixer_for
from formatting_tool.apply.fixers import LeaveAlone
from formatting_tool.extract import derive_master_spec
from formatting_tool.guidelines import BrandGuidelines
from formatting_tool.linemetrics import ShapeKey, TextBounds
from formatting_tool.models import (
    DeckProfile,
    Geometry,
    ParagraphProfile,
    RunProfile,
    Severity,
    ShapeProfile,
    SlideProfile,
)
from formatting_tool.rules.base import RuleContext
from formatting_tool.rules.textframe import (
    AnchorBlocksFitRule,
    AutofitScaleRule,
    TextInsetRule,
)

WIDE, TALL = 13.333, 7.5


class _Metrics:
    """A renderer that says where the text was actually drawn."""

    def __init__(self, drawn: dict):
        self.drawn = drawn
        self.available = True

    def lines(self, key):                     # pragma: no cover - not read here
        return None

    def bounds(self, key: ShapeKey):
        return self.drawn.get(key.shape_id)


def _shape(name, top=1.0, height=0.5, **kwargs):
    size = kwargs.pop("size_pt", 14.0)
    text = kwargs.pop("text", "some copy")
    return ShapeProfile(
        shape_id=kwargs.pop("shape_id", abs(hash(name)) % 9999),
        name=name, shape_type="TEXT_BOX (17)",
        geometry=Geometry(left_in=1.0, top_in=top, width_in=3.0,
                          height_in=height),
        text=text,
        paragraphs=[ParagraphProfile(
            text=text, runs=[RunProfile(text=text, size_pt=size)])],
        **kwargs,
    )


def _ctx(*shapes):
    deck = DeckProfile(path="deck.pptx", width_in=WIDE, height_in=TALL,
                       slides=[SlideProfile(number=1, shapes=list(shapes))])
    master = DeckProfile(path="master.pptx", width_in=WIDE, height_in=TALL)
    return RuleContext(deck=deck,
                       spec=derive_master_spec(master, BrandGuidelines()))


# --------------------------------------------------------------------------- #
# An anchor that stops a box being grown
# --------------------------------------------------------------------------- #

def test_copy_that_will_not_fit_a_box_nothing_can_grow_is_reported() -> None:
    """The finding the applier has been asking for. `fix_text_overflow` refuses
    a middle-anchored box because growing one moves the copy instead of giving
    it room, and its refusal ends "anchor the text to the top first" -- a
    correction nothing could make, because nothing read the anchor."""
    shape = _shape("Caption 1", shape_id=11, vertical_anchor="middle")
    drawn = {11: TextBounds(left_in=1.0, top_in=1.0, width_in=3.0, height_in=0.9)}

    [issue] = list(AnchorBlocksFitRule(metrics=_Metrics(drawn)).check(
        _ctx(shape)))
    assert "0.40in past its box" in issue.message
    assert issue.found == "middle-anchored"
    assert "grown to fit" in issue.expected


def test_a_top_anchored_box_that_overflows_is_the_overflow_rules_finding(
) -> None:
    """This rule is not about overflow. It is about an overflow that cannot be
    corrected, and a top-anchored box can be grown."""
    shape = _shape("Caption 1", shape_id=11, vertical_anchor="top")
    drawn = {11: TextBounds(left_in=1.0, top_in=1.0, width_in=3.0, height_in=0.9)}
    assert list(AnchorBlocksFitRule(metrics=_Metrics(drawn)).check(
        _ctx(shape))) == []


def test_a_middle_anchored_box_whose_copy_fits_is_nobodys_finding() -> None:
    """Middle-anchoring is how most labels on most decks are set. On its own it
    is a preference, and reporting it would bury the case that is not."""
    shape = _shape("Label 1", shape_id=11, vertical_anchor="middle")
    drawn = {11: TextBounds(left_in=1.0, top_in=1.05, width_in=3.0, height_in=0.4)}
    assert list(AnchorBlocksFitRule(metrics=_Metrics(drawn)).check(
        _ctx(shape))) == []


def test_without_a_renderer_the_anchor_rule_says_nothing() -> None:
    """Whether copy fits is not in the file."""
    shape = _shape("Caption 1", shape_id=11, vertical_anchor="middle")
    assert list(AnchorBlocksFitRule().check(_ctx(shape))) == []


def test_the_anchor_fix_lifts_the_text_and_says_where_it_was() -> None:
    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.5))
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    box.text_frame.paragraphs[0].add_run().text = "copy"

    shape = _shape("Caption 1", shape_id=11, vertical_anchor="middle")
    drawn = {11: TextBounds(left_in=1.0, top_in=1.0, width_in=3.0, height_in=0.9)}
    [issue] = list(AnchorBlocksFitRule(metrics=_Metrics(drawn)).check(
        _ctx(shape)))

    detail = fixer_for(issue)(box, issue, None)
    assert "anchored the text to the top, from middle" in detail
    assert box.text_frame.vertical_anchor == MSO_ANCHOR.TOP

    # And a box already anchored top is said to be, rather than silently
    # reported as corrected.
    with pytest.raises(LeaveAlone):
        fixer_for(issue)(box, issue, None)


# --------------------------------------------------------------------------- #
# A box drawing its text at a size the file does not admit to
# --------------------------------------------------------------------------- #

def test_the_amount_of_the_shrink_is_the_finding_not_the_setting() -> None:
    """`size.autofit_shrink` reports the setting and always has. The amount is
    what defeats every size rule in the tool, all of which read the file and
    see 14 where a reader sees 8.8."""
    shape = _shape("Heading 1", shape_id=11, size_pt=14.0, autofit_scale=0.625)

    [issue] = list(AutofitScaleRule().check(_ctx(shape)))
    assert "drawing its 14pt copy at 8.8pt" in issue.message
    assert issue.found == "62% of 14pt"
    assert issue.severity is Severity.ERROR       # under four fifths


def test_a_box_that_is_set_to_shrink_but_is_not_shrinking_is_not_reported(
) -> None:
    """A trap waiting for the next copy edit is not a defect today, and
    reporting it would bury the boxes that are actually shrinking."""
    assert list(AutofitScaleRule().check(
        _ctx(_shape("Heading 1", shape_id=11, autofit_scale=1.0)))) == []
    assert list(AutofitScaleRule().check(
        _ctx(_shape("Heading 1", shape_id=11, autofit_scale=None)))) == []
    # And a scale so near 1.0 that nobody could see it is rounding, not a
    # finding.
    assert list(AutofitScaleRule().check(
        _ctx(_shape("Heading 1", shape_id=11, autofit_scale=0.98)))) == []


def test_the_autofit_fix_writes_down_what_is_already_on_the_screen() -> None:
    """Nothing about the slide changes, which is why this is safe to do without
    a designer: the deck renders exactly as it did and the tool can finally see
    what it renders as."""
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.5))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "A heading"
    run.font.size = Pt(14)
    body = box.text_frame._txBody.find(qn("a:bodyPr"))
    body.append(body.makeelement(qn("a:normAutofit"),
                                 {"fontScale": "62500"}))

    shape = _shape("Heading 1", shape_id=11, size_pt=14.0, autofit_scale=0.625)
    [issue] = list(AutofitScaleRule().check(_ctx(shape)))

    detail = fixer_for(issue)(box, issue, None)
    assert "62% size they were already being drawn at" in detail
    assert round(box.text_frame.paragraphs[0].runs[0].font.size.pt) == 9
    assert body.findall(qn("a:normAutofit")) == []


def test_inherited_type_has_no_size_to_write_the_drawn_one_over() -> None:
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.5))
    box.text_frame.paragraphs[0].add_run().text = "A heading"
    body = box.text_frame._txBody.find(qn("a:bodyPr"))
    body.append(body.makeelement(qn("a:normAutofit"), {"fontScale": "62500"}))

    shape = _shape("Heading 1", shape_id=11, size_pt=14.0, autofit_scale=0.625)
    [issue] = list(AutofitScaleRule().check(_ctx(shape)))

    with pytest.raises(LeaveAlone) as refusal:
        fixer_for(issue)(box, issue, None)
    assert "inherited" in str(refusal.value)
    # And the shrink is still there, because nothing was written down.
    assert body.findall(qn("a:normAutofit"))


# --------------------------------------------------------------------------- #
# Copy that starts at a different inset from the row it repeats with
# --------------------------------------------------------------------------- #

def _cards(*insets):
    return [
        ShapeProfile(
            shape_id=10 + index, name=f"Card {index}", shape_type="TEXT_BOX (17)",
            geometry=Geometry(left_in=0.6 + index * 3.6, top_in=3.0,
                              width_in=3.4, height_in=1.2),
            text="Card", text_margins=inset,
            paragraphs=[ParagraphProfile(
                text="Card", runs=[RunProfile(text="Card", size_pt=12.0)])],
        )
        for index, inset in enumerate(insets)
    ]


def test_a_card_whose_copy_starts_further_in_than_its_row_is_reported() -> None:
    """The misalignment no alignment rule can see: every box is exactly where
    it should be, and what a reader sees is text that does not line up."""
    inset = (0.1, 0.05, 0.1, 0.05)
    odd = (0.3, 0.05, 0.3, 0.05)

    [issue] = list(TextInsetRule().check(_ctx(*_cards(inset, inset, odd))))
    assert issue.shape == "Card 2"
    assert issue.expected == "0.1in at the sides and 0.05in top and bottom"
    assert issue.found == "0.3in at the sides and 0.05in top and bottom"


def test_a_row_with_no_majority_is_a_set_of_decisions_rather_than_a_drift(
) -> None:
    """Two cards at one inset and two at another is a pair of decisions
    somebody made, and setting the row to whichever came first is a coin toss a
    designer has to check anyway."""
    a, b = (0.1, 0.05, 0.1, 0.05), (0.3, 0.05, 0.3, 0.05)
    assert list(TextInsetRule().check(_ctx(*_cards(a, a, b, b)))) == []


def test_two_shapes_are_not_a_row() -> None:
    a, b = (0.1, 0.05, 0.1, 0.05), (0.3, 0.05, 0.3, 0.05)
    assert list(TextInsetRule().check(_ctx(*_cards(a, b)))) == []


def test_shapes_of_different_sizes_are_not_a_repeated_component() -> None:
    """The grouping `typography._heading_rows` uses, and deliberately the same:
    two ways of deciding what counts as a row would report two different rows
    on one slide."""
    cards = _cards((0.1, 0.05, 0.1, 0.05), (0.1, 0.05, 0.1, 0.05),
                   (0.3, 0.05, 0.3, 0.05))
    cards[2].geometry = Geometry(left_in=7.8, top_in=3.0, width_in=2.0,
                                 height_in=1.2)
    assert list(TextInsetRule().check(_ctx(*cards))) == []


def test_the_inset_fix_sets_all_four_sides() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.margin_left = Inches(0.3)
    box.text_frame.margin_right = Inches(0.3)
    box.text_frame.paragraphs[0].add_run().text = "Card"

    inset = (0.1, 0.05, 0.1, 0.05)
    [issue] = list(TextInsetRule().check(
        _ctx(*_cards(inset, inset, (0.3, 0.05, 0.3, 0.05)))))

    detail = fixer_for(issue)(box, issue, None)
    assert "set its text insets to" in detail
    assert round(box.text_frame.margin_left.inches, 2) == 0.1
    assert round(box.text_frame.margin_top.inches, 2) == 0.05
    assert round(box.text_frame.margin_right.inches, 2) == 0.1
    assert round(box.text_frame.margin_bottom.inches, 2) == 0.05


# --------------------------------------------------------------------------- #
# A table's colours
#
# "Any icon, text or shape whose colour is off the palette should be changed" --
# and a table was none of the three as far as the palette rules were concerned.
# Its copy lives on its cells and `ShapeProfile.table` keeps those out of
# `children` on purpose, so no rule reading runs saw a word of it and no rule
# reading fills saw a cell. A deck's tables went out in whatever colours they
# arrived in, measured by nothing.
# --------------------------------------------------------------------------- #

def _table_shape(cells, name="Table 1"):
    from formatting_tool.models import TableProfile
    shape = ShapeProfile(
        shape_id=42, name=name, shape_type="TABLE (19)",
        geometry=Geometry(left_in=1.0, top_in=1.0, width_in=8.0, height_in=2.0),
        text="Region",
    )
    shape.table = TableProfile(rows=1, columns=len(cells), cells=list(cells))
    return shape


def _cell(column, fill=None, ink=None):
    from formatting_tool.models import TableCell
    runs = [RunProfile(text="Region", color_hex=ink, size_pt=11.0)] if ink else []
    return TableCell(
        row=0, column=column, fill_hex=fill, text="Region" if ink else "",
        paragraphs=[ParagraphProfile(text="Region", runs=runs)] if runs else [],
    )


def _palette_ctx(shape):
    """A context whose master states one palette, so off-palette means something."""
    master = DeckProfile(
        path="master.pptx", width_in=WIDE, height_in=TALL,
        theme_colors={"dk1": "1A1A1A", "lt1": "FFFFFF", "accent1": "1E2761"},
    )
    deck = DeckProfile(path="deck.pptx", width_in=WIDE, height_in=TALL,
                       slides=[SlideProfile(number=1, shapes=[shape])])
    return RuleContext(deck=deck,
                       spec=derive_master_spec(master, BrandGuidelines()))


def test_a_cells_text_colour_is_measured_like_any_other() -> None:
    from formatting_tool.rules.colors import OffPaletteTextRule

    shape = _table_shape([_cell(0, fill="1E2761", ink="FF00AA")])
    found = list(OffPaletteTextRule().check(_palette_ctx(shape)))

    assert [i.found for i in found] == ["#FF00AA"]
    assert found[0].shape == "Table 1"      # addressed to the table, not a cell


def test_a_cells_fill_is_measured_like_any_other() -> None:
    """A table shaded in a colour nobody approved read as a shape with no fill
    at all: the colour is on the cells, not on the graphic frame."""
    from formatting_tool.rules.colors import OffPaletteShapeRule

    shape = _table_shape([_cell(0, fill="FF00AA"), _cell(1, fill="1E2761")])
    found = [i for i in OffPaletteShapeRule().check(_palette_ctx(shape))
             if "FF00AA" in (i.found or "")]
    assert found, "the off-palette cell fill was not measured"


def test_a_cell_on_the_palette_is_left_alone() -> None:
    from formatting_tool.rules.colors import OffPaletteShapeRule, OffPaletteTextRule

    shape = _table_shape([_cell(0, fill="1E2761", ink="FFFFFF")])
    ctx = _palette_ctx(shape)
    assert list(OffPaletteTextRule().check(ctx)) == []
    assert list(OffPaletteShapeRule().check(ctx)) == []
