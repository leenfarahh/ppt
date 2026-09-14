"""Clearing a hardcoded typeface that is not the whole of what a run says.

`font.family.theme_drift` reports a run that hardcodes the theme's own
typeface, so removing the hardcode should change nothing. That holds only when
the Latin face is everything the run says about its typeface. On a real deck it
was not:

    <a:latin typeface="Cambria" .../>      <- the fixer removed this
    <a:cs    typeface="Arial"   .../>      <- and left this

Asked directly, PowerPoint reported the range's font as 'Cambria' before and as
'' -- meaning mixed -- afterwards, and the text grew from one line of 232.5pt
to two of 170.9pt. A heading that fitted its box stopped fitting it, and the
collision that followed moved a column of shapes off the page.
"""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from formatting_tool.apply.fixers import (
    LeaveAlone,
    _other_script_faces,
    fix_theme_font_drift,
)
from formatting_tool.models import Category, Issue, Severity

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _shape_with(latin: str, **slots: str):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "POOL OF INDEPENDENT ASSESSORS"
    run.font.size = Pt(14)
    run.font.name = latin
    rPr = run._r.find(A_NS + "rPr")
    for slot, face in slots.items():
        node = rPr.makeelement(A_NS + slot, {"typeface": face})
        rPr.append(node)
    return box


def _issue(found="Cambria"):
    return Issue(
        category=Category.FONT_FAMILY,
        severity=Severity.WARNING,
        rule_id="font.family.theme_drift",
        message=f"1 run(s) hardcode {found!r}, which is the theme typeface.",
        found=f"explicit {found}",
        expected="inherited from the layout",
    )


def test_another_script_pinning_a_different_face_is_found() -> None:
    shape = _shape_with("Cambria", cs="Arial")

    assert _other_script_faces(shape, "Cambria") == {"Arial"}


def test_a_theme_reference_is_not_a_pin() -> None:
    """`+mn-ea` already says "inherit", so it is not something left behind."""
    shape = _shape_with("Cambria", ea="+mn-ea")

    assert _other_script_faces(shape, "Cambria") == set()


def test_the_same_face_in_another_slot_is_not_a_conflict() -> None:
    shape = _shape_with("Cambria", cs="Cambria")

    assert _other_script_faces(shape, "Cambria") == set()


def test_nothing_else_pinned_means_nothing_to_report() -> None:
    shape = _shape_with("Cambria")

    assert _other_script_faces(shape, "Cambria") == set()


def test_the_fix_stands_down_when_another_script_is_pinned() -> None:
    """Neither half is safe on its own: leaving `a:cs` changes the rendered
    face, and clearing it would change the face Arabic renders in. What the
    run needs is a decision about both, which is a designer's call."""
    shape = _shape_with("Cambria", cs="Arial")

    with pytest.raises(LeaveAlone) as raised:
        fix_theme_font_drift(shape, _issue(), ctx=None)

    assert "another" in str(raised.value)
    # And the hardcode is still there, untouched.
    assert shape.text_frame.paragraphs[0].runs[0].font.name == "Cambria"


def test_the_fix_still_clears_a_plain_hardcode() -> None:
    """The case it was written for, which must keep working."""
    shape = _shape_with("Cambria")

    detail = fix_theme_font_drift(shape, _issue(), ctx=None)

    assert detail and "cleared" in detail
    assert shape.text_frame.paragraphs[0].runs[0].font.name is None
