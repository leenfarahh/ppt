"""Three ways the rebuild moved things it had no business moving.

All three came off the same Arabic deck, and all three have the same shape: a
rule written for one slide that was right about that slide and wrong about the
rest of the world.

1. `lift_imagery` pushed a row of flag roundels out of the panel holding them
   and down onto the footer. It was written to move an icon off artwork it was
   invisible against, and the test for "on the artwork" was touching it.
2. `series._ordered` put a run in left-to-right reading order whatever language
   it was in, and `claim_runs` fills run to run in reading order and then
   DELETES the boxes it emptied. On an Arabic deck every label landed in the
   mirror-image cell.
3. `_copy_text` returned silently when the target could not hold text, and the
   caller deleted the source anyway.
"""

from __future__ import annotations

from formatting_tool.rebuild import series
from formatting_tool.rebuild.builder import lift_imagery

EMU = 914400
CANVAS = (13.333, 7.5)


class Box:
    def __init__(self, left, top, width, height, text="", kind="TEXT_BOX (17)",
                 placeholder=False):
        self.left = int(left * EMU)
        self.top = int(top * EMU)
        self.width = int(width * EMU)
        self.height = int(height * EMU)
        self._text = text
        self.shape_type = kind
        self.is_placeholder = placeholder
        self.name = text or kind

    @property
    def has_text_frame(self):
        return True

    @property
    def text_frame(self):
        return self

    @property
    def text(self):
        return self._text

    def top_in(self):
        return round(self.top / EMU, 2)


class Slide:
    def __init__(self, shapes, layout_shapes):
        self.shapes = shapes
        self.slide_layout = type("L", (), {"shapes": layout_shapes})()


# --------------------------------------------------------------------------- #
# 1. The flag roundels
# --------------------------------------------------------------------------- #

def _full_page_picture():
    """The master behind the slide that broke this: a photograph clipped to
    the whole page, which every shape on the slide therefore overlaps."""
    return [Box(0.0, 0.0, 13.33, 7.5, kind="PICTURE (13)")]


def _flag_row():
    """Five roundels in a panel near the bottom, with the source line below."""
    flags = [Box(x, 6.15, 0.22, 0.22, kind="PICTURE (13)")
             for x in (7.55, 7.85, 8.15, 8.45, 8.75)]
    panel = Box(7.30, 5.95, 1.90, 0.62)                   # no copy
    footer = Box(0.56, 7.05, 10.35, 0.30, "Ministry of the Interior | Source")
    return flags, [*flags, panel, footer]


def test_a_row_of_flags_is_not_dragged_down_onto_the_footer() -> None:
    """The reported bug. They sat in a panel at the bottom of the page and
    came out below it, on the source line."""
    flags, shapes = _flag_row()
    before = [flag.top_in() for flag in flags]

    moved = lift_imagery(Slide(shapes, _full_page_picture()), CANVAS)

    assert moved == []
    assert [flag.top_in() for flag in flags] == before


def test_a_move_that_would_not_clear_the_artwork_is_refused() -> None:
    """The discriminator. This exists to get a graphic off artwork it is
    invisible against; a destination inside the same artwork achieves nothing
    and only takes the graphic out of whatever was holding it."""
    icon = Box(6.0, 3.0, 0.4, 0.4, kind="PICTURE (13)")
    copy = Box(5.0, 4.0, 3.0, 0.4, "a heading further down the same picture")

    assert lift_imagery(Slide([icon, copy], _full_page_picture()), CANVAS) == []
    assert icon.top_in() == 3.0


def test_a_graphic_only_touching_the_artwork_is_not_on_it() -> None:
    """Touching was the old test. An icon resting against the edge of a band
    is not invisible against it."""
    band = [Box(0.0, 0.0, 13.33, 3.0, kind="FREEFORM (5)")]
    icon = Box(4.0, 2.9, 0.5, 0.5, kind="PICTURE (13)")   # 0.1in of 0.5 inside
    copy = Box(3.5, 4.5, 3.0, 0.4, "a heading below it")

    assert lift_imagery(Slide([icon, copy], band), CANVAS) == []


def test_a_graphic_too_far_from_its_copy_is_not_labelling_it() -> None:
    """A label sits directly above what it labels. Anything that has to cross
    the page to find copy has found the footer."""
    band = [Box(0.0, 0.0, 13.33, 3.0, kind="FREEFORM (5)")]
    icon = Box(4.0, 1.0, 0.5, 0.5, kind="PICTURE (13)")
    far = Box(3.5, 6.8, 3.0, 0.4, "the source line at the bottom")

    assert lift_imagery(Slide([icon, far], band), CANVAS) == []


def test_an_icon_buried_in_a_band_is_still_lifted_clear_of_it() -> None:
    """The case the function was written for has to keep working: dark line
    art well inside a dark band, with the copy it labels below."""
    band = [Box(0.0, 0.0, 13.33, 3.3, kind="FREEFORM (5)")]
    icon = Box(4.0, 2.0, 0.5, 0.5, kind="PICTURE (13)")
    copy = Box(3.5, 4.0, 3.0, 0.4, "the heading it labels")

    moved = lift_imagery(Slide([icon, copy], band), CANVAS)

    assert len(moved) == 1
    assert icon.top_in() >= 3.3                 # out of the band
    assert icon.top / EMU + icon.height / EMU <= 4.0    # above its copy


def _straddling_row():
    """A row of four roundels on a band, over copy that reaches under only the
    first two of them. Judged one at a time, half the row moves."""
    band = [Box(0.0, 0.0, 13.33, 3.0, kind="FREEFORM (5)")]
    row = [Box(x, 2.0, 0.3, 0.3, kind="PICTURE (13)")
           for x in (3.0, 3.4, 3.8, 4.2)]
    copy = Box(2.5, 3.8, 1.4, 0.4, "the heading they sit over")
    return band, row, copy


def test_a_row_moves_together_or_not_at_all() -> None:
    """Three of five moved on the real slide, because separate roundels touch
    nothing and each was judged alone. A rearranged row is worse than either
    outcome: a reader cannot tell it from one that was always ragged."""
    band, row, copy = _straddling_row()

    moved = lift_imagery(Slide([*row, copy], band), CANVAS)

    assert len(moved) == 4                               # all of it, not some
    assert len({flag.top_in() for flag in row}) == 1     # one row, one answer


def test_without_the_row_held_together_it_would_be_torn_in_half(
    monkeypatch,
) -> None:
    """The bug this guards, shown rather than asserted about: with neighbours
    judged separately the same row comes out at two different heights."""
    from formatting_tool.rebuild import builder

    monkeypatch.setattr(builder, "_side_by_side", lambda box, other: False)
    band, row, copy = _straddling_row()

    builder.lift_imagery(Slide([*row, copy], band), CANVAS)

    assert len({flag.top_in() for flag in row}) == 2     # torn


# --------------------------------------------------------------------------- #
# 2. Reading order
# --------------------------------------------------------------------------- #

def _row(labels):
    """Six chevrons across the page. In Arabic the FIRST is the rightmost."""
    return [Box(11.0 - i * 1.85, 1.3, 1.75, 0.45, text)
            for i, text in enumerate(labels)]


LABELS = ["1-business-models", "2-strategy", "3-annual-plan",
          "4-execution", "5-monitoring", "6-review"]


def test_a_right_to_left_run_is_read_right_to_left() -> None:
    run = series.find_series(_row(LABELS), shortest=3, rtl=True)[0]

    assert [shape.text for shape in run.shapes] == LABELS


def test_a_left_to_right_run_is_unchanged() -> None:
    """An English deck must read exactly as it did before this existed."""
    run = series.find_series(_row(LABELS), shortest=3, rtl=False)[0]

    assert [shape.text for shape in run.shapes] == LABELS[::-1]


def test_the_two_directions_are_exact_opposites() -> None:
    row = _row(LABELS)
    ltr = series.find_series(row, shortest=3, rtl=False)[0]
    rtl = series.find_series(row, shortest=3, rtl=True)[0]

    assert [s.text for s in ltr.shapes] == [s.text for s in rtl.shapes][::-1]


def test_a_tall_list_still_reads_down_its_columns() -> None:
    """Right to left turns the columns round. It does not turn the rows round:
    an Arabic agenda still runs top to bottom."""
    # Two columns of three, the RIGHT column first in Arabic.
    cells = []
    for column, left in enumerate((6.0, 1.0)):          # right column first
        for row, top in enumerate((1.0, 2.0, 3.0)):
            cells.append(Box(left, top, 3.0, 0.6, f"c{column}r{row}"))

    run = series.find_series(cells, shortest=3, rtl=True)[0]

    assert [s.text for s in run.shapes] == [
        "c0r0", "c0r1", "c0r2", "c1r0", "c1r1", "c1r2",
    ]


def test_a_deck_of_arabic_copy_is_read_as_right_to_left() -> None:
    from formatting_tool.rebuild.builder import _reads_rtl

    class _Shape:
        def __init__(self, text):
            self._text = text
            self.has_text_frame = True

        @property
        def text_frame(self):
            return self

        @property
        def text(self):
            return self._text

    def deck(*texts):
        slide = type("S", (), {"shapes": [_Shape(t) for t in texts]})()
        return type("P", (), {"slides": [slide]})()

    assert _reads_rtl(deck("نماذج الأعمال", "تطوير الاستراتيجية", "Annexes"))
    assert not _reads_rtl(deck("Business models", "Strategy", "الملاحق"))
    assert not _reads_rtl(type("P", (), {"slides": []})())


# --------------------------------------------------------------------------- #
# 3. Copy deleted after going nowhere
# --------------------------------------------------------------------------- #

def test_copy_into_a_region_that_cannot_hold_it_reports_failure() -> None:
    """The caller deletes the box the copy came out of, so "it landed" and "it
    did not" have to be tellable apart."""
    from formatting_tool.rebuild.builder import _copy_text

    class _NoFrame:
        has_text_frame = False
        name = "Table Placeholder 4"

    source = Box(1.0, 1.0, 3.0, 0.5, "words worth keeping")
    assert _copy_text(source, _NoFrame()) is False


def test_copy_that_lands_reports_success() -> None:
    from formatting_tool.rebuild.builder import _copy_text
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    source = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    source.text_frame.text = "words worth keeping"
    target = slide.shapes.add_textbox(Inches(5), Inches(1), Inches(3), Inches(1))

    assert _copy_text(source, target) is True
    assert target.text_frame.text == "words worth keeping"
