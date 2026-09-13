"""Taking back one applied fix, without disturbing the others.

A designer applies forty corrections, looks at the renders, and wants one of
them back the way it was. Per change, not per slide: the slide usually carries
several fixes and only one of them is the argument.

UNDO IS A REPLAY, NOT A REVERSE, which is what these tests are really about.
The deck is written again from the untouched original without the fixes on the
undo list, rather than the written file being edited to drag one shape back.
The difference is visible in the third test here: a cohort move carries shapes
no finding named, so reversing "the fix" would put one shape back and leave
the rest of the column where the fix took it -- a state no run of this tool
would ever produce. Replaying puts the whole column back, because on that run
the move never happened.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formatting_tool.apply import apply_fixes
from formatting_tool.models import Category, Issue, Severity, Source

MARGINS = "top 0.4in, right 0.92in, bottom 0.73in, left 0.92in safe margin"


def _issue(rule_id: str, shape: str, shape_id: int, **kwargs) -> Issue:
    issue = Issue(
        category=Category.SPACE,
        severity=Severity.WARNING,
        message=kwargs.pop("message", f"{rule_id} finding"),
        source=Source.RULE,
        rule_id=rule_id,
        slide=1,
        shape=shape,
        shape_id=shape_id,
        deck="messy.pptx",
        **kwargs,
    )
    issue.id = issue.fingerprint()
    return issue


def _deck(tmp_path: Path, lefts: list[float]):
    """A slide of boxes at the given left edges, one per row."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    made = []
    for index, left in enumerate(lefts):
        box = slide.shapes.add_textbox(
            Inches(left), Inches(1.0 + index * 0.8), Inches(4.0), Inches(0.6)
        )
        box.name = f"Row {index}"
        made.append(box)
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))
    return deck, made


def _lefts(path: Path) -> dict[str, float]:
    from pptx import Presentation

    return {
        s.name: round(s.left / 914400, 2)
        for s in Presentation(str(path)).slides[0].shapes
    }


# --------------------------------------------------------------------------- #
# One change, taken back
# --------------------------------------------------------------------------- #

def test_an_undone_fix_is_not_applied_and_the_others_still_are(
    tmp_path: Path,
) -> None:
    """The whole feature in one case: two corrections, one taken back."""
    deck, made = _deck(tmp_path, [-1.0, -2.0])
    issues = [
        _issue("space.off_canvas", "Row 0", made[0].shape_id),
        _issue("space.off_canvas", "Row 1", made[1].shape_id),
    ]
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(deck, issues, out, selected=None, undone=[issues[0].id])

    assert len(result.applied) == 1
    assert result.undone == [issues[0].id]
    lefts = _lefts(out)
    assert lefts["Row 0"] == -1.0        # held back, exactly as it arrived
    assert lefts["Row 1"] == 0.0         # applied, as it would have been


def test_undoing_every_fix_gives_back_the_deck_as_it_arrived(
    tmp_path: Path,
) -> None:
    """Not an approximation of it. The original is never modified, so a run
    that applies nothing writes the file that was uploaded."""
    deck, made = _deck(tmp_path, [-1.0, -2.0])
    issues = [
        _issue("space.off_canvas", "Row 0", made[0].shape_id),
        _issue("space.off_canvas", "Row 1", made[1].shape_id),
    ]
    out = tmp_path / "fixed.pptx"
    before = _lefts(deck)

    result = apply_fixes(
        deck, issues, out, selected=None, undone=[i.id for i in issues]
    )

    assert result.applied == []
    assert _lefts(out) == before


def test_undoing_a_set_move_puts_the_whole_set_back(tmp_path: Path) -> None:
    """The case a reverse cannot do. One finding named "Row 0"; the fix moved
    three shapes, because a column moves whole or not at all. Undoing it has
    to put all three back, and a replay does that without being told: on that
    run the move never happened."""
    deck, made = _deck(tmp_path, [0.48, 0.48, 0.48])
    issue = _issue(
        "space.safe_margin", "Row 0", made[0].shape_id, expected=MARGINS
    )
    out = tmp_path / "fixed.pptx"

    applied = apply_fixes(deck, [issue], out, selected=[issue.id])
    assert _lefts(out) == {"Row 0": 0.92, "Row 1": 0.92, "Row 2": 0.92}

    undone = apply_fixes(
        deck, [issue], out, selected=[issue.id], undone=[issue.id]
    )

    assert undone.applied == []
    assert _lefts(out) == {"Row 0": 0.48, "Row 1": 0.48, "Row 2": 0.48}
    assert applied.output == undone.output      # the same file, written again


def test_undoing_one_of_several_leaves_the_rest_where_the_fixes_put_them(
    tmp_path: Path,
) -> None:
    """Two shapes moved by two findings, one taken back. The fix that stands
    is not re-measured against the undone one: each run is measured against
    the original deck, which is the state the report describes."""
    deck, made = _deck(tmp_path, [-1.0, -1.5])
    issues = [
        _issue("space.off_canvas", "Row 0", made[0].shape_id),
        _issue("space.off_canvas", "Row 1", made[1].shape_id),
    ]
    out = tmp_path / "fixed.pptx"

    apply_fixes(deck, issues, out, selected=None)
    assert _lefts(out) == {"Row 0": 0.0, "Row 1": 0.0}

    apply_fixes(deck, issues, out, selected=None, undone=[issues[1].id])

    assert _lefts(out) == {"Row 0": 0.0, "Row 1": -1.5}


def test_an_undo_can_be_taken_back_by_not_asking_for_it(tmp_path: Path) -> None:
    """Redo, which needs no machinery of its own: the fix is on the next run
    unless the next run is told to hold it back."""
    deck, made = _deck(tmp_path, [-1.0])
    issue = _issue("space.off_canvas", "Row 0", made[0].shape_id)
    out = tmp_path / "fixed.pptx"

    apply_fixes(deck, [issue], out, selected=None, undone=[issue.id])
    assert _lefts(out)["Row 0"] == -1.0

    apply_fixes(deck, [issue], out, selected=None, undone=[])

    assert _lefts(out)["Row 0"] == 0.0


def test_the_original_deck_is_still_never_modified(tmp_path: Path) -> None:
    """The property the whole design rests on: undo replays from this file, so
    if a run could touch it there would be nothing to replay from."""
    deck, made = _deck(tmp_path, [-1.0])
    issue = _issue("space.off_canvas", "Row 0", made[0].shape_id)
    before = deck.read_bytes()

    apply_fixes(deck, [issue], tmp_path / "a.pptx", selected=None)
    apply_fixes(
        deck, [issue], tmp_path / "b.pptx", selected=None, undone=[issue.id]
    )

    assert deck.read_bytes() == before


# --------------------------------------------------------------------------- #
# Ids that name nothing
# --------------------------------------------------------------------------- #

def test_an_undo_id_that_holds_nothing_back_is_reported_not_raised(
    tmp_path: Path,
) -> None:
    """A page with a button on it will undo the same change twice. That is not
    an error and must not cost the run -- but it is not silence either: an id
    matching nothing usually means the selection and the report have come
    apart, and the designer is looking at a change that is still there."""
    deck, made = _deck(tmp_path, [-1.0])
    issue = _issue("space.off_canvas", "Row 0", made[0].shape_id)
    out = tmp_path / "fixed.pptx"

    result = apply_fixes(
        deck, [issue], out, selected=None, undone=[issue.id, "deadbeef"]
    )

    assert result.undone_unknown == ["deadbeef"]
    assert _lefts(out)["Row 0"] == -1.0          # the real one still held back


def test_a_selection_id_that_names_nothing_is_still_an_error(
    tmp_path: Path,
) -> None:
    """Undo is forgiving about an id; the tick list is not, and neither has
    changed. Applying the rest of a selection that has come apart from its
    report would hide that it had."""
    from formatting_tool.apply import ApplyError

    deck, made = _deck(tmp_path, [-1.0])
    issue = _issue("space.off_canvas", "Row 0", made[0].shape_id)

    with pytest.raises(ApplyError, match="no finding with id"):
        apply_fixes(deck, [issue], tmp_path / "out.pptx", selected=["deadbeef"])
