"""The undo the designer actually presses: /api/undo, over the wire.

The applier's half of this is covered in `test_undo`. What is left is the part
the page depends on and nothing else tests: that the session remembers the run
it made, so taking one change back replays THAT run without it rather than
some other one, and that undone changes accumulate while a restore takes one
off the list.

The server is started on an ephemeral port with a session put straight into
its table. Going through the upload and the check first would be a slower test
of `validate`, which has its own.
"""

from __future__ import annotations

import json
import threading
import urllib.request
from pathlib import Path

import pytest

from formatting_tool.models import Category, Issue, Severity, Source


def _issue(shape: str, shape_id: int) -> Issue:
    issue = Issue(
        category=Category.SPACE,
        severity=Severity.WARNING,
        message="space.off_canvas finding",
        source=Source.RULE,
        rule_id="space.off_canvas",
        slide=1,
        shape=shape,
        shape_id=shape_id,
        deck="messy.pptx",
    )
    issue.id = issue.fingerprint()
    return issue


@pytest.fixture()
def wired(tmp_path: Path):
    """A running server with one session on it, and the two findings in it."""
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    from formatting_tool.web import server as web

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = []
    for index, left in enumerate((-1.0, -2.0)):
        box = slide.shapes.add_textbox(
            Inches(left), Inches(1.0 + index), Inches(4.0), Inches(0.6)
        )
        box.name = f"Row {index}"
        shapes.append(box)
    deck = tmp_path / "messy.pptx"
    prs.save(str(deck))

    issues = [_issue(s.name, s.shape_id) for s in shapes]
    session = web._Session(
        id="test-session",
        directory=tmp_path,
        master=tmp_path / "master.pptx",
        deck=deck,
        report={"issues": [i.to_dict() for i in issues]},
    )
    web._SESSIONS[session.id] = session

    server = web._Server(("127.0.0.1", 0), web._Handler, tmp_path)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{server.server_address[1]}"
    try:
        yield base, session, issues
    finally:
        server.shutdown()
        server.server_close()
        web._SESSIONS.pop(session.id, None)


def _post(base: str, route: str, body: dict) -> tuple[int, dict]:
    request = urllib.request.Request(
        f"{base}{route}",
        data=json.dumps(body).encode(),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request) as response:
            return response.status, json.loads(response.read())
    except urllib.error.HTTPError as error:      # the page reads these too
        return error.code, json.loads(error.read())


def _lefts(path: Path) -> dict[str, float]:
    from pptx import Presentation

    return {
        s.name: round(s.left / 914400, 2)
        for s in Presentation(str(path)).slides[0].shapes
    }


def test_one_change_is_taken_back_and_the_rest_stand(wired) -> None:
    base, session, issues = wired

    code, applied = _post(
        base, "/api/apply", {"session": session.id, "fix": [i.id for i in issues]}
    )
    assert code == 200 and len(applied["applied"]) == 2

    code, undone = _post(
        base, "/api/undo", {"session": session.id, "undo": [issues[0].id]}
    )

    assert code == 200
    assert len(undone["applied"]) == 1
    assert [u["id"] for u in undone["undone"]] == [issues[0].id]
    # The held-back change is described from the report, since this run never
    # made it and so has nothing of its own to say about it.
    assert undone["undone"][0]["shape"] == "Row 0"
    assert _lefts(session.fixed) == {"Row 0": -1.0, "Row 1": 0.0}


def test_undone_changes_accumulate_and_a_restore_takes_one_off(wired) -> None:
    """A designer takes one back, looks, takes another back. The second undo
    must not quietly reinstate the first."""
    base, session, issues = wired
    _post(base, "/api/apply", {"session": session.id, "fix": [i.id for i in issues]})

    _post(base, "/api/undo", {"session": session.id, "undo": [issues[0].id]})
    _, both = _post(base, "/api/undo", {"session": session.id, "undo": [issues[1].id]})

    assert {u["id"] for u in both["undone"]} == {i.id for i in issues}
    assert both["applied"] == []
    assert _lefts(session.fixed) == {"Row 0": -1.0, "Row 1": -2.0}

    _, back = _post(base, "/api/undo", {"session": session.id, "redo": [issues[0].id]})

    assert [u["id"] for u in back["undone"]] == [issues[1].id]
    assert _lefts(session.fixed) == {"Row 0": 0.0, "Row 1": -2.0}


def test_applying_again_starts_a_fresh_run(wired) -> None:
    """The tick list the designer just sent is what they want. Carrying an
    earlier undo forward would silently drop a fix they can see they asked
    for."""
    base, session, issues = wired
    _post(base, "/api/apply", {"session": session.id, "fix": [i.id for i in issues]})
    _post(base, "/api/undo", {"session": session.id, "undo": [issues[0].id]})

    _, again = _post(
        base, "/api/apply", {"session": session.id, "fix": [i.id for i in issues]}
    )

    assert again["undone"] == []
    assert len(again["applied"]) == 2
    assert _lefts(session.fixed) == {"Row 0": 0.0, "Row 1": 0.0}


def test_the_renders_of_the_corrected_deck_are_dropped_when_it_is_rewritten(
    wired,
) -> None:
    """A stale after image is the worst thing this page can show: the renders
    are the evidence a designer accepts the result on. `_render_into` caches
    because driving PowerPoint costs seconds, so a run that writes the file
    again has to throw that cache away -- true of an undo, and already true of
    a second apply."""
    base, session, issues = wired
    _post(base, "/api/apply", {"session": session.id, "fix": [i.id for i in issues]})
    after = session.directory / "after"
    after.mkdir(exist_ok=True)
    (after / "1.png").write_bytes(b"a render of a deck that no longer exists")

    _post(base, "/api/undo", {"session": session.id, "undo": [issues[0].id]})

    assert not (after / "1.png").exists()
    # The before images are of the upload, which no run touches, so they stay.
    before = session.directory / "before"
    before.mkdir(exist_ok=True)
    (before / "1.png").write_bytes(b"still true")
    _post(base, "/api/undo", {"session": session.id, "undo": [issues[1].id]})
    assert (before / "1.png").exists()


def test_every_change_the_page_lists_carries_an_id_to_undo_by() -> None:
    """A second-round change is measured on the written deck by `_recheck`,
    which builds findings rather than reading them off a report, so it has no
    id of its own. The page was sent null for those, and an Undo addressed to
    null does nothing when it is pressed -- silently, which from where the
    designer is standing is a broken button."""
    from formatting_tool.apply import FixOutcome
    from formatting_tool.web.server import _outcome_id

    issue = _issue("Row 0", 7)
    issue.id = None

    key = _outcome_id(FixOutcome(issue=issue, applied=True, detail="moved it"))

    assert key
    # The same id `apply_fixes` matches a second-round finding on, so undoing
    # by it holds back the change the page is pointing at.
    assert key == issue.fingerprint()


def test_a_forced_render_ignores_what_is_cached(wired, monkeypatch) -> None:
    """The cache is right nearly always, and "nearly always" is not something
    a designer can check from the outside. The button is the way to say: render
    this one again, now."""
    from formatting_tool.web import server as web

    base, session, issues = wired
    _post(base, "/api/apply", {"session": session.id, "fix": [i.id for i in issues]})

    asked = []

    def fake_render(deck, slides=None):
        from formatting_tool.render import SlideImages

        asked.append(sorted(slides or []))
        directory = session.directory / "rendered"
        directory.mkdir(exist_ok=True)
        made = {}
        for number in (slides or [1]):
            path = directory / f"Slide{number}.PNG"
            path.write_bytes(b"fresh")
            made[number] = path
        return SlideImages(renderer="fake", directory=directory, images=made)

    monkeypatch.setattr(web, "render_deck", fake_render)

    stale = session.directory / "after"
    stale.mkdir(exist_ok=True)
    (stale / "1.png").write_bytes(b"stale")

    code, payload = _post(
        base, "/api/preview", {"session": session.id, "slides": [1], "force": [1]}
    )

    assert code == 200
    assert payload["slides"][0]["slide"] == 1
    assert (stale / "1.png").read_bytes() == b"fresh"
    assert asked == [[1], [1]]      # the before pair and the after pair


def test_an_undo_before_anything_was_applied_is_refused(wired) -> None:
    base, session, issues = wired

    code, payload = _post(
        base, "/api/undo", {"session": session.id, "undo": [issues[0].id]}
    )

    assert code == 400
    assert "nothing has been applied" in payload["error"]
