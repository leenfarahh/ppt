"""A transform written onto a placeholder that was inheriting one.

A placeholder that states nothing takes everything from its layout. The moment
anything is written onto it that stops being true for the part that was
written, and stays true for the part that was not -- only now there is an
`a:xfrm` saying otherwise. Set a placeholder's height and python-pptx writes an
extent with no offset, and PowerPoint reads a missing offset as (0, 0).

On a real deck the title of a page did exactly that: `space.text_overflow` grew
it to fit its own copy, and it arrived in the top-left corner with its first
letter off the edge of the slide.

What makes it worth a test of its own is that nothing could see it. python-pptx
resolves a missing offset by walking to the layout, so every check made through
it read the inherited 0.68in and reported no movement. The file said one thing
and the object model said another.
"""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches

from formatting_tool.apply.applier import _whole_box

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _title_of(deck):
    return deck.slides[0].shapes.title


def _xfrm(shape):
    spPr = shape._element.spPr
    return None if spPr is None else spPr.find(A_NS + "xfrm")


def _inheriting_title():
    """A title placeholder that states no box of its own."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "A title"
    return presentation


def test_setting_only_the_size_leaves_half_a_transform() -> None:
    """The defect, reproduced against python-pptx itself so the test fails if
    the library ever starts writing the whole box."""
    title = _title_of(_inheriting_title())
    assert _xfrm(title) is None

    title.height = Inches(1.2)

    xfrm = _xfrm(title)
    assert xfrm is not None
    assert xfrm.find(A_NS + "ext") is not None
    assert xfrm.find(A_NS + "off") is None      # the missing half


def test_the_missing_offset_is_written_from_what_it_inherited() -> None:
    title = _title_of(_inheriting_title())
    inherited = (title.left, title.top, title.width, title.height)

    title.height = Inches(1.2)
    _whole_box(title, inherited)

    off = _xfrm(title).find(A_NS + "off")
    assert off is not None
    assert int(off.get("x")) == inherited[0]
    assert int(off.get("y")) == inherited[1]
    # And the size the fix asked for is untouched.
    assert title.height == Inches(1.2)


def test_the_missing_extent_is_written_too() -> None:
    title = _title_of(_inheriting_title())
    inherited = (title.left, title.top, title.width, title.height)

    title.left = Inches(2.0)
    _whole_box(title, inherited)

    ext = _xfrm(title).find(A_NS + "ext")
    assert ext is not None
    assert int(ext.get("cx")) == inherited[2]
    assert int(ext.get("cy")) == inherited[3]
    assert title.left == Inches(2.0)


def test_a_shape_that_states_nothing_is_left_stating_nothing() -> None:
    """Inheriting everything is correct and must not be written out: doing so
    would freeze the layout's geometry onto every slide it touched."""
    title = _title_of(_inheriting_title())

    _whole_box(title, (title.left, title.top, title.width, title.height))

    assert _xfrm(title) is None


def test_a_complete_transform_is_left_alone() -> None:
    title = _title_of(_inheriting_title())
    title.left, title.top = Inches(1.0), Inches(1.0)
    title.width, title.height = Inches(4.0), Inches(1.0)
    before = _xfrm(title)

    _whole_box(title, (0, 0, 0, 0))

    assert int(before.find(A_NS + "off").get("x")) == Inches(1.0)
    assert int(before.find(A_NS + "ext").get("cx")) == Inches(4.0)
