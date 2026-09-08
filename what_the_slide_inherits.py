"""What a slide draws itself, and what it inherits from its layout.

    python what_the_slide_inherits.py DECK.pptx [MORE.pptx ...]

What a designed slide shows is often not on the slide. A section divider can
carry a full-bleed photograph, the connector lines joining its icons and the
panels behind its copy while its own shape tree holds nothing but a title --
all of it inherited from the layout. Point that slide at a different layout
and every one of those is gone, with nothing to report, because no shape was
lost: there was never a shape.

Run it on a messy deck to see how much of it is inherited, and on a rebuilt
one to check that what was inherited came across.
"""

import sys
from pathlib import Path

from pptx import Presentation

PICTURE = 13
GROUP = 6


def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == GROUP:
            yield from walk(shape.shapes)


def is_placeholder(shape) -> bool:
    try:
        return bool(shape.is_placeholder)
    except Exception:
        return False


def describe(shape) -> str:
    if is_placeholder(shape):
        token = str(shape.placeholder_format.type)
        holds = shape._element.find(
            ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
        ) is not None
        return f"placeholder {token} ({'holds an image' if holds else 'empty'})"
    return str(shape.shape_type)


def main(path: str) -> None:
    prs = Presentation(path)
    name = Path(path).name
    print(f"\n{name}\n" + "=" * len(name))

    own_total = 0
    inherited_total = 0
    for number, slide in enumerate(prs.slides, 1):
        own = [s for s in walk(slide.shapes) if not is_placeholder(s)]
        layout = slide.slide_layout
        inherited = [s for s in layout.shapes if not is_placeholder(s)]
        own_total += len(own)
        inherited_total += len(inherited)

        print(
            f"  slide {number:>3}  layout {layout.name!r:<38} "
            f"{len(own)} of its own, {len(inherited)} inherited"
        )
        for shape in own:
            if shape.shape_type == PICTURE:
                print(f"           - {shape.name!r}: {describe(shape)}")
        for shape in inherited:
            print(f"           ~ {shape.name!r}: {describe(shape)}  ON THE LAYOUT")

    print(
        f"\n  TOTAL  {own_total} shape(s) the slides draw themselves, "
        f"{inherited_total} inherited from layouts"
    )
    print(
        "\n  A shape on a slide survives a rebuild on its own. One on a layout"
        "\n  is inherited, and is copied onto the slide before the layouts swap"
        "\n  -- unless it repeats across layouts, which makes it brand furniture."
    )


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit("usage: python what_the_slide_inherits.py DECK.pptx [...]")
    for argument in sys.argv[1:]:
        main(argument)
